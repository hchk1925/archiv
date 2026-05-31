#!/usr/bin/env python3
"""
fill_remaining_fate.py — Doplnit zbývající prázdná season_fate.

Pravidla:
  - T-řádek v poslední sezóně (S2013_14) bez next → 'setrval?'
  - T-řádek mimo poslední sezónu bez next-matche + bez fate → 'zanik'
    (klub mizí; konzervativní default, dle §8)
  - Vynechat KVAL/L15/L25/L35/L45 (D30 — má být prázdné)
"""
import openpyxl, glob, os, re, sys
from collections import defaultdict

DATA = sys.argv[1] if len(sys.argv) > 1 else 'data'
NON_DATA = {'NOTES', 'SYSTEM', 'SERIES', 'CLUBS', 'META', 'PYRAMIDA', 'TODO'}
LEVEL_RE = re.compile(r'L(\d+)')


def hidx(row0):
    return {h: i for i, h in enumerate(row0) if h is not None}


def lvl(v):
    if v is None:
        return None
    m = LEVEL_RE.search(str(v))
    return int(m.group(1)) if m else None


def season_of(p):
    m = re.search(r'S(\d{4})_(\d{2})', os.path.basename(p))
    return f"S{m.group(1)}_{m.group(2)}"


def load_clubs(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['CLUBS']
    rows = list(ws.iter_rows(values_only=True))
    H = hidx(rows[0])
    out = {}
    for r in rows[1:]:
        if not r or r[H['club_id']] is None:
            continue
        out[r[H['club_id']]] = {
            'prev': r[H['prev_club_id']] if 'prev_club_id' in H else None,
        }
    wb.close()
    return out


def main():
    files = sorted(glob.glob(os.path.join(DATA, 'S*_FINAL.xlsx')))
    seasons = [season_of(p) for p in files]
    s_by = dict(zip(seasons, files))
    last_season = seasons[-1]
    # mapa next-season prev sets per sezóna
    next_prevs = {}
    for i in range(len(seasons) - 1):
        nxt = load_clubs(s_by[seasons[i + 1]])
        next_prevs[seasons[i]] = {nc['prev'] for nc in nxt.values() if nc['prev']}
    next_prevs[last_season] = None

    tot_setr = tot_zanik = 0
    for path in files:
        sid = season_of(path)
        nps = next_prevs[sid]
        wb = openpyxl.load_workbook(path)
        changed = False
        for sh in wb.sheetnames:
            if sh in NON_DATA:
                continue
            if sh == 'KVAL' or sh.startswith('KVAL'):
                continue
            dws = wb[sh]
            rows = list(dws.iter_rows())
            if not rows:
                continue
            DH = hidx([c.value for c in rows[0]])
            if 'club_id' not in DH or 'season_fate' not in DH:
                continue
            rti = DH.get('row_type')
            sfi, dci = DH['season_fate'], DH['club_id']
            lvi = DH.get('level')
            for row in rows[1:]:
                if rti is not None and row[rti].value != 'T':
                    continue
                if row[sfi].value:
                    continue
                cid = row[dci].value
                if not cid:
                    continue
                ln = lvl(row[lvi].value) if lvi is not None else None
                if ln in (15, 25, 35, 45):
                    continue
                if nps is None:
                    row[sfi].value = 'setrval?'
                    tot_setr += 1
                elif cid not in nps:
                    row[sfi].value = 'zanik'
                    tot_zanik += 1
                else:
                    # má next ale fate prázdné — fallback setrval
                    row[sfi].value = 'setrval'
                    tot_setr += 1
                changed = True
        if changed:
            wb.save(path)
        wb.close()
    print(f"=== fill_remaining_fate ===")
    print(f"  setrval/setrval?: {tot_setr}")
    print(f"  zanik:            {tot_zanik}")


if __name__ == '__main__':
    main()
