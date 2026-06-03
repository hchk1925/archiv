#!/usr/bin/env python3
"""
fix_1949_50_top.py — oprava nejvyšší soutěže 1949/50 proti PDF (in-place).

PDF (Celostátní mistrovská soutěž): ATK Praha 1. místo, 24 bodů = MISTR. Čísla
v datech sedí, ale M badge byl chybně u Zdar LTC Praha (3. místo) — přesun na ATK.
Plus oficiální názvy do raw_name. Idempotentní.
"""
import openpyxl

PATH = 'data/S1949_50_FINAL.xlsx'

# clean_name → (oficiální raw_name z PDF, badge)
RAW = {
    'ATK Praha': ('ATK Praha [M]', 'M'),
    'Vítkovické Železárny': ('ZSJ Vítkovické železárny [N]', 'N'),
    'Zdar LTC Praha': ('ZSJ Zdar LTC Praha', ''),
    'Stadion České Budějovice': ('ZSJ OD Stadion České Budějovice', ''),
    'GZ Královo Pole': ('ZSJ GZ Královo Pole [N]', 'N'),
    'NV Bratislava': ('ZJ Sokol NV Bratislava', ''),
    'Bratrství Sparta Praha': ('ZSJ Bratrství Sparta', ''),
    'Zbrojovka Brno': ('ZSJ Zbrojovka Brno-Židenice', ''),
}


def main():
    wb = openpyxl.load_workbook(PATH)
    # 10_liga: M badge → ATK, pryč ze Zdar LTC
    ws = wb['10_liga']
    H = {c.value: i for i, c in enumerate(ws[1])}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, H['row_type'] + 1).value != 'T':
            continue
        nm = ws.cell(r, H['club_name'] + 1).value
        if nm == 'ATK Praha':
            ws.cell(r, H['note'] + 1).value = 'M'
        elif nm == 'Zdar LTC Praha':
            ws.cell(r, H['note'] + 1).value = None
    # CLUBS: raw_name + entry_note (M/N)
    cl = wb['CLUBS']
    Hc = {c.value: i for i, c in enumerate(cl[1])}
    for r in range(2, cl.max_row + 1):
        nm = cl.cell(r, Hc['clean_name'] + 1).value
        if nm in RAW:
            raw, badge = RAW[nm]
            cl.cell(r, Hc['raw_name'] + 1).value = raw
            if 'entry_note' in Hc:
                cl.cell(r, Hc['entry_note'] + 1).value = badge or None
    # SYSTEM node 0001: oficiální název do note
    sy = wb['SYSTEM']
    for r in range(2, sy.max_row + 1):
        if sy.cell(r, 1).value == 'NODE_S1949_50_0001':
            ni = [c.value for c in sy[1]].index('note')
            sy.cell(r, ni + 1).value = 'Oficiální název 1949/50: Celostátní mistrovská soutěž (CMS).'
    # META
    me = wb['META']
    mk = {me.cell(r, 1).value: r for r in range(2, me.max_row + 1)}
    def setm(k, v):
        if k in mk: me.cell(mk[k], 2).value = v
        else: me.append([k, v])
    setm('mistr', 'ATK Praha')
    setm('source', 'PDF sources/CZE1 + CZE2_3_nizsi (dobový tisk)')
    setm('note', 'Mistr ATK Praha (1. místo CMS, 24 b). M badge opraven (byl chybně '
                 'u Zdar LTC Praha, 3. místo). Navázáno na S1948_49 (prev_club_id doplněn). '
                 'Nováčci: Vítkovice, GZ Královo Pole (z divize).')
    # NOTES
    no = wb['NOTES']
    note = ('Nejvyšší soutěž 1949/50 = „Celostátní mistrovská soutěž (CMS)". Mistr '
            'ATK Praha (1. místo). LTC Praha (3.) byl v té době rozpouštěn — část '
            'reprezentantů zatčena v procesu 1950.')
    if no.max_row >= 1 and note not in {no.cell(r, 1).value for r in range(1, no.max_row + 1)}:
        no.append([note])
    wb.save(PATH)
    print("✔ 1949/50 nejvyšší soutěž: M badge → ATK Praha, oficiální názvy, META/NOTES")


if __name__ == '__main__':
    main()
