import openpyxl, json

fn = 'work/S1988_89_FINAL.xlsx'
wb = openpyxl.load_workbook(fn, data_only=True, read_only=True)
nodes={}
for r in wb['SYSTEM'].iter_rows(values_only=True):
    if not r or not r[0] or r[0]=='node_id': continue
    nodes[r[0]]={'name':r[1],'type':r[2],'level':r[3]}

def parse_sheet(sn):
    ws=wb[sn];blocks=[];cur=None
    for r in ws.iter_rows(values_only=True):
        if not r or not r[0] or r[0]=='row_type': continue
        if r[0]=='H':
            if cur and cur['rows']: blocks.append(cur)
            cur={'name':r[1] or '','rows':[]}
        elif r[0]=='T' and cur is not None:
            cur['rows'].append({'pos':r[2],'club':r[3],'note':r[4] or '','gp':r[5],'w':r[6],'d':r[7],'l':r[8],'gf':r[9],'ga':r[11],'body':r[12],'fate':r[20] or ''})
    if cur and cur['rows']: blocks.append(cur)
    return blocks

# Play-off serie
series_raw={}
for r in wb['SERIES'].iter_rows(values_only=True):
    if not r or not r[0] or r[0]=='series_id': continue
    sid=r[0];nid=r[1];name=r[4];sscore=r[7];gscore=r[6]
    if sid not in series_raw:
        ni=nodes.get(nid,{}); series_raw[sid]={'node':ni.get('name','?'),'level':ni.get('level',''),'teams':[]}
    series_raw[sid]['teams'].append({'name':name,'sscore':sscore,'games':gscore})

extra = parse_sheet('10_liga')
po_series=[{'node':s['node'],'t1':s['teams'][0],'t2':s['teams'][1]} for s in series_raw.values() if s['level']=='L10' and len(s['teams'])>=2]

# === ODVOZENI NAVAZNOSTI: kdo je v play-off, kdo v udrzeni ===
playoff_clubs=set()
for s in po_series:
    if 'tvrtfinále' in s['node']:  # ctvrtfinale = vstup do play-off
        playoff_clubs.add(s['t1']['name']); playoff_clubs.add(s['t2']['name'])
udrzeni_clubs=set()
for b in extra:
    if 'udržení' in b['name'].lower() and 'Prolínací' not in b['name']:
        for r in b['rows']: udrzeni_clubs.add(r['club'])

def dest_label(club):
    if club in playoff_clubs: return ('play-off', 'd-po')
    if club in udrzeni_clubs: return ('skupina o udržení', 'd-rel')
    return ('', '')

with open('/tmp/s1988_data.json','w') as f:
    json.dump({'extra':extra,'po':po_series,
               'playoff_clubs':list(playoff_clubs),'udrzeni_clubs':list(udrzeni_clubs)}, f, ensure_ascii=False)

print("Play-off kluby:", sorted(playoff_clubs))
print("\nUdrzeni kluby:", sorted(udrzeni_clubs))
print("\n=== Zakl.cast s navaznosti ===")
for b in extra:
    if 'ákladní' in b['name']:
        for r in b['rows']:
            d,_=dest_label(r['club'])
            print(f"  {r['pos']}. {r['club'][:28]:<28} → {d}")
wb.close()
