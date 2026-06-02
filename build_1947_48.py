#!/usr/bin/env python3
"""
build_1947_48.py — postaví S1947_48_FINAL.xlsx z PDF sources/CZE1/CZ-1947-48.pdf
a naváže chain dopředu na existující S1948_49.

Rozsah: nejvyšší soutěž (Státní liga, 2 skupiny po 6 + finále) — jako S1948_49
je i tato sezóna PARTIAL (republiková úroveň). Nižší/regiony = pozdější chirurgie.

Ověřeno z PDF: sumy GF=GA (A 172=172, B 161=161), body=2·V+R sedí, mistr LTC Praha
(finále 7:1 a 13:5 nad I. ČLTK).
"""
import openpyxl
from openpyxl.utils import get_column_letter

S = 'S1947_48'
OUT = f'data/{S}_FINAL.xlsx'

# (pos, clean_name, GP,W,D,L,GF,GA,PTS, city)
GROUP_A = [
    (1, 'I. ČLTK Praha', 5, 5, 0, 0, 61, 6, 10, 'Praha'),
    (2, 'HC Stadion Podolí', 5, 3, 0, 2, 52, 21, 6, 'Praha'),
    (3, 'SK Prostějov', 5, 3, 0, 2, 28, 24, 6, 'Prostějov'),
    (4, 'SK Horácká Slavia Třebíč', 5, 2, 1, 2, 13, 29, 5, 'Třebíč'),
    (5, 'VŠ Bratislava', 5, 1, 1, 3, 12, 29, 3, 'Bratislava'),
    (6, 'ŠK Banská Bystrica', 5, 0, 0, 5, 6, 63, 0, 'Banská Bystrica'),
]
GROUP_B = [
    (1, 'LTC Praha', 5, 5, 0, 0, 55, 7, 10, 'Praha'),
    (2, 'ŠK Bratislava', 5, 3, 1, 1, 28, 21, 7, 'Bratislava'),
    (3, 'AC Sparta', 5, 3, 0, 2, 23, 18, 6, 'Praha'),
    (4, 'AC Stadion České Budějovice', 5, 2, 1, 2, 32, 16, 5, 'České Budějovice'),
    (5, 'ČSK Říčany', 5, 1, 0, 4, 10, 31, 2, 'Říčany'),
    (6, 'BK Havlíčkův Brod', 5, 0, 0, 5, 13, 68, 0, 'Havlíčkův Brod'),
]
# fate (dle doložené návaznosti na 1948/49 — viz META.note tam); '' = neurčeno (pozdější)
FATE = {
    'LTC Praha': 'setrval', 'ŠK Bratislava': 'setrval', 'AC Sparta': 'setrval',
    'AC Stadion České Budějovice': 'setrval', 'SK Prostějov': 'setrval',
    'I. ČLTK Praha': 'setrval', 'SK Horácká Slavia Třebíč': 'sloučení',
    'HC Stadion Podolí': 'zánik',
}

HDR10 = ['row_type', 'block_name', 'pos', 'club_name', 'note', 'GP', 'W', 'D', 'L',
         'GF', ':', 'GA', 'PTS', 'comp_path', 'node_id', 'level', 'club_id',
         'prev_club_id', 'dest_node_id', 'dest_type', 'season_fate', 'tr_id', 'district']


def cid(n):
    return f'CLUB_{S}_{n:04d}'


def build():
    wb = openpyxl.Workbook()
    # club_id mapping: A 1..6 -> 0001..0006, B 1..6 -> 0007..0012
    clubs = []           # (club_id, clean_name, city, level, fate)
    rows10 = [HDR10]
    tr = 1
    cidx = 1
    NODE_A = f'NODE_{S}_0002'
    NODE_B = f'NODE_{S}_0003'

    def add_group(block, node, teams):
        nonlocal tr, cidx
        rows10.append(['H', block, '', '', '', '', '', '', '', '', '', '', '', '',
                       node, 'L10', '', '', '', '', '', '', ''])
        for (pos, nm, gp, w, d, l, gf, ga, pts, city) in teams:
            c = cid(cidx)
            note = 'M' if nm == 'LTC Praha' else ''
            fate = FATE.get(nm, '')
            rows10.append(['T', '', pos, nm, note, gp, w, d, l, gf, ':', ga, pts,
                           block, node, 'L10', c, '', '', '', fate,
                           f'TR_{S}_{tr:05d}', ''])
            clubs.append((c, nm, city, 'L10', fate))
            tr += 1
            cidx += 1

    add_group('Státní liga / Skupina A', NODE_A, GROUP_A)
    add_group('Státní liga / Skupina B', NODE_B, GROUP_B)

    ws = wb.active
    ws.title = '10_liga'
    for r in rows10:
        ws.append(r)

    # SERIES — finále o titul
    se = wb.create_sheet('SERIES')
    se.append(['series_id', 'node_id', 'club_id', 'raw_name', 'clean_name', 'side',
               'game_scores', 'series_score', 'dest_node_id', 'note'])
    fin = f'NODE_{S}_0004'
    ltc = cid(7); clt = cid(1)
    se.append([f'SER_{S}_0001', fin, ltc, 'LTC Praha', 'LTC Praha', '1',
               '7:1, 13:5', '2:0', '', 'Mistr 1947/48'])
    se.append([f'SER_{S}_0001', fin, clt, 'I. ČLTK Praha', 'I. ČLTK Praha', '2',
               '1:7, 5:13', '0:2', '', ''])

    # SYSTEM
    sy = wb.create_sheet('SYSTEM')
    sy.append(['node_id', 'name', 'competition_type', 'level', 'region',
               'parent_node_id', 'feeds_into', 'feeds_into_loser', 'scoring',
               'status', 'note', 'prev_node_id'])
    root = f'NODE_{S}_0001'
    sy.append([root, 'Státní liga', 'league', 'L10', 'REG_CS', '', '', '', '2-1-0', '', '', ''])
    sy.append([NODE_A, 'Státní liga / Skupina A', 'group', 'L10', 'REG_CS', root, '', '', '2-1-0', '', '', ''])
    sy.append([NODE_B, 'Státní liga / Skupina B', 'group', 'L10', 'REG_CS', root, '', '', '2-1-0', '', '', ''])
    sy.append([fin, 'Státní liga / finále', 'final_group', 'L10', 'REG_CS', root, '', '', '2-1-0', '', 'LTC Praha – I. ČLTK Praha 7:1, 13:5', ''])

    # CLUBS (formát jako S1948_49: club_id, clean_name, raw_name, sheet, entry_note,
    #        level, district, prev_club_id, change_note, city)
    cl = wb.create_sheet('CLUBS')
    cl.append(['club_id', 'clean_name', 'raw_name', 'sheet', 'entry_note', 'level',
               'district', 'prev_club_id', 'change_note', 'city'])
    for (c, nm, city, lvl, fate) in clubs:
        raw = nm + (' [M]' if nm == 'LTC Praha' else '')
        note = 'Nejvyšší soutěž z PDF (dobový tisk). prev_club_id (na 1946/47) zatím prázdný — starší sezóny přijdou později.'
        cl.append([c, nm, raw, '10_liga', 'M' if nm == 'LTC Praha' else '', lvl,
                   '', '', note, city])

    # META
    me = wb.create_sheet('META')
    me.append(['key', 'value'])
    for k, v in [
        ('season_id', S), ('season_label', '1947/48'),
        ('era_note', 'Československo, Státní liga (5. ročník). 2 skupiny po 6 + finále o titul.'),
        ('scoring_system', '2-1-0'),
        ('source', 'PDF sources/CZE1/CZ-1947-48.pdf (dobový tisk)'),
        ('total_clubs', '12'), ('total_nodes', '4'),
        ('status', 'PARTIAL — pouze nejvyšší soutěž (Státní liga). Nižší a regiony zatím nezpracovány.'),
        ('mistr', 'LTC Praha'),
        ('note', 'Pátý ročník. Dvě skupiny po 6, vítězové skupin (LTC, I. ČLTK) hráli finále — LTC mistr (7:1, 13:5). ŠK Banská Bystrica během sezóny vyloučena z ligy (kontumace). Sumy GF=GA i body ověřeny z PDF.'),
        ('chain_note', 'Navázáno dopředu na S1948_49 (prev_club_id doplněn tam u doložených klubů). Prev na S1946_47 zatím prázdný.'),
    ]:
        me.append([k, v])

    # NOTES
    no = wb.create_sheet('NOTES')
    no.append(['note'])
    for t in [
        'Nejvyšší soutěž (Státní liga) zpracována z PDF (dobový tisk).',
        'Všechny pražské oddíly (I. ČLTK, Podolí, LTC, Sparta) hrály domácí utkání na ZS Štvanice.',
        'Kvůli nepřízni počasí mnoho utkání přeloženo na umělé ledy v Brně, Bratislavě a Ostravě.',
        'ŠK Banská Bystrica vyloučena z ligy (nedostavila se k utkáním), výsledky zkontumovány 3:0.',
        'Nižší a krajské soutěže 1947/48 ZATÍM NEZPRACOVÁNY — přijdou později (chirurgicky).',
    ]:
        no.append([t])

    # šířky pro 10_liga
    for j, w in enumerate([6, 26, 5, 30, 5, 4, 4, 4, 4, 5, 3, 5, 5, 26, 20, 6, 20, 20, 18, 12, 12, 18, 10], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    wb.save(OUT)
    print(f"✔ {OUT}: 10_liga {len(rows10)-1} řádků, CLUBS {len(clubs)}, SERIES 2, SYSTEM 4")


# ── navázání 1948/49 → 1947/48 (doložené předchůdce) ──
LINKS_4849 = {                       # 1948/49 club_id : 1947/48 club_id
    'CLUB_S1948_49_0001': cid(7),    # LTC Praha
    'CLUB_S1948_49_0002': cid(8),    # ŠK Bratislava
    'CLUB_S1948_49_0004': cid(10),   # AC Stadion České Budějovice
    'CLUB_S1948_49_0006': cid(9),    # AC Sparta Bubeneč <- AC Sparta
    'CLUB_S1948_49_0007': cid(3),    # Sokol Prostějov <- SK Prostějov
    'CLUB_S1948_49_0008': cid(1),    # I. ČLTK Praha
    'CLUB_S1948_49_0005': cid(4),    # Sokol Zbrojovka Židenice <- (sloučení) Horácká Slavia Třebíč
}


def link_4849():
    path = 'data/S1948_49_FINAL.xlsx'
    wb = openpyxl.load_workbook(path)
    ws = wb['CLUBS']
    hdr = [c.value for c in ws[1]]
    ci, pi, ni = hdr.index('club_id'), hdr.index('prev_club_id'), hdr.index('change_note')
    n = 0
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, ci + 1).value
        if c in LINKS_4849 and not ws.cell(r, pi + 1).value:
            ws.cell(r, pi + 1).value = LINKS_4849[c]
            ws.cell(r, ni + 1).value = 'prev_club_id doplněn ze S1947_48 (nejvyšší soutěž z PDF).'
            n += 1
    wb.save(path)
    print(f"✔ S1948_49: doplněno {n} prev_club_id na S1947_48")


if __name__ == '__main__':
    build()
    link_4849()
