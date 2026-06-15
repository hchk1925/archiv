#!/usr/bin/env python3
"""
expand_names.py — „blbuvzdorné" plné názvy soutěží.

Rozbalí kódové tokeny v názvech uzlů (SYSTEM list) na čitelné názvy:
  L10/L100 → nejvyšší soutěž, L15 → kvalifikace, L20 → 2. úroveň (oblastní),
  L30 → 3. úroveň (krajský přebor), L40 → 4. úroveň, L50 → 5. úroveň,
  „Tyrš.kraj" → „Tyršův kraj", odstraní kódový prefix „30_" apod.

Transformace je IDEMPOTENTNÍ (po rozbalení už token není) a MINIMÁLNÍ
(čisté názvy jako „I. liga" nechá být). Funkce pretty() se používá i k zobrazení.

  python expand_names.py --preview   → statistika + docs/NAZVY_PREVIEW.csv (nezapisuje)
  python expand_names.py --apply     → zapíše nové názvy do data/S*_FINAL.xlsx
"""
import csv
import glob
import os
import re
import sys
import openpyxl

DATA = os.path.join(os.path.dirname(__file__), 'data')

# editovatelný slovník úrovní (lze upravit a skript spustit znovu)
LEVEL_NAMES = {
    'L100': 'nejvyšší soutěž',
    'L10': 'nejvyšší soutěž',
    'L15': 'kvalifikace',
    'L20': '2. úroveň (oblastní)',
    'L30': '3. úroveň (krajský přebor)',
}
# kódová slova → plné názvy
WORD_MAP = [
    (r'\bPHA\s+mesto\b', 'Praha město'),
    (r'\bPHA\s+venkov\b', 'Praha venkov'),
    (r'\bSVK\b', 'Slovensko'),
    (r'\bcelost\b', 'celostátní'),
]
_LTOK = re.compile(r'\bL(1[05]|[1-9]0|100)\b')
_PHASE = (r'(úroveň(?:\s*\([^)]*\))?)\s+'
          r'(skupina|Skupina|finále|Finále|semifinále|Semifinále|'
          r'čtvrtfinále|Čtvrtfinále|předkolo|baráž|play\-?off|'
          r'kvalifikace|základní|Základní|sk\.)')


def _level_label(tok):
    if tok in LEVEL_NAMES:
        return LEVEL_NAMES[tok]
    n = tok[1:-1] if tok != 'L15' else '15'      # Ln0 → n
    return f'{n}. úroveň'


def pretty(name):
    """Vrať čitelný plný název. Bez kódových artefaktů → vrací (skoro) beze změny."""
    s = str(name or '').strip()
    if not s:
        return s
    s = re.sub(r'\bTyrš\.\s*', 'Tyršův ', s)         # zkratka
    s = re.sub(r'^\s*[1-5]0_', '', s)                # kódový prefix „30_…"
    s = s.replace('_', ' ')                          # podtržítka → mezery
    for pat, repl in WORD_MAP:                       # kódová slova
        s = re.sub(pat, repl, s)
    if _LTOK.search(s):                              # L-token → ", N. úroveň"
        s = _LTOK.sub(lambda m: '\x00' + _level_label('L' + m.group(1)), s)
        s = re.sub(r'\s*\x00\s*', ', ', s)
        s = re.sub(_PHASE, r'\1, \2', s)             # čárka před fází
    s = re.sub(r',\s*,', ',', s)
    s = re.sub(r'\s+', ' ', s).strip().strip(',').strip()
    return s


def iter_system_rows(ws):
    head = [c.value for c in ws[1]]
    try:
        ni = head.index('name')
    except ValueError:
        return
    for row in ws.iter_rows(min_row=2):
        if ni < len(row):
            yield row[ni]


def run(apply):
    changes = {}       # (before -> after) -> [count, example_sid]
    files = sorted(glob.glob(os.path.join(DATA, 'S*_FINAL.xlsx')))
    touched_files = 0
    total_changes = 0
    for path in files:
        sid = re.search(r'S(\d{4}_\d{2})', os.path.basename(path)).group(1)
        wb = openpyxl.load_workbook(path)
        if 'SYSTEM' not in wb.sheetnames:
            wb.close(); continue
        ws = wb['SYSTEM']
        dirty = False
        for cell in iter_system_rows(ws):
            old = cell.value
            new = pretty(old)
            if new != (old or '').strip() and new != old:
                key = (str(old), new)
                if key not in changes:
                    changes[key] = [0, sid]
                changes[key][0] += 1
                total_changes += 1
                if apply:
                    cell.value = new
                    dirty = True
        if apply and dirty:
            wb.save(path); touched_files += 1
        wb.close()
    return changes, total_changes, touched_files, len(files)


def main():
    apply = '--apply' in sys.argv
    changes, total, touched, nfiles = run(apply)
    distinct = sorted(changes.items(), key=lambda kv: -kv[1][0])
    print(f"{'ZAPSÁNO' if apply else 'PREVIEW'}: {total} změn názvů "
          f"({len(distinct)} unikátních) napříč {nfiles} sezónami.")
    if apply:
        print(f"  upraveno souborů: {touched}")
    print("\nTop 15 nejčastějších změn (před → po):")
    for (before, after), (cnt, sid) in distinct[:15]:
        print(f"  {cnt:5}×  {before[:42]:42}  →  {after}")
    if not apply:
        os.makedirs('docs', exist_ok=True)
        with open('docs/NAZVY_PREVIEW.csv', 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f)
            w.writerow(['pocet', 'priklad_sezona', 'pred', 'po'])
            for (before, after), (cnt, sid) in distinct:
                w.writerow([cnt, sid, before, after])
        print("\nKompletní seznam změn: docs/NAZVY_PREVIEW.csv")


if __name__ == '__main__':
    main()
