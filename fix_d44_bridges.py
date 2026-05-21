#!/usr/bin/env python3
"""
fix_d44_bridges.py — Identity můstky přes organizační reformy (D44).

Pro klub v sezóně N+1 BEZ prev_club_id najde předchůdce v sezóně N
podle pravidla D44: identita = (město, jádro názvu) bez ohledu na
organizační prefix (Sokol/ZSJ/DSO/DŠO/DSJ/TJ/HC/VTJ/ASD/SK/...).

- Striktní: stejné city + stejné name-core
- B-tým ↔ B-tým, A-tým ↔ A-tým
- Pokud >1 kandidát ve stejném městě, vybere podle name+sheet shody
- Doplní TBD: flag pro pozdější verifikaci
- Synchronizuje do datových listů
"""
import openpyxl, glob, os, re, sys
from collections import defaultdict

DATA = sys.argv[1] if len(sys.argv) > 1 else 'data'
NON_DATA = {'NOTES', 'SYSTEM', 'SERIES', 'CLUBS', 'META', 'PYRAMIDA', 'TODO'}
ORG_PREFIX = re.compile(
    r'^(TJ|HC|VTJ|ASD|ZTJ|DSO|DŠO|DSJ|ZSJ|Sokol|SK|AC|AFK|SKP|KLH|KS|BK|HO|RH|ČH|VSJ|PDA|DA|VŠ|VŠS)\.?\s+',
    re.I)
B_SUF = re.compile(r'\s+(B|II|III|IV)\s*$')


def hidx(row0):
    return {h: i for i, h in enumerate(row0) if h is not None}


def core(name):
    """D44 jádro: bez org. prefixů, závorek, B/II/.. suffixu."""
    if not name:
        return ''
    s = str(name).strip()
    s = re.sub(r'\s*\([^)]*\)\s*', ' ', s)
    prev = None
    while prev != s:
        prev = s
        s = ORG_PREFIX.sub('', s).strip()
    s = B_SUF.sub('', s).strip()
    return re.sub(r'\s+', ' ', s).strip().lower()


def is_b(name):
    return bool(B_SUF.search(str(name or '')))


def season_of(p):
    m = re.search(r'S(\d{4})_(\d{2})', os.path.basename(p))
    return f"S{m.group(1)}_{m.group(2)}"


def load_clubs(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['CLUBS']
    rows = list(ws.iter_rows(values_only=True))
    H = hidx(rows[0])
    out = {}
    for r in rows[1:]:
        if not r or r[H['club_id']] is None:
            continue
        out[r[H['club_id']]] = {
            'clean': r[H['clean_name']],
            'sheet': r[H.get('sheet', -1)] if 'sheet' in H else None,
            'city': r[H['city']] if 'city' in H else None,
            'prev': r[H['prev_club_id']] if 'prev_club_id' in H else None,
        }
    wb.close()
    return out


def main():
    files = sorted(glob.glob(os.path.join(DATA, 'S*_FINAL.xlsx')))
    seasons = [season_of(p) for p in files]
    s_by = dict(zip(seasons, files))
    clubs_by = {s: load_clubs(p) for s, p in s_by.items()}

    bridged = ambig = 0
    by_season_fix = defaultdict(dict)            # season → {cid: prev_cid}
    for i in range(1, len(seasons)):
        sid = seasons[i]
        psid = seasons[i - 1]
        cur, prev = clubs_by[sid], clubs_by[psid]
        # předchůdce kandidáti: index podle (city, core, is_b)
        prev_idx = defaultdict(list)
        for pcid, pc in prev.items():
            key = (str(pc['city'] or '').strip().lower(),
                   core(pc['clean']), is_b(pc['clean']))
            if key[0] and key[1]:
                prev_idx[key].append(pcid)
        # pro každý klub bez prev v aktuální sezóně
        for cid, c in cur.items():
            if c['prev']:
                continue
            ck = (str(c['city'] or '').strip().lower(),
                  core(c['clean']), is_b(c['clean']))
            if not ck[0] or not ck[1]:
                continue
            cands = prev_idx.get(ck, [])
            if not cands:
                continue
            if len(cands) == 1:
                by_season_fix[sid][cid] = cands[0]
                bridged += 1
            else:
                # tie-breaker: stejný sheet
                same_sheet = [pc for pc in cands
                              if prev[pc]['sheet'] == c['sheet']]
                if len(same_sheet) == 1:
                    by_season_fix[sid][cid] = same_sheet[0]
                    bridged += 1
                else:
                    ambig += 1

    print(f"D44 bridges to apply: {bridged}  (ambiguózní {ambig} → ponecháno)")

    # APLIKACE
    for sid, fix_map in by_season_fix.items():
        path = s_by[sid]
        wb = openpyxl.load_workbook(path)
        ws = wb['CLUBS']
        rows = list(ws.iter_rows())
        H = hidx([c.value for c in rows[0]])
        ci, pi = H['club_id'], H['prev_club_id']
        cni = H.get('change_note')
        for row in rows[1:]:
            cid = row[ci].value
            if cid in fix_map:
                tgt = fix_map[cid]
                row[pi].value = tgt
                if cni is not None:
                    old = row[cni].value
                    tag = (f"TBD: D44 můstek prev_club_id → {tgt} "
                           f"(stejné city+jádro, org. reforma) [fix_d44_bridges]")
                    row[cni].value = (str(old) + ' || ' + tag) if old else tag
        # sync do datových listů
        for sh in wb.sheetnames:
            if sh in NON_DATA:
                continue
            dws = wb[sh]
            drows = list(dws.iter_rows())
            if not drows:
                continue
            DH = hidx([c.value for c in drows[0]])
            if 'club_id' not in DH or 'prev_club_id' not in DH:
                continue
            for row in drows[1:]:
                if row[DH['club_id']].value in fix_map:
                    row[DH['prev_club_id']].value = fix_map[row[DH['club_id']].value]
        wb.save(path)
        wb.close()
        print(f"  ✓ {sid}: +{len(fix_map)} prev linků")


if __name__ == '__main__':
    main()
