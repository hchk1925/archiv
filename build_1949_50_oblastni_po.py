#!/usr/bin/env python3
"""
build_1949_50_oblastni_po.py — doplní oblastní play-off 1949/50 z DS PDF (in-place).

V datech chyběla závěrečná fáze oblastní soutěže (čtvrtfinále → semifinálové
skupiny A/B → finálová skupina). PDF ji má kompletní; VÍTĚZ oblastní soutěže =
JTO Sokol I Pardubice. Play-off týmy = vítězové krajských skupin (reuse club_id).
Idempotentní.
"""
import openpyxl

PATH = 'data/S1949_50_FINAL.xlsx'
S = 'S1949_50'
# (clean_name v datech, club_id, GP,W,D,L,GF,GA,PTS, note)
SF_A = [('ZSJ Meteor Rovnost České Budějovice','CLUB_S1949_50_0030',3,2,0,1,21,10,4,''),
        ('Sokol Pardubice I','CLUB_S1949_50_0018',3,2,0,1,16,10,4,''),
        ('ZSJ ZMP Praha I','CLUB_S1949_50_0046',3,2,0,1,16,11,4,''),
        ('Sokol Jimlín','CLUB_S1949_50_0014',3,0,0,3,7,29,0,'')]
SF_B = [('ZSJ ČSSZ Prostějov','CLUB_S1949_50_0052',2,2,0,0,18,5,4,''),
        ('Sokol Tatry Poprad','CLUB_S1949_50_0064',2,1,0,1,11,15,2,''),
        ('ZSJ Spartak Zbrojovka Brno II - Židenice','CLUB_S1949_50_0058',2,0,0,2,7,16,0,'')]
FINAL = [('Sokol Pardubice I','CLUB_S1949_50_0018',3,3,0,0,21,8,6,'M-oblast'),
         ('ZSJ ČSSZ Prostějov','CLUB_S1949_50_0052',3,2,0,1,18,10,4,''),
         ('ZSJ Meteor Rovnost České Budějovice','CLUB_S1949_50_0030',3,0,1,2,8,15,1,''),
         ('Sokol Tatry Poprad','CLUB_S1949_50_0064',3,0,1,2,4,18,1,'')]
# čtvrtfinále → SERIES: (a_name,a_cid,b_name,b_cid, score, winner_note)
QF = [('Slavia Karlovy Vary','CLUB_S1949_50_0009','Sokol Jimlín','CLUB_S1949_50_0014','2:5','postup Jimlín'),
      ('Sokol Pardubice I','CLUB_S1949_50_0018','Sokol Železný Brod','CLUB_S1949_50_0024','6:1','postup Pardubice'),
      ('ZSJ Meteor Rovnost České Budějovice','CLUB_S1949_50_0030','Felbabka','CLUB_S1949_50_0036','5:4','postup Meteor ČB'),
      ('ZSJ Kara Starý Kolín','CLUB_S1949_50_0041','ZSJ ZMP Praha I','CLUB_S1949_50_0046','3:8','postup ZMP Praha')]
NOTES = [
    'Oblastní soutěž 1949/50 — závěr: čtvrtfinále → semifinálové skupiny A/B → finálová skupina. VÍTĚZ OBLASTNÍ SOUTĚŽE: JTO Sokol I Pardubice (postup do nejvyšší soutěže 1950/51).',
    'Čtvrtfinále: Jimlín, Pardubice, Meteor ČB, ZMP Praha postoupili; Prostějov, Zbrojovka Brno-Židenice II. a Tatry Poprad postoupili přímo do semifinále.',
    'Krajské I. třídy: vítězové finálových skupin (Strašnice, Důl Anna Rynholec, II. Smíchov, SONP Kladno II. ad.) postoupili do oblastního mistrovství 1950/51 — viz PDF.',
    'Krajské soutěže (listy 30*) pocházejí z původního almanachu; DS PDF 1949/50 je nepokrývá (jen oblastní úroveň), proto nebyly proti tomuto PDF křížově ověřeny. Část tabulek III. tříd má neúplná GF/GA ve zdroji.',
]


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb['20_oblastni']
    H = {c.value: i for i, c in enumerate(ws[1])}
    # už hotovo?
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == 'H' and 'Finálová' in str(r[1] or ''):
            print("  · play-off už existuje");
            return
    import re
    sy = wb['SYSTEM']
    nnums = [int(m.group(1)) for r in range(2, sy.max_row + 1)
             if (v := sy.cell(r, 1).value) and (m := re.search(r'_(\d+)$', str(v)))]
    nidx = max(nnums) + 1
    # parent = oblastní uzel
    parent = ''
    for r in range(2, sy.max_row + 1):
        if 'Oblastní' in str(sy.cell(r, 2).value or '') and not str(sy.cell(r, 2).value).count('/'):
            parent = sy.cell(r, 1).value; break

    def node(name, ctype):
        nonlocal nidx
        n = f'NODE_{S}_{nidx:04d}'; nidx += 1
        sy.append([n, name, ctype, 'L20', 'REG_CS', parent, '', '', '2-1-0', '', '', '', '', ''])
        return n
    npo = node('Oblastní soutěž / play-off', 'playoff')
    nsfa = node('Oblastní soutěž / Semifinálová skupina A', 'final_group')
    nsfb = node('Oblastní soutěž / Semifinálová skupina B', 'final_group')
    nfin = node('Oblastní soutěž / Finálová skupina', 'final_group')

    tr = [900]
    cRT, cBN = H['row_type'], H['block_name']

    def block(bname, nd, teams):
        ws.append(['H', bname] + ['']*(cBN-1) + ['']*(len(ws[1])-cBN-1))  # placeholder
        # přepiš celý řádek korektně
        row = ws[ws.max_row]
        for c in row: c.value = None
        ws.cell(ws.max_row, cRT+1).value = 'H'
        ws.cell(ws.max_row, cBN+1).value = bname
        ws.cell(ws.max_row, H['node_id']+1).value = nd
        ws.cell(ws.max_row, H['level']+1).value = 'L20'
        for i, t in enumerate(teams, 1):
            nm, cid, gp, w, d, l, gf, ga, pts, nt = t
            tr[0] += 1
            vals = {'row_type':'T','pos':i,'club_name':nm,'note':nt,'GP':gp,'W':w,'D':d,
                    'L':l,'GF':gf,':':':','GA':ga,'PTS':pts,'comp_path':bname,'node_id':nd,
                    'level':'L20','club_id':cid,'tr_id':f'TR_{S}_{tr[0]:05d}'}
            ws.append([''] * len(ws[1]))
            for k, v in vals.items():
                ws.cell(ws.max_row, H[k]+1).value = v

    block('Oblastní soutěž / Semifinálová skupina A', nsfa, SF_A)
    block('Oblastní soutěž / Semifinálová skupina B', nsfb, SF_B)
    block('Oblastní soutěž / Finálová skupina', nfin, FINAL)

    # čtvrtfinále → SERIES
    se = wb['SERIES']
    for k, (an, ac, bn, bc, sc, note) in enumerate(QF, 1):
        sid = f'SER_{S}_09{k:02d}'
        se.append([sid, npo, ac, an, an, '1', sc, '', '', f'Oblastní ČF: {note}'])
        se.append([sid, npo, bc, bn, bn, '2', sc, '', '', ''])

    # NOTES + META
    no = wb['NOTES']
    have = {no.cell(r, 1).value for r in range(2, no.max_row + 1)}
    for t in NOTES:
        if t not in have:
            no.append([t])
    me = wb['META']
    mk = {me.cell(r, 1).value: r for r in range(2, me.max_row + 1)}
    if 'mistr_oblast' in mk: me.cell(mk['mistr_oblast'], 2).value = 'JTO Sokol I Pardubice'
    else: me.append(['mistr_oblast', 'JTO Sokol I Pardubice'])

    wb.save(PATH)
    print("✔ oblastní play-off (ČF + semifinálové A/B + finálová) — vítěz Pardubice")


if __name__ == '__main__':
    main()
