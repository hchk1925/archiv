#!/usr/bin/env python3
"""
krajske_check.py — worklist nekonzistencí krajských soutěží (listy 30*) napříč
všemi sezónami → docs/KRAJSKE_CHECK.md.

Krajské nemají externí zdroj (DS PDF je nepokrývá), takže standings se NEopravují
automaticky — tento report slouží pro ruční průchod. Hlásí jen REÁLNÉ chyby:
  - GF≠GA při kompletních datech (chyba skóre),
  - PTS≠2·V+R při V+R+P=GP (chyba bodů),
  - V+R+P≠GP (chyba zápasů).
„Neúplná GF/GA" (řídká zdrojová data) se NEhlásí — to není chyba.
"""
import openpyxl, glob, re, datetime


def num(x):
    return x if isinstance(x, int) else None


def main():
    out = ["# Krajské soutěže — worklist nekonzistencí (k ručnímu průchodu)\n",
           f"_{datetime.datetime.now():%Y-%m-%d %H:%M} · generuje `krajske_check.py`_\n",
           "Krajské nemají externí zdroj (PDF je nepokrývá) → opravit ručně dle "
           "vlastních podkladů. Hlásí jen reálné chyby (ne řídká neúplná data).\n"]
    tot = 0
    for p in sorted(glob.glob('data/S*_FINAL.xlsx')):
        sid = re.search(r'S(\d{4}_\d{2})', p).group(1)
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        season_issues = []
        for sh in wb.sheetnames:
            if not sh.startswith('30'):
                continue
            ws = wb[sh]
            blk = ''
            rows = []

            def flush():
                if not rows:
                    return
                gfs = [r[4] for r in rows if isinstance(r[4], int)]
                gas = [r[5] for r in rows if isinstance(r[5], int)]
                fl = []
                if len(gfs) == len(rows) and len(gas) == len(rows) and sum(gfs) != sum(gas):
                    fl.append(f"GF{sum(gfs)}≠GA{sum(gas)}")
                wbad = sum(1 for r in rows if None not in r[:4] and r[1]+r[2]+r[3] != r[0])
                if wbad:
                    fl.append(f"V+R+P≠GP×{wbad}")
                pbad = sum(1 for r in rows if None not in (r[6], r[1], r[2])
                           and None not in r[:4] and r[1]+r[2]+r[3] == r[0]
                           and r[6] != 2*r[1]+r[2])
                if pbad:
                    fl.append(f"PTS≠2V+R×{pbad}")
                if fl:
                    season_issues.append(f"  - `{sh}` [{blk[:34]}]: {', '.join(fl)}")

            for r in ws.iter_rows(values_only=True):
                if r and r[0] == 'H':
                    flush(); blk = str(r[1] or ''); rows = []
                elif r and r[0] == 'T':
                    rows.append((num(r[5]), num(r[6]), num(r[7]), num(r[8]),
                                 num(r[9]), num(r[11]), num(r[12])))
            flush()
        wb.close()
        if season_issues:
            out.append(f"\n## {sid}  ({len(season_issues)} skupin)\n")
            out += season_issues
            tot += len(season_issues)
    out.insert(3, f"\n**Celkem skupin s nálezem: {tot}** (napříč všemi sezónami)\n")
    with open('docs/KRAJSKE_CHECK.md', 'w', encoding='utf-8') as f:
        f.write('\n'.join(out) + '\n')
    print(f"docs/KRAJSKE_CHECK.md — {tot} skupin s nálezem")


if __name__ == '__main__':
    main()
