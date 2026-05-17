#!/usr/bin/env python3
"""Audit continuity across all season workbooks.

Reports:
  - prev_club_id coverage & broken links (CLUBS sheet, season N -> N-1)
  - clubs in season N+1 with no prev_club_id but a name-match in season N (missing links)
  - season_fate coverage & per-club consistency (data sheets)
  - SYSTEM feeds_into / feeds_into_loser coverage
"""
import openpyxl, glob, os, re, sys
from collections import defaultdict

FILES = sorted(glob.glob('data/S*_FINAL.xlsx'))

DATA_SHEET_RE = re.compile(r'^(10_|20_|30|40|KVAL|50)')
NON_DATA = {'NOTES', 'SYSTEM', 'SERIES', 'CLUBS', 'META', 'PYRAMIDA'}


def season_key(path):
    m = re.search(r'S(\d{4})_(\d{2})', path)
    return int(m.group(1))


def load_clubs(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['CLUBS']
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    idx = {h: i for i, h in enumerate(hdr)}
    clubs = {}
    for r in rows[1:]:
        if not r or r[idx['club_id']] is None:
            continue
        cid = r[idx['club_id']]
        clubs[cid] = {
            'clean_name': r[idx['clean_name']],
            'sheet': r[idx['sheet']],
            'level': r[idx.get('level', 5)] if 'level' in idx else None,
            'prev_club_id': r[idx['prev_club_id']],
            'city': r[idx['city']] if 'city' in idx else None,
        }
    wb.close()
    return clubs


def load_fate(path):
    """Return {club_id: set(fate)} and counts from data sheets."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    cid_fate = defaultdict(set)
    total_t = 0
    fate_set = 0
    for sh in wb.sheetnames:
        if sh in NON_DATA:
            continue
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hdr = list(rows[0])
        idx = {h: i for i, h in enumerate(hdr)}
        if 'club_id' not in idx or 'season_fate' not in idx:
            continue
        rt_i = idx.get('row_type')
        for r in rows[1:]:
            if not r:
                continue
            rt = r[rt_i] if rt_i is not None else None
            cid = r[idx['club_id']]
            fate = r[idx['season_fate']]
            if rt == 'T' and cid:
                total_t += 1
                if fate:
                    fate_set += 1
            if cid and fate:
                cid_fate[cid].add(str(fate).strip())
    wb.close()
    return cid_fate, total_t, fate_set


def load_system(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['SYSTEM']
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    idx = {h: i for i, h in enumerate(hdr)}
    total = 0
    feeds = 0
    for r in rows[1:]:
        if not r or r[idx['node_id']] is None:
            continue
        total += 1
        if r[idx.get('feeds_into')] or r[idx.get('feeds_into_loser')]:
            feeds += 1
    wb.close()
    return total, feeds


print(f"{'Season':>9} | {'#clubs':>6} | {'prevSet':>7} | {'prevBrk':>7} | "
      f"{'miss':>5} | {'fateT':>6}/{'tot':<6} | {'fInconst':>8} | {'sysFeeds':>8}/{'sys':<5}")
print("-" * 100)

prev_clubs = None
prev_path = None
grand = defaultdict(int)

for path in FILES:
    base = os.path.basename(path)
    clubs = load_clubs(path)
    cid_fate, tot_t, fate_t = load_fate(path)
    sys_tot, sys_feeds = load_system(path)

    n = len(clubs)
    prev_set = sum(1 for c in clubs.values() if c['prev_club_id'])
    prev_brk = 0
    miss = 0
    if prev_clubs is not None:
        prev_ids = set(prev_clubs.keys())
        prev_by_name = defaultdict(list)
        for pid, pc in prev_clubs.items():
            prev_by_name[(str(pc['clean_name']).strip().lower())].append(pid)
        for cid, c in clubs.items():
            pcid = c['prev_club_id']
            if pcid:
                if pcid not in prev_ids:
                    prev_brk += 1
            else:
                nm = str(c['clean_name']).strip().lower()
                if nm in prev_by_name:
                    miss += 1

    f_inconst = sum(1 for cid, fs in cid_fate.items() if len(fs) > 1)

    grand['clubs'] += n
    grand['prev_set'] += prev_set
    grand['prev_brk'] += prev_brk
    grand['miss'] += miss
    grand['fate_t'] += fate_t
    grand['tot_t'] += tot_t
    grand['f_inconst'] += f_inconst
    grand['sys_feeds'] += sys_feeds
    grand['sys_tot'] += sys_tot

    print(f"{base[1:8]:>9} | {n:>6} | {prev_set:>7} | {prev_brk:>7} | "
          f"{miss:>5} | {fate_t:>6}/{tot_t:<6} | {f_inconst:>8} | {sys_feeds:>8}/{sys_tot:<5}")

    prev_clubs = clubs
    prev_path = path

print("-" * 100)
print(f"TOTALS: clubs={grand['clubs']} prev_set={grand['prev_set']} "
      f"prev_broken={grand['prev_brk']} missing_links={grand['miss']} "
      f"fate={grand['fate_t']}/{grand['tot_t']} fate_inconsistent={grand['f_inconst']} "
      f"sys_feeds={grand['sys_feeds']}/{grand['sys_tot']}")
