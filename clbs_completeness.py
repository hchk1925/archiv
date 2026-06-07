#!/usr/bin/env python3
"""
clbs_completeness.py — doplní CHYBĚJÍCÍ kluby nižších (okresních) soutěží do
per-sezónního xlsx z kolegova indexu 'clbs' (jen názvy + pořadí + příznak N/M/S,
BEZ statistik). Záměr: základ úplnosti k pozdějšímu ručnímu doladění.

Použití:
    python3 clbs_completeness.py <sezona.xlsx> <kolega.xlsx> [--write]

Bere jen kraje (ne I.liga/kvalifikace o ligu). Týmy, které už v sezoně máme
(dle jména), nepřidává. Nové okresní soutěže zakládá jako uzly L30.
"""
import sys, re, shutil
import openpyxl

# (klíčová slova v názvu listu) -> kandidátní kódy našich kraj-sheetů.
# Pokrývá starou éru krajů 1949-1960 (PHAM/JHCK/...) i reformu 1960+
# (PRAH/STRC/...). Z kandidátů se za běhu vybere ten, jehož sheet v sezoně
# existuje, takže se mapování samo přizpůsobí éře.
KRAJ_KW = [
    (('venkov',), ('PHAV', 'STRC')),
    (('středočesk', 'středolab', 'kladn', 'kladen'), ('STRC', 'PHAV')),
    (('tyrš', 'praha', 'pražsk'), ('PHAM', 'PRAH')),
    (('jihočesk',), ('JHCK',)),
    (('plzeň', 'západočesk', 'šumav'), ('PLZN', 'ZAPC')),
    (('karlovar',), ('KVRY', 'ZAPC')),
    (('ústeck', 'severočesk', 'severozápadočesk'), ('USTE', 'SVRC')),
    (('libereck',), ('LIBE', 'SVRC')),
    (('pardubick', 'východočesk', 'středolab'), ('PARD', 'VYCH')),
    (('hradeck', 'královéhradeck'), ('HRAD', 'VYCH')),
    (('vysočin',), ('VYSO', 'JHMR', 'VYCH', 'JIHL')),
    (('jihlavsk',), ('JIHL', 'VYSO')),
    (('brněnsk', 'jihomorav'), ('BRNO', 'JHMR')),
    (('gottwaldov',), ('GOTT',)),
    (('olomouck',), ('OLOM',)),
    (('ostravsk', 'slezsk', 'severomorav'), ('OSTR', 'SVMR')),
]
SKIP_LIST = {'I.liga', 'kvalifikace o ligu', 'SVK'}

def kraj_sheet(lst, sheetnames):
    """Vrať (sheet_name, okr_level) pro daný list, nebo (None, None)."""
    s = str(lst or '').lower()
    krsheets = [sn for sn in sheetnames if sn[:2].isdigit()]
    for kws, codes in KRAJ_KW:
        if any(k in s for k in kws):
            for code in codes:
                sn = next((x for x in krsheets if x[2:] == code), None)
                if sn:
                    return sn, f"L{int(sn[:2]) + 10}"
    return None, None

STD_HEADER = ['row_type','block_name','pos','club_name','note','GP','W','D','L',
              'GF',':','GA','PTS','comp_path','node_id','level','club_id',
              'prev_club_id','dest_node_id','dest_type','season_fate','tr_id','district']


def norm(s):
    s = re.sub(r'\s+', ' ', str(s or '')).strip()
    s = re.sub(r'\s*[\(\[](N|S|M)[\)\]]\s*$', '', s)
    return s.lower()


def main():
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    season_f, coll_f = sys.argv[1], sys.argv[2]
    do_write = '--write' in sys.argv

    cwb = openpyxl.load_workbook(coll_f, read_only=True, data_only=True)
    if 'clbs' not in cwb.sheetnames:
        print("!! kolega nemá list 'clbs'"); sys.exit(1)
    crows = list(cwb['clbs'].iter_rows(values_only=True))
    ch = {str(v): i for i, v in enumerate(crows[0])}
    def cc(r, name):
        return r[ch[name]] if name in ch else None
    cwb.close()

    wb = openpyxl.load_workbook(season_f)
    season_id = re.search(r'S(\d{4}_\d{2})', season_f).group(0)

    # všechny existující názvy klubů v sezoně (dedup)
    have = set()
    for sn in wb.sheetnames:
        rows = list(wb[sn].iter_rows(values_only=True))
        if not rows:
            continue
        h = {str(v): i for i, v in enumerate(rows[0])}
        if 'club_name' not in h:
            continue
        for r in rows[1:]:
            if r and r[h.get('row_type', 0)] == 'T' and r[h['club_name']]:
                have.add(norm(r[h['club_name']]))

    def maxid(prefix):
        mx = 0
        for sn in wb.sheetnames:
            for row in wb[sn].iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        for g in re.findall(prefix + r'(\d+)', v):
                            mx = max(mx, int(g))
        return mx
    club_n = maxid(f'CLUB_{season_id}_'); tr_n = maxid(f'TR_{season_id}_')
    node_n = maxid(f'NODE_{season_id}_')

    sys_ws = wb['SYSTEM']
    SH = {str(v): i for i, v in enumerate(next(sys_ws.iter_rows(values_only=True)))}
    # region/kontejner per kraj sheet
    def kraj_region_parent(sheet_code):
        # najdi krajský kontejner v SYSTEM přes region z existujícího bloku listu
        rows = list(wb[sheet_code].iter_rows(values_only=True))
        h = {str(v): i for i, v in enumerate(rows[0])}
        nid = None
        for r in rows[1:]:
            if r[h['row_type']] == 'H':
                nid = r[h['node_id']]; break
        reg = None; parent = None
        for sr in sys_ws.iter_rows(values_only=True):
            if sr[SH['node_id']] == nid:
                reg = sr[SH['region']]
                parent = sr[SH['parent_node_id']] or nid
                break
        return reg, parent

    # seskup clbs řádky podle (List, Soutěž, Skupina) — jen kraje
    from collections import OrderedDict
    groups = OrderedDict()
    for r in crows[1:]:
        lst = str(cc(r, 'List (oblast)') or '').strip()
        if lst in SKIP_LIST or kraj_sheet(lst, wb.sheetnames)[0] is None:
            continue
        sou = str(cc(r, 'Soutěž') or '').strip()
        # přeskoč soutěže, které už v sezoně modelujeme (KP, kvalifikace) —
        # pár nenamatchovaných jmen jsou varianty, nechceme duplicitní uzly
        sl = sou.lower()
        if 'krajský přebor' in sl or 'kvalifik' in sl or 'postup' in sl:
            continue
        sk = cc(r, 'Skupina/část')
        sk = str(sk).strip() if sk is not None else ''
        groups.setdefault((lst, sou, sk), []).append(r)

    new_clubs = []; new_sys = []
    sheet_adds = {}  # sheet_code -> list of out-rows to append
    report = []

    def mint(name, level, parent, region):
        nonlocal node_n
        node_n += 1; nid = f"NODE_{season_id}_{node_n:04d}"
        new_sys.append((nid, name, 'league', level, region, parent))
        return nid

    for (lst, sou, sk), rws in groups.items():
        sheet, okr = kraj_sheet(lst, wb.sheetnames)
        if not sheet:
            continue
        # jen kluby, které ještě nemáme
        todo = [r for r in rws if norm(cc(r, 'Klub')) not in have]
        if not todo:
            continue
        reg, parent = kraj_region_parent(sheet)
        label = f"{sou} / {sk}" if sk else sou
        node = mint(f"{lst} – {label}", okr, parent, reg)
        out = [['H', label, None, None, None, None, None, None, None, None, None,
                None, None, None, node, okr, None, None, None, None, None, None, None]]
        for r in todo:
            nm = str(cc(r, 'Klub')).strip()
            pori = cc(r, 'Pořadí')
            flag = cc(r, 'Příznak []') or cc(r, 'Anotace ()')
            note = flag if str(flag) in ('N', 'S', 'M') else None
            club_n += 1; cid = f"CLUB_{season_id}_{club_n:04d}"
            new_clubs.append((cid, nm, note, sheet, okr))
            tr_n += 1
            out.append(['T', None, pori, nm, note, None, None, None, None, None,
                        ':', None, None, label, node, okr, cid, None, None, None,
                        None, f"TR_{season_id}_{tr_n:05d}", None])
            have.add(norm(nm))
        sheet_adds.setdefault(sheet, []).extend(out)
        report.append((sheet, label, len(todo)))

    # report
    bys = {}
    for sheet, label, n in report:
        bys.setdefault(sheet, []).append((label, n))
    total = sum(n for _, _, n in report)
    print(f"== {season_id}: doplnění okresních klubů z clbs (jen názvy) ==")
    for sheet in sorted(bys):
        print(f"  {sheet}: +{sum(n for _,n in bys[sheet])} klubů")
        for label, n in bys[sheet]:
            print(f"      {label}: {n}")
    print(f"  CELKEM +{total} klubů, +{len(new_sys)} uzlů")

    if not do_write:
        print("\n(dry-run — přidej --write)"); return

    shutil.copy(season_f, season_f + '.bak')
    for sheet, out in sheet_adds.items():
        ws = wb[sheet]
        for row in out:
            ws.append(row)
    for nid, name, ctype, level, region, parent in new_sys:
        sys_ws.append([nid, name, ctype, level, region, parent, None, None, None,
                       None, 'Okresní soutěž – základ úplnosti z clbs (jen názvy).',
                       None, None, None])
    # CLUBS – pořadí sloupců se mezi érami liší (stará vs moderní hlavička),
    # proto mapujeme podle NÁZVŮ sloupců, ne pozičně.
    cws = wb['CLUBS']
    chdr = [str(c.value) for c in next(cws.iter_rows())]
    ci = {name: i for i, name in enumerate(chdr)}
    desc = 'Okresní soutěž – základ úplnosti z clbs (jen názvy, bez statistik).'
    for cid, name, note, sheet, level in new_clubs:
        clean = re.sub(r'\s*[\(\[](N|S|M)[\)\]]\s*$', '', name).strip()
        vals = {'club_id': cid, 'raw_name': name, 'clean_name': clean,
                'sheet': sheet, 'entry_note': note, 'level': level,
                'district': None, 'prev_club_id': None, 'change_note': desc,
                'city': None}
        cws.append([vals.get(col) for col in chdr])
    wb.save(season_f)
    print(f"\n✓ {season_f}: +{total} klubů, +{len(new_sys)} uzlů (záloha {season_f}.bak)")


if __name__ == '__main__':
    main()
