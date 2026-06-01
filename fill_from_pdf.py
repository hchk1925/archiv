#!/usr/bin/env python3
"""
fill_from_pdf.py — doplnění/oprava poškozených extraligových tabulek z
originálních PDF (zdroj CZE1: CZ-YYYY-YY.pdf — „jednička v úplnosti a přesnosti").

Řeší sezóny, kde xlsx měl chybějící/posunuté hodnoty (viz docs/DATA_QUALITY.md):
  1980/81, 1981/82 — celý řádek posunutý (GA chyběla, PTS držel GA);
                     formát 4 sloupce V/VP/PP/P (bez remíz).
  2000/01, 2001/02 — chyběly GA a PTS; formát 5 sloupců V/VP/R/PP/P.

Mapování do 3-sloupcové schématu W/D/L:
  W = výhry celkem (řádné + v prodloužení), D = remízy, L = prohry celkem.
  GF, GA, PTS se berou přímo z PDF (oficiální).

Pojistky proti špatnému zarovnání (skript NIC nezapíše, pokud neprojdou):
  - počet T-řádků extraligy == počet řádků z PDF a každý GP sedí;
  - 1980/81,1981/82: současná „PTS" buňka v xlsx == GA z PDF (drží posunutou GA);
  - 2000/01,2001/02: GF v xlsx == GF z PDF;
  - interní kontrola PDF: suma GF == suma GA, V+R+P == GP.

Spouštěj:  python3 fill_from_pdf.py [PDF_DIR] [--write]
"""
import openpyxl, glob, os, re, sys
from pdfminer.high_level import extract_text

PDF_DIR = next((a for a in sys.argv[1:] if not a.startswith('--')), 'sources/CZE1')
WRITE = '--write' in sys.argv

ROW = re.compile(r'(?<!\d)(\d{2})\s+((?:\d{1,3}\s+){3,5})(\d{1,3}):(\d{1,3})\s+(\d{1,3})(?!\d)')
# flexibilní: GP, 2-5 výsledků, GF:GA (případné PTS ignorujeme — dopočítáme)
ROW_FLEX = re.compile(r'(?<!\d)(\d{1,2})\s+((?:\d{1,3}\s+){2,5}?)(\d{1,3}):(\d{1,3})(?!\d)')

# season_xlsx: (pdf_basename, gp, n_teams, mode)
SEASONS = {
    'S1980_81': ('CZ-1980-81', 44, 12, 'shift'),
    'S1981_82': ('CZ-1981-82', 44, 12, 'shift'),
    'S2000_01': ('CZ-2000-01', 52, 14, 'ga_pts'),
    'S2001_02': ('CZ-2001-02', 52, 14, 'ga_pts'),
    # rozdělené ligy/nadstavba: xlsx „GA" sloupec drží skutečné GF → klíč párování
    'S1991_92': ('CZ-1991-92', None, None, 'regroup'),
    'S1992_93': ('CZ-1992-93', None, None, 'regroup'),
    'S1993_94': ('CZ-1993-94', None, None, 'regroup'),
    'S1994_95': ('CZ-1994-95', None, None, 'regroup'),
    'S1995_96': ('CZ-1995-96', None, None, 'regroup'),
}


def parse_pdf(basename, gp, n):
    txt = extract_text(os.path.join(PDF_DIR, basename + '.pdf'))
    rows = []
    for m in ROW.finditer(txt):
        if int(m.group(1)) != gp:
            continue
        outs = [int(x) for x in m.group(2).split()]
        rows.append({'outs': outs, 'gf': int(m.group(3)),
                     'ga': int(m.group(4)), 'pts': int(m.group(5))})
        if len(rows) == n:
            break
    return rows


def parse_pdf_all(basename):
    """Všechny tabulkové řádky (libovolné GP) v pořadí dokumentu, PTS dopočteme."""
    txt = extract_text(os.path.join(PDF_DIR, basename + '.pdf'))
    rows = []
    for m in ROW_FLEX.finditer(txt):
        outs = [int(x) for x in m.group(2).split()]
        rows.append({'gp': int(m.group(1)), 'outs': outs,
                     'gf': int(m.group(3)), 'ga': int(m.group(4))})
    return rows


def derive_wdl(outs):
    if len(outs) == 3:                  # V R P (bez prodloužení — mini-skupiny)
        return outs[0], outs[1], outs[2]
    if len(outs) == 4:                  # V VP PP P (bez remíz)
        return outs[0] + outs[1], 0, outs[2] + outs[3]
    if len(outs) == 5:                  # V VP R PP P
        return outs[0] + outs[1], outs[2], outs[3] + outs[4]
    raise ValueError(f"neočekávaný počet sloupců: {outs}")


def hcol(ws, name):
    for j, c in enumerate(next(ws.iter_rows(max_row=1)), 1):
        if c.value == name:
            return j
    raise KeyError(name)


def process_regroup(sid):
    """Sezóny s rozdělenými skupinami: xlsx 'GA' sloupec drží skutečné GF.
    Páruj každý T-řádek na PDF přes (GP, GF) v pořadí dokumentu."""
    basename, _gp, _n, _mode = SEASONS[sid]
    pdf = parse_pdf_all(basename)
    for r in pdf:
        r['used'] = False
    path = f'data/{sid}_FINAL.xlsx'
    wb = openpyxl.load_workbook(path)
    ws = wb['10_liga']
    cRT, cGP = hcol(ws, 'row_type'), hcol(ws, 'GP')
    cW, cD, cL = hcol(ws, 'W'), hcol(ws, 'D'), hcol(ws, 'L')
    cGF, cGA, cPTS = hcol(ws, 'GF'), hcol(ws, 'GA'), hcol(ws, 'PTS')
    def take(gp, gf, ga=None):
        return next((r for r in pdf if not r['used'] and r['gp'] == gp
                     and r['gf'] == gf and (ga is None or r['ga'] == ga)), None)

    changes, skips = [], []
    for row in ws.iter_rows():
        if row[cRT - 1].value != 'T':
            continue
        gp = row[cGP - 1].value
        gf_col, ga_col = row[cGF - 1].value, row[cGA - 1].value
        if gp is None:
            continue
        # 1) GF/GA už správné (sedí na PDF)? Pak ještě srovnej W/D/L/PTS,
        #    protože i správný řádek mohl mít vypadlý sloupec výher v prodl.
        if isinstance(gf_col, int) and isinstance(ga_col, int):
            ok = take(gp, gf_col, ga_col)
            if ok is not None:
                ok['used'] = True
                w, d, l = derive_wdl(ok['outs'])
                if w + d + l == gp:
                    pts = 2 * w + d
                    old = (row[cW-1].value, row[cD-1].value, row[cL-1].value,
                           row[cPTS-1].value)
                    if (w, d, l, pts) != old:
                        changes.append((row[cRT-1].row,
                                        old + (gf_col, ga_col),
                                        (w, d, l, pts, gf_col, ga_col)))
                        if WRITE:
                            (row[cW-1].value, row[cD-1].value, row[cL-1].value,
                             row[cPTS-1].value) = w, d, l, pts
                continue
        # 2) posunutá varianta: skutečné GF leží v GA sloupci
        if not isinstance(ga_col, int):
            continue
        cand = take(gp, ga_col)
        if cand is None:
            skips.append((row[cRT-1].row, gp, ga_col, row[3].value))
            continue
        cand['used'] = True
        w, d, l = derive_wdl(cand['outs'])
        if w + d + l != gp:
            skips.append((row[cRT-1].row, gp, ga_col, 'V+R+P≠GP'))
            continue
        pts = 2 * w + d                       # 2-1-0 éra (VP=2, R=1)
        old = (row[cW-1].value, row[cD-1].value, row[cL-1].value,
               gf_col, ga_col, row[cPTS-1].value)
        new = (w, d, l, cand['gf'], cand['ga'], pts)
        if old != new:
            changes.append((row[cRT-1].row, old, new))
            if WRITE:
                (row[cW-1].value, row[cD-1].value, row[cL-1].value,
                 row[cGF-1].value, row[cGA-1].value, row[cPTS-1].value) = new
    if WRITE and changes:
        wb.save(path)
    wb.close()
    for rr, gp, gf, nm in skips:
        print(f"   ⚠ nepárováno ř.{rr}: GP={gp} GF={gf} '{nm}'")
    return changes


def process(sid):
    basename, gp, n, mode = SEASONS[sid]
    if mode == 'regroup':
        return process_regroup(sid)
    pdf = parse_pdf(basename, gp, n)
    assert len(pdf) == n, f"{sid}: PDF dalo {len(pdf)} řádků, čekáno {n}"
    sgf, sga = sum(r['gf'] for r in pdf), sum(r['ga'] for r in pdf)
    assert sgf == sga, f"{sid}: suma GF({sgf})≠GA({sga}) — PDF parse podezřelý"
    for r in pdf:
        w, d, l = derive_wdl(r['outs'])
        assert w + d + l == gp, f"{sid}: V+R+P≠GP u {r}"
        r['w'], r['d'], r['l'] = w, d, l

    path = f'data/{sid}_FINAL.xlsx'
    wb = openpyxl.load_workbook(path)
    ws = wb['10_liga']
    cW, cD, cL = hcol(ws, 'W'), hcol(ws, 'D'), hcol(ws, 'L')
    cGF, cGA, cPTS = hcol(ws, 'GF'), hcol(ws, 'GA'), hcol(ws, 'PTS')
    cRT = hcol(ws, 'row_type')
    trows = [row for row in ws.iter_rows() if row[cRT - 1].value == 'T']
    assert len(trows) == n, f"{sid}: xlsx má {len(trows)} T-řádků, čekáno {n}"

    # pojistky zarovnání
    for row, pr in zip(trows, pdf):
        if mode == 'ga_pts':
            xgf = row[cGF - 1].value
            assert xgf == pr['gf'], (f"{sid}: GF nesedí (xlsx {xgf} vs PDF "
                                     f"{pr['gf']}) — špatné zarovnání, NEPÍŠU")
        else:  # shift: před opravou drží GA sloupec 'PTS', po opravě sloupec 'GA'
            assert pr['ga'] in (row[cGA - 1].value, row[cPTS - 1].value), (
                f"{sid}: kontrolní GA {pr['ga']} není v GA/PTS buňce "
                f"({row[cGA-1].value}/{row[cPTS-1].value}) — NEPÍŠU")

    changes = []
    for row, pr in zip(trows, pdf):
        old = (row[cW-1].value, row[cD-1].value, row[cL-1].value,
               row[cGF-1].value, row[cGA-1].value, row[cPTS-1].value)
        new = (pr['w'], pr['d'], pr['l'], pr['gf'], pr['ga'], pr['pts'])
        if old != new:
            changes.append((row[cRT-1].row, old, new))
            if WRITE:
                row[cW-1].value, row[cD-1].value, row[cL-1].value = pr['w'], pr['d'], pr['l']
                row[cGF-1].value, row[cGA-1].value, row[cPTS-1].value = pr['gf'], pr['ga'], pr['pts']
    if WRITE and changes:
        wb.save(path)
    wb.close()
    return changes


def main():
    print(f"PDF_DIR={PDF_DIR}  WRITE={WRITE}\n")
    for sid in SEASONS:
        ch = process(sid)
        print(f"=== {sid}: {len(ch)} řádků {'ZAPSÁNO' if WRITE else '(dry-run)'} ===")
        for rr, old, new in ch[:3]:
            print(f"   ř.{rr}  W/D/L/GF/GA/PTS {old} -> {new}")
        if len(ch) > 3:
            print(f"   … a dalších {len(ch)-3}")
    if not WRITE:
        print("\n(dry-run) Pro zápis spusť: python3 fill_from_pdf.py", PDF_DIR, "--write")


if __name__ == '__main__':
    main()
