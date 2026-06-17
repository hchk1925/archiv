#!/usr/bin/env python3
"""
solve_flags.py — interaktivní dořešení oflagovaných osudů (postup/sestup u víceFázových
soutěží, které fix_fates.py neuměl rozhodnout: chybí feeds_into i tabulky sesterských fází).

Projde flag po flagu, ukáže kontext (tabulka + sesterské fáze) a nabídne varianty.
Ty ťukneš číslo, napíšeš vlastní text (specifická situace), nebo dáš Enter = skip
(flag zůstane). 'q' = uložit a konec. Zápis do data/S*_FINAL.xlsx průběžně (po sezóně).

  python solve_flags.py --list           # jen vypíše všechny flagy (nic nemění)
  python solve_flags.py                   # interaktivní řešení všech sezón
  python solve_flags.py --season 1950_51  # jen jedna sezóna
"""
import collections
import glob
import os
import re
import sys
import openpyxl

import fix_fates as ff   # sdílené: load_system, sibling_label, lvl, DATA, NON_DATA

OPTS = [('1', 'play off'), ('2', 'o udržení'), ('3', 'o umístění'),
        ('4', 'kvalifikace'), ('5', 'prolínací'), ('6', 'postup (přímý)'),
        ('7', 'sestup (přímý)'), ('8', 'udržel se'), ('9', 'setrval'),
        ('0', 'zůstal')]
VAL = {'1': 'play off', '2': 'o udržení', '3': 'o umístění', '4': 'kvalifikace',
       '5': 'prolínací', '6': 'postup', '7': 'sestup', '8': 'udržel se',
       '9': 'setrval', '0': 'zůstal'}


def _poskey(p):
    m = re.match(r'\d+', str(p) if p is not None else '')
    return int(m.group()) if m else 999


def analyze(path):
    """Vrať flagy = oflagované soutěže s kontextem (stejná podmínka jako fix_fates)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    nodes = ff.load_system(wb)
    members = collections.defaultdict(set)
    rowsby = collections.defaultdict(list)   # node -> [(sheet,row,pos,cid,club,fate)]
    for sh in wb.sheetnames:
        if sh in ff.NON_DATA:
            continue
        rr = list(wb[sh].iter_rows(values_only=True))
        if not rr:
            continue
        H = {c: j for j, c in enumerate(rr[0]) if c}
        if 'season_fate' not in H or 'node_id' not in H or 'row_type' not in H:
            continue
        for i, r in enumerate(rr[1:], start=2):
            if r[H['row_type']] != 'T':
                continue
            nid = r[H['node_id']]
            if not nid:
                continue
            cid = r[H.get('club_id', -1)] if 'club_id' in H else None
            pos = r[H.get('pos', -1)] if 'pos' in H else None
            club = r[H.get('club_name', -1)] if 'club_name' in H else ''
            members[nid].add(cid)
            rowsby[nid].append((sh, i, pos, cid, club, r[H['season_fate']]))
    wb.close()

    def root_of(nid):
        seen = set()
        while nid in nodes and nodes[nid]['parent'] in nodes and nodes[nid]['parent'] not in seen:
            seen.add(nid); nid = nodes[nid]['parent']
        return nid
    comps = collections.defaultdict(list)
    for nid in rowsby:
        comps[root_of(nid)].append(nid)

    items = []
    for root, phase_ids in comps.items():
        tphases = [p for p in phase_ids if rowsby[p]]
        if len(tphases) < 2:
            continue
        base = max(tphases, key=lambda p: len(members[p]))
        bn = nodes.get(base, {})
        wlab = ff.sibling_label(nodes[bn['feeds']]) if bn.get('feeds') in nodes else None
        llab = ff.sibling_label(nodes[bn['loser']]) if bn.get('loser') in nodes else None
        if wlab or llab:
            continue                       # feeds → vyřešeno jinde
        dest = {}
        for s in [p for p in tphases if p != base]:
            lab = ff.sibling_label(nodes.get(s, {}))
            if lab:
                for cid in members[s]:
                    if cid in members[base]:
                        dest.setdefault(cid, lab)
        if dest:
            continue                       # členství → vyřešeno jinde
        unresolved = [r for r in rowsby[base] if r[5] in ('postup', 'sestup')]
        if not unresolved:
            continue
        sibs = [{'name': nodes[s]['name'], 'type': nodes[s]['competition_type'],
                 'pos': sorted([x[2] for x in rowsby[s]], key=_poskey)}
                for s in tphases if s != base]
        items.append({'comp': bn.get('name', base), 'level': bn.get('level', ''),
                      'base': sorted(rowsby[base], key=lambda v: _poskey(v[2])),
                      'unresolved': unresolved, 'sibs': sibs})
    return items


def list_all(paths):
    n = 0
    for path in paths:
        sid = re.search(r'S(\d{4}_\d{2})', os.path.basename(path)).group(1)
        for it in analyze(path):
            n += 1
            print(f"{sid}  {str(it['comp'])[:46]:46} ({len(it['unresolved'])} týmů)")
    print(f"\nCelkem flagů: {n}")


def ask(prompt):
    try:
        return input(prompt).strip()
    except EOFError:
        return 'q'


def solve(paths):
    print("Řešení flagů. U každého týmu: číslo varianty / vlastní text / Enter=skip / q=konec.\n")
    for path in paths:
        sid = re.search(r'S(\d{4}_\d{2})', os.path.basename(path)).group(1)
        items = analyze(path)
        if not items:
            continue
        wb = openpyxl.load_workbook(path)
        fatecol = {}
        for sh in wb.sheetnames:
            if sh in ff.NON_DATA:
                continue
            head = [c.value for c in wb[sh][1]]
            if 'season_fate' in head:
                fatecol[sh] = head.index('season_fate') + 1
        dirty = False
        for it in items:
            print('=' * 72)
            print(f"SEZÓNA {sid}  ·  {it['comp']}  ({it['level']})")
            print("  Tabulka (základní fáze):")
            for (sh, i, pos, cid, club, fate) in it['base']:
                star = '  ←?' if fate in ('postup', 'sestup') else ''
                print(f"     {str(pos):>3}. {str(club)[:30]:30} [{fate}]{star}")
            if it['sibs']:
                print("  Další fáze v soutěži:")
                for s in it['sibs']:
                    print(f"     • {str(s['name'])[:40]:40} [{s['type']}] poz {s['pos']}")
            print("  Varianty: " + " | ".join(f"{k}={v}" for k, v in OPTS))
            for (sh, i, pos, cid, club, fate) in it['unresolved']:
                a = ask(f"   → {pos}. {club}  (teď: {fate})  volba: ")
                if a == 'q':
                    if dirty:
                        wb.save(path)
                    print("Uloženo, končím.")
                    return
                if a == '':
                    continue                      # skip → flag zůstává
                new = VAL.get(a, a)               # číslo → hodnota, jinak vlastní text
                wb[sh].cell(row=i, column=fatecol[sh]).value = new
                dirty = True
                print(f"      ✓ {new}")
        if dirty:
            wb.save(path)
            print(f"[{sid}] uloženo.")
        wb.close()
    print("\nHotovo — všechny sezóny projity.")


def main():
    one = None
    if '--season' in sys.argv:
        one = sys.argv[sys.argv.index('--season') + 1]
    paths = ([os.path.join(ff.DATA, f'S{one}_FINAL.xlsx')] if one
             else sorted(glob.glob(os.path.join(ff.DATA, 'S*_FINAL.xlsx'))))
    if '--list' in sys.argv:
        list_all(paths)
    else:
        solve(paths)


if __name__ == '__main__':
    main()
