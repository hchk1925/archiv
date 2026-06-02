#!/usr/bin/env python3
"""
enrich_1947_48.py — doplní do existující S1947_48 vše ostatní, co PDF popisuje
(mimo už hotovou Státní ligu). NEpřepisuje sešit od nuly (zachová TODO i pole fází).

Přidává:
  - Mistrovství Slovenska (Poprad, 28.–29. 2. 1948) — reálná tabulka 3 týmů
    (Spišská N. Ves vyloučena) → nový list 30SVK + SYSTEM uzel + CLUBS.
  - Bohaté NOTES: anulace divize/žup (počasí), Středočeský divizní pohár (náhradní
    vyřazovák), nedohraná kvalifikace + přechod na 8člennou ligu 1948/49,
    postup do MS divize (Těšetice–Ostravská Slavia), slovenský titulový zápas,
    sloučení Pokojovice+Okříšky, známý výsledek VS župy.
Idempotentní (kontroluje existenci listu/uzlu/klubů).
"""
import openpyxl
from openpyxl.utils import get_column_letter

PATH = 'data/S1947_48_FINAL.xlsx'
S = 'S1947_48'
HDR10 = ['row_type', 'block_name', 'pos', 'club_name', 'note', 'GP', 'W', 'D', 'L',
         'GF', ':', 'GA', 'PTS', 'comp_path', 'node_id', 'level', 'club_id',
         'prev_club_id', 'dest_node_id', 'dest_type', 'season_fate', 'tr_id', 'district']

# Mistrovství Slovenska — (pos, name, GP,W,D,L,GF,GA,PTS, city, note)
SVK = [
    (1, 'HC Tatry Poprad', 2, 2, 0, 0, 37, 2, 4, 'Poprad', 'M'),
    (2, 'TTS Trenčín', 2, 1, 0, 1, 9, 18, 2, 'Trenčín', ''),
    (3, 'ŠK Kežmarok', 2, 0, 0, 2, 5, 31, 0, 'Kežmarok', ''),
]
NODE_SVK = f'NODE_{S}_0005'

NOTES_NEW = [
    'Nižší soutěže (divize i župní) byly 13. 2. 1948 Českým svazem prohlášeny za nesehrané a ANULOVÁNY pro nepřízeň počasí (zimní stadiony odmítly finanční riziko).',
    'Výjimka: mistři župy slezské a hanácké směli sehrát zápas o 6. účastníka skupiny A Moravskoslezské divize.',
    'Mistrovství Slovenska (28.–29. 2. 1948, Poprad): 1. HC Tatry Poprad, 2. TTS Trenčín, 3. ŠK Kežmarok; AC Spišská Nová Ves vyloučena ze soutěže.',
    'Středočeská divizní soutěž nahrazena vyřazovacím turnajem (Středočeský divizní pohár) na ZS Štvanice (1. kolo, čtvrtfinále, semifinále). Finále se nehrálo (mělo být až příští sezónu Slavia–Starokolínský, ale nestalo se); výsledek 2. semifinále není doložen.',
    'Kvalifikace o Státní ligu (turnaj divizních vítězů) se pro počasí nedohrála.',
    'Přechod na jednoskupinovou Státní ligu o 8 účastnících: poslední 2 z každé skupiny sestoupily; bez kvalifikace zařazen nově vzniklý ATK Praha (→ 9 týmů), pak HC Stadion Podolí ukončil činnost → 1948/49 hrálo 8 klubů.',
    'O postup do Moravskoslezské divize: SK Těšetice – SK Ostravská Slavia 1:15; odveta nehrána (Těšetice vzdaly) → postoupila SK Ostravská Slavia.',
    'Po sezóně sloučení HC Viktoria HPH Pokojovice s SK Okříšky.',
    'Východoslovenská župa — jediný doložený výsledek: AC Veľký Šariš – ŠK Slávia Prešov 2:9.',
    'Neoficiální mistr Slovenska (zápas dvou bratislavských klubů ze Státní ligy): ŠK Bratislava – VŠ Bratislava 4:3 (11. 12. 1947).',
    'Původní (nesehrané) rozlosování všech divizí je v podkladech doloženo včetně soupisek skupin.',
]


def main():
    wb = openpyxl.load_workbook(PATH)

    # ── 1) Mistrovství Slovenska — datový list 30SVK ──
    if '30SVK' not in wb.sheetnames:
        ws = wb.create_sheet('30SVK')
        ws.append(HDR10)
        ws.append(['H', 'Mistrovství Slovenska', '', '', '', '', '', '', '', '', '',
                   '', '', '', NODE_SVK, 'L30', '', '', '', '', '', '', ''])
        tr = 1
        for (pos, nm, gp, w, d, l, gf, ga, pts, city, note) in SVK:
            cid = f'CLUB_{S}_{12 + pos:04d}'
            ws.append(['T', '', pos, nm, note, gp, w, d, l, gf, ':', ga, pts,
                       'Mistrovství Slovenska', NODE_SVK, 'L30', cid, '', '', '',
                       '', f'TR_{S}_1{tr:04d}', ''])
            tr += 1
        for j, wdt in enumerate([6, 22, 5, 24, 5, 4, 4, 4, 4, 5, 3, 5, 5, 22, 20, 6, 20, 20, 18, 12, 12, 18, 10], 1):
            ws.column_dimensions[get_column_letter(j)].width = wdt
        print("  ✔ list 30SVK (Mistrovství Slovenska) přidán")

    # ── 2) SYSTEM uzel ──
    sy = wb['SYSTEM']
    have = {sy.cell(r, 1).value for r in range(2, sy.max_row + 1)}
    if NODE_SVK not in have:
        sy.append([NODE_SVK, 'Mistrovství Slovenska', 'regional_championship', 'L30',
                   'REG_SVK', '', '', '', '2-1-0', '', '28.–29. 2. 1948, Poprad. '
                   'AC Spišská Nová Ves vyloučena.', '', '', ''])
        print("  ✔ SYSTEM uzel Mistrovství Slovenska")

    # ── 3) CLUBS ──
    cl = wb['CLUBS']
    hdr = [c.value for c in cl[1]]
    ci = hdr.index('club_id')
    have_c = {cl.cell(r, ci + 1).value for r in range(2, cl.max_row + 1)}
    svk_clubs = [(13, 'HC Tatry Poprad', 'Poprad', 'M'),
                 (14, 'TTS Trenčín', 'Trenčín', ''),
                 (15, 'ŠK Kežmarok', 'Kežmarok', ''),
                 (16, 'AC Spišská Nová Ves', 'Spišská Nová Ves', '')]
    # CLUBS pořadí: club_id, clean_name, raw_name, sheet, entry_note, level, district, prev_club_id, change_note, city
    for (n, nm, city, m) in svk_clubs:
        cid = f'CLUB_{S}_{n:04d}'
        if cid in have_c:
            continue
        note = 'Mistrovství Slovenska (Poprad). prev_club_id zatím prázdný.'
        if nm == 'AC Spišská Nová Ves':
            note = 'Mistrovství Slovenska — vyloučena ze soutěže (bez záznamu v tabulce).'
        cl.append([cid, nm, nm + (' [M-SVK]' if m else ''),
                   '30SVK' if nm != 'AC Spišská Nová Ves' else '', m, 'L30', '', '',
                   note, city])
    print("  ✔ CLUBS Slovenska doplněny")

    # ── 4) NOTES (přidat chybějící) ──
    no = wb['NOTES']
    existing = {no.cell(r, 1).value for r in range(2, no.max_row + 1)}
    for t in NOTES_NEW:
        if t not in existing:
            no.append([t])
    print("  ✔ NOTES rozšířeny")

    # ── 5) META ──
    me = wb['META']
    mk = {me.cell(r, 1).value: r for r in range(2, me.max_row + 1)}
    def setmeta(k, v):
        if k in mk:
            me.cell(mk[k], 2).value = v
        else:
            me.append([k, v])
    setmeta('mistr_slovenska', 'HC Tatry Poprad')
    setmeta('total_clubs', '16')
    setmeta('total_nodes', '5')
    setmeta('status', 'PARTIAL — Státní liga + Mistrovství Slovenska (tabulky). '
                      'Divize a župní soutěže anulovány (počasí) — jen poznámky.')
    print("  ✔ META aktualizována")

    wb.save(PATH)
    print(f"=== {PATH} obohaceno ===")


if __name__ == '__main__':
    main()
