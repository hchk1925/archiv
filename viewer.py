#!/usr/bin/env python3
"""
viewer.py — základní prohlížeč: vybranou sezónu z xlsx zobrazí jako PDF (tabulky).
Otevře se v samostatném okně z rozcestníku Almanach.py. Potřebuje openpyxl + reportlab.
"""
import collections
import glob
import io
import os
import re
import subprocess
import sys
import tempfile
import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl

NON_DATA = {'META', 'CLUBS', 'SYSTEM', 'NOTES', 'TODO', 'CHANGES', 'README', 'KOSILKA'}


def _open_file(path):
    try:
        if sys.platform.startswith('win'):
            os.startfile(path)  # noqa
        elif sys.platform == 'darwin':
            subprocess.run(['open', path])
        else:
            subprocess.run(['xdg-open', path])
    except Exception:
        pass


def _lvlnum(v):
    m = re.search(r'(\d+)', str(v or ''))
    return int(m.group(1)) if m else 9999


def _poskey(p):
    m = re.match(r'\d+', str(p) if p is not None else '')
    return int(m.group()) if m else 999


def season_pdf_bytes(path, label):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    Table, TableStyle, KeepTogether)
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fn, fb = 'Helvetica', 'Helvetica-Bold'
    for d in ('/usr/share/fonts/truetype/dejavu', r'C:\Windows\Fonts',
              '/Library/Fonts', '/usr/share/fonts'):
        reg = os.path.join(d, 'DejaVuSans.ttf'); bold = os.path.join(d, 'DejaVuSans-Bold.ttf')
        if os.path.exists(reg) and os.path.exists(bold):
            try:
                pdfmetrics.registerFont(TTFont('DV', reg))
                pdfmetrics.registerFont(TTFont('DVB', bold))
                fn, fb = 'DV', 'DVB'
                break
            except Exception:
                pass

    wb = openpyxl.load_workbook(path, data_only=True)
    nodes = {}
    if 'SYSTEM' in wb.sheetnames:
        sr = list(wb['SYSTEM'].iter_rows(values_only=True)); SH = {c: j for j, c in enumerate(sr[0]) if c}
        for r in sr[1:]:
            nid = r[SH['node_id']] if 'node_id' in SH else None
            if nid:
                nodes[nid] = {'name': r[SH.get('name', -1)] if 'name' in SH else '',
                              'level': r[SH.get('level', -1)] if 'level' in SH else ''}
    byn = collections.defaultdict(list)
    for sh in wb.sheetnames:
        if sh in NON_DATA:
            continue
        rr = list(wb[sh].iter_rows(values_only=True))
        if not rr:
            continue
        H = {c: j for j, c in enumerate(rr[0]) if c}
        if 'node_id' not in H or 'row_type' not in H:
            continue
        for r in rr[1:]:
            if r[H['row_type']] != 'T':
                continue
            nid = r[H['node_id']]
            if not nid:
                continue
            g = lambda k: (r[H[k]] if k in H else None)
            byn[nid].append((g('pos'), g('club_name'), g('GP'), g('W'), g('D'),
                             g('L'), g('GF'), g('GA'), g('PTS'), g('season_fate')))
    wb.close()

    order = sorted(byn.keys(), key=lambda n: (_lvlnum(nodes.get(n, {}).get('level')),
                                              str(nodes.get(n, {}).get('name', ''))))
    NAVY = colors.HexColor('#1a3050')
    H1 = ParagraphStyle('H1', fontName=fb, fontSize=16, textColor=NAVY, spaceAfter=2)
    sub = ParagraphStyle('sub', fontName=fn, fontSize=9, textColor=colors.HexColor('#666'), spaceAfter=8)
    H2 = ParagraphStyle('H2', fontName=fb, fontSize=10.5, textColor=NAVY, spaceBefore=8, spaceAfter=3)
    cell = ParagraphStyle('cell', fontName=fn, fontSize=7, leading=8)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=13 * mm, rightMargin=13 * mm,
                            topMargin=13 * mm, bottomMargin=12 * mm, title=f'Almanach {label}')
    el = [Paragraph(f'Almanach — sezóna {label}', H1),
          Paragraph(f'{len(order)} soutěží / tabulek', sub)]
    for nid in order:
        nm = nodes.get(nid, {}).get('name') or nid
        lv = nodes.get(nid, {}).get('level') or ''
        data = [['#', 'Klub', 'GP', 'W', 'D', 'L', 'GF:GA', 'PTS', 'Fate']]
        for i, (pos, club, gp, w, dd, l, gf, ga, pts, fate) in enumerate(
                sorted(byn[nid], key=lambda x: _poskey(x[0])), 1):
            dpos = pos if pos not in (None, '') else i
            data.append([dpos, Paragraph(str(club or ''), cell), gp or '', w or '',
                         dd or '', l or '', f'{gf if gf not in (None,"") else ""}:'
                         f'{ga if ga not in (None,"") else ""}', pts or '',
                         Paragraph(str(fate or ''), cell)])
        t = Table(data, colWidths=[8 * mm, 66 * mm, 9 * mm, 8 * mm, 8 * mm, 8 * mm,
                                   16 * mm, 9 * mm, 30 * mm])
        t.setStyle(TableStyle([
            ('FONT', (0, 0), (-1, -1), fn, 7), ('FONT', (0, 0), (-1, 0), fb, 7),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e8edf5')),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#cccccc')),
            ('ALIGN', (2, 0), (-1, -1), 'CENTER'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
        el.append(KeepTogether([Paragraph(f'{nm}   ({lv})', H2), t, Spacer(1, 3)]))
    doc.build(el)
    return buf.getvalue()


class Viewer(tk.Toplevel):
    def __init__(self, master, data_dir):
        super().__init__(master)
        self.title('Prohlížeč sezón — PDF')
        self.geometry('420x180')
        self.data_dir = data_dir
        self.seasons = [re.search(r'S(\d{4}_\d{2})', os.path.basename(p)).group(1)
                        for p in sorted(glob.glob(os.path.join(data_dir, 'S*_FINAL.xlsx')))]
        frm = ttk.Frame(self); frm.pack(expand=True)
        ttk.Label(frm, text='Vyber sezónu a zobraz ji jako PDF:',
                  font=('TkDefaultFont', 11, 'bold')).pack(pady=(14, 8))
        row = ttk.Frame(frm); row.pack()
        ttk.Label(row, text='Sezóna:').pack(side='left')
        self.sv = tk.StringVar(value=self.seasons[0] if self.seasons else '')
        ttk.Combobox(row, textvariable=self.sv, values=self.seasons, width=12,
                     state='readonly').pack(side='left', padx=4)
        ttk.Button(row, text='📄 Zobrazit PDF', command=self.show).pack(side='left', padx=6)
        self.status = ttk.Label(frm, text='', foreground='#666'); self.status.pack(pady=10)

    def show(self):
        sid = self.sv.get()
        if not sid:
            return
        path = os.path.join(self.data_dir, f'S{sid}_FINAL.xlsx')
        self.status.config(text='Generuji PDF…'); self.update_idletasks()
        try:
            pdf = season_pdf_bytes(path, sid.replace('_', '/'))
        except ImportError:
            messagebox.showerror('Chybí reportlab',
                                 'Pro PDF je potřeba knihovna reportlab.\n'
                                 'Nainstaluj ji:  pip install reportlab')
            self.status.config(text=''); return
        except Exception as e:
            messagebox.showerror('PDF', f'Nepodařilo se vytvořit PDF:\n{e}')
            self.status.config(text=''); return
        out = os.path.join(tempfile.gettempdir(), f'almanach_{sid}.pdf')
        with open(out, 'wb') as f:
            f.write(pdf)
        _open_file(out)
        self.status.config(text=f'Otevřeno: {out}')


def main():
    root = tk.Tk(); root.withdraw()
    from tkinter import filedialog
    dd = None
    for c in ('xlsx', 'data', '.'):
        if glob.glob(os.path.join(c, 'S*_FINAL.xlsx')):
            dd = c; break
    if not dd:
        dd = filedialog.askdirectory(title='Vyber složku se sešity S*_FINAL.xlsx')
    if dd and glob.glob(os.path.join(dd, 'S*_FINAL.xlsx')):
        v = Viewer(root, dd)
        v.protocol('WM_DELETE_WINDOW', root.destroy)
        root.mainloop()


if __name__ == '__main__':
    main()
