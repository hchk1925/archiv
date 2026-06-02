#!/usr/bin/env python3
"""
fix_missing_links.py — doplnění chybějících prev_club_id (rozrod klubů).

Pro každou dvojici po sobě jdoucích sezón najde klub bez prev_club_id, jehož
clean_name se PŘESNĚ shoduje s právě jedním klubem předchozí sezóny, který:
  - není B/juniorský tým (vlastní identita dle D20),
  - ještě není ničím předchůdcem (volný).

Takový pár je jednoznačné pokračování klubu → doplní prev_club_id.
Konzervativní: žádné fuzzy/přejmenování (to vyžaduje registr, řeší se ručně).
Re-runnable.
"""
import openpyxl, glob, re, sys
from collections import defaultdict

DATA = 'data'
WRITE = '--write' in sys.argv
B = re.compile(r'(\bB\b|\bII+\b|\bIII\b|\bC\b|"B"|"C"|\bjun)', re.I)


def read_clubs(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['CLUBS']
    rows = list(ws.iter_rows(values_only=True))
    H = {h: i for i, h in enumerate(rows[0])}
    out = {}
    for r in rows[1:]:
        if not r or r[H['club_id']] is None:
            continue
        out[r[H['club_id']]] = {'name': str(r[H['clean_name']] or '').strip(),
                                'prev': r[H['prev_club_id']]}
    wb.close()
    return out


def main():
    files = sorted(glob.glob(f'{DATA}/S*_FINAL.xlsx'))
    prev = None
    plan = []                      # (path, club_id, name, prev_id)
    for path in files:
        cur = read_clubs(path)
        if prev is not None:
            byname = defaultdict(list)
            for pid, pc in prev.items():
                byname[pc['name'].lower()].append(pid)
            claimed = {c['prev'] for c in cur.values() if c['prev']}
            for cid, c in cur.items():
                if c['prev'] or B.search(c['name']):
                    continue
                cands = [p for p in byname.get(c['name'].lower(), [])
                         if p not in claimed and not B.search(prev[p]['name'])]
                if len(cands) == 1:
                    plan.append((path, cid, c['name'], cands[0]))
        prev = cur

    by_file = defaultdict(list)
    for path, cid, nm, pid in plan:
        by_file[path].append((cid, nm, pid))
    for path, items in by_file.items():
        if WRITE:
            wb = openpyxl.load_workbook(path)
            ws = wb['CLUBS']
            hdr = [c.value for c in ws[1]]
            ci, pi = hdr.index('club_id'), hdr.index('prev_club_id')
            byid = {ws.cell(r, ci+1).value: r for r in range(2, ws.max_row+1)}
            for cid, nm, pid in items:
                ws.cell(byid[cid], pi+1).value = pid
            wb.save(path)
        for cid, nm, pid in items:
            print(f"  {'✔' if WRITE else '·'} {path.split('/')[-1][1:8]} "
                  f"{nm[:32]:32} prev <- {pid}")
    print(f"\n=== {'ZAPSÁNO' if WRITE else 'dry-run'}: {len(plan)} linků ===")


if __name__ == '__main__':
    main()
