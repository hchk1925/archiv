#!/usr/bin/env python3
"""
solve_flags_gui.py — Tkinter appka na dořešení oflagovaných osudů (postup/sestup
u víceFázových soutěží, které automat neuměl rozhodnout).

Samostatný soubor (potřebuje jen Python + openpyxl; tkinter je součást Pythonu).
Dej tento soubor do složky, kde máš sešity S*_FINAL.xlsx (nebo do složky, která má
podsložku data/), a spusť:

    python solve_flags_gui.py        (na Windows dvojklik na run_solve.bat)

Co umí:
  • projíždíš flagy jeden po druhém, vidíš tabulku + další fáze soutěže (kontext),
  • u sporného týmu ťukneš variantu (tlačítko) / napíšeš vlastní / Přeskočíš,
  • PRŮBĚŽNĚ se to zapisuje do xlsx a stav se ukládá (solve_progress.json),
  • můžeš kdykoliv zavřít a vrátit se — naváže tam, kde jsi skončil,
  • tlačítko „Zabalit do ZIP" vytvoří zip se všemi sešity + vysvětlujícím textem.
"""
import datetime
import glob
import json
import os
import re
import shutil
import sys
import zipfile
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import collections
import openpyxl

NON_DATA = {'META', 'CLUBS', 'SYSTEM', 'NOTES', 'TODO', 'CHANGES', 'README'}
PLAYOFF_T = {'playoff_round', 'playoff', 'final_group', 'final_series', 'series'}
RELEG_T = {'relegation_group', 'relegation_playoff', 'relegation_series'}
CLASS_T = {'classification'}
QUAL_T = {'qualification_group', 'qualification_series', 'qualification', 'baraz'}


def _lvl(v):
    m = re.search(r'(\d+)', str(v or ''))
    return int(m.group(1)) if m else None


def _poskey(p):
    m = re.match(r'\d+', str(p) if p is not None else '')
    return int(m.group()) if m else 999


def _sibling_label(node):
    t = str(node.get('competition_type') or '')
    nm = str(node.get('name') or '').lower()
    if t in PLAYOFF_T or any(w in nm for w in
            ('play off', 'playoff', 'play-off', 'finále', 'čtvrtfinále',
             'semifinále', 'předkolo')):
        return 'play off'
    if t in RELEG_T or 'udržení' in nm or 'záchran' in nm:
        return 'o udržení'
    if t in CLASS_T or 'o umístění' in nm or re.search(r'\bo \d+\.', nm):
        return 'o umístění'
    if t in QUAL_T or 'kvalif' in nm or 'baráž' in nm or 'prolín' in nm:
        return 'prolínací' if ('baráž' in nm or 'prolín' in nm) else 'kvalifikace'
    return None


def _load_system(wb):
    ws = wb['SYSTEM']
    rows = list(ws.iter_rows(values_only=True))
    H = {c: j for j, c in enumerate(rows[0]) if c}
    nodes = {}
    for r in rows[1:]:
        nid = r[H['node_id']] if 'node_id' in H else None
        if not nid:
            continue
        nodes[nid] = {
            'node_id': nid,
            'name': r[H.get('name', -1)] if 'name' in H else '',
            'competition_type': r[H.get('competition_type', -1)] if 'competition_type' in H else '',
            'level': r[H.get('level', -1)] if 'level' in H else '',
            'parent': r[H.get('parent_node_id', -1)] if 'parent_node_id' in H else None,
            'feeds': r[H.get('feeds_into', -1)] if 'feeds_into' in H else None,
            'loser': r[H.get('feeds_into_loser', -1)] if 'feeds_into_loser' in H else None,
        }
    return nodes


def analyze(path):
    """Oflagované soutěže (víc fází, postup/sestup nešel auto-rozhodnout) + kontext."""
    wb = openpyxl.load_workbook(path, data_only=True)
    nodes = _load_system(wb)
    members = collections.defaultdict(set)
    rowsby = collections.defaultdict(list)
    for sh in wb.sheetnames:
        if sh in NON_DATA:
            continue
        rr = list(wb[sh].iter_rows(values_only=True))
        if not rr:
            continue
        H = {c: j for j, c in enumerate(rr[0]) if c}
        if 'season_fate' not in H or 'node_id' not in H or 'row_type' not in H:
            continue
        for i, r in enumerate(rr[1:], start=2):
            if r[H['row_type']] != 'T':
                continue
            nid = r[H['node_id']]
            if not nid:
                continue
            cid = r[H.get('club_id', -1)] if 'club_id' in H else None
            pos = r[H.get('pos', -1)] if 'pos' in H else None
            club = r[H.get('club_name', -1)] if 'club_name' in H else ''
            members[nid].add(cid)
            rowsby[nid].append((sh, i, pos, cid, club, r[H['season_fate']]))
    wb.close()

    def root_of(nid):
        seen = set()
        while nid in nodes and nodes[nid]['parent'] in nodes and nodes[nid]['parent'] not in seen:
            seen.add(nid); nid = nodes[nid]['parent']
        return nid
    comps = collections.defaultdict(list)
    for nid in rowsby:
        comps[root_of(nid)].append(nid)

    items = []
    for root, phase_ids in comps.items():
        tphases = [p for p in phase_ids if rowsby[p]]
        if len(tphases) < 2:
            continue
        base = max(tphases, key=lambda p: len(members[p]))
        bn = nodes.get(base, {})
        wlab = _sibling_label(nodes[bn['feeds']]) if bn.get('feeds') in nodes else None
        llab = _sibling_label(nodes[bn['loser']]) if bn.get('loser') in nodes else None
        if wlab or llab:
            continue
        dest = {}
        for s in [p for p in tphases if p != base]:
            lab = _sibling_label(nodes.get(s, {}))
            if lab:
                for cid in members[s]:
                    if cid in members[base]:
                        dest.setdefault(cid, lab)
        if dest:
            continue
        unresolved = [r for r in rowsby[base] if r[5] in ('postup', 'sestup')]
        if not unresolved:
            continue
        ordered = sorted(tphases, key=lambda p: (_lvl(nodes.get(p, {}).get('level')) or 0,
                                                 str(nodes.get(p, {}).get('name', ''))))
        phases = [{'name': nodes.get(p, {}).get('name', p),
                   'type': nodes.get(p, {}).get('competition_type', ''),
                   'rows': sorted(rowsby[p], key=lambda v: _poskey(v[2]))}
                  for p in ([base] + [p for p in ordered if p != base])]
        items.append({'comp': bn.get('name', base), 'level': bn.get('level', ''),
                      'base': sorted(rowsby[base], key=lambda v: _poskey(v[2])),
                      'unresolved': unresolved, 'phases': phases})
    return items


TORZO_PHASE = re.compile(
    r'play.?off|finále|čtvrtfinále|semifinále|předkolo|o umístění|o udržení|'
    r'skupina o|baráž|kvalif|nadstavba|o \d+\.\s*m[ií]sto|prolín', re.I)
TORZO_TYPES = {'league', 'group', 'regional_championship', ''}


def scan_torzo(path, thresh=0.8):
    """Tabulky-torza: mají vyplněné pořadí, ale odehráno < thresh jednoho kola (N-1)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    nodes = _load_system(wb)
    rowsby = collections.defaultdict(list)
    for sh in wb.sheetnames:
        if sh in NON_DATA:
            continue
        rr = list(wb[sh].iter_rows(values_only=True))
        if not rr:
            continue
        H = {c: j for j, c in enumerate(rr[0]) if c}
        if 'node_id' not in H or 'row_type' not in H:
            continue
        for i, r in enumerate(rr[1:], start=2):
            if r[H['row_type']] != 'T':
                continue
            nid = r[H['node_id']]
            if not nid:
                continue
            pos = r[H.get('pos', -1)] if 'pos' in H else None
            gp = r[H.get('GP', -1)] if 'GP' in H else None
            club = r[H.get('club_name', -1)] if 'club_name' in H else ''
            fate = r[H.get('season_fate', -1)] if 'season_fate' in H else ''
            rowsby[nid].append((sh, i, pos, gp, club, fate))
    wb.close()
    items = []
    for nid, rows in rowsby.items():
        nd = nodes.get(nid, {})
        if str(nd.get('competition_type') or '') not in TORZO_TYPES:
            continue
        if TORZO_PHASE.search(str(nd.get('name') or '')):
            continue
        N = len(rows)
        if N < 4 or N > 18:                          # ne registry/agregáty
            continue
        nposs = sum(1 for r in rows if r[2] not in (None, ''))
        if nposs < max(3, 0.6 * N):                  # pořadí reálně není
            continue
        gps = [r[3] for r in rows if isinstance(r[3], (int, float))]
        if len(gps) < max(3, 0.5 * N):               # málo GP dat → nelze posoudit
            continue
        comp = (sum(gps) / len(gps)) / (N - 1) if N > 1 else 1.0
        if comp < thresh:
            items.append({'name': str(nd.get('name') or nid), 'level': nd.get('level', ''),
                          'node': nid, 'pct': round(comp * 100),
                          'rows': sorted(rows, key=lambda v: _poskey(v[2])),
                          'sheet': rows[0][0]})
    return items


VARIANTS = [('play off', 'play off'), ('o udržení', 'o udržení'),
            ('o umístění', 'o umístění'), ('kvalifikace', 'kvalifikace'),
            ('prolínací', 'prolínací'), ('postup (přímý)', 'postup'),
            ('sestup (přímý)', 'sestup'), ('udržel se', 'udržel se'),
            ('setrval', 'setrval'), ('zůstal', 'zůstal')]
SKIP = '__skip__'
CHAT = '__chat__'


def _app_workspace():
    """Jako .exe: vedle exe vznikne přehledná složka pojmenovaná podle programu
    (Almanach), do ní se rozbalí data. Mimo .exe: None."""
    if getattr(sys, 'frozen', False):
        exe_dir = os.path.dirname(sys.executable)
        name = os.path.splitext(os.path.basename(sys.executable))[0]
        return os.path.join(exe_dir, name)
    return None


def find_data_dir():
    # pracovní sešity hledá ve složce 'xlsx' (preferovaně) nebo 'data'; i jako .exe
    base = (os.path.dirname(sys.executable) if getattr(sys, 'frozen', False)
            else os.path.abspath(os.path.dirname(__file__)))
    roots = []
    ws = _app_workspace()
    if ws:
        roots.append(ws)
    roots += [base, os.getcwd()]
    cands = []
    for r in roots:
        cands += [os.path.join(r, 'xlsx'), os.path.join(r, 'data'), r]
    for cand in cands:
        if glob.glob(os.path.join(cand, 'S*_FINAL.xlsx')):
            return cand
    return None


def season_ids(data_dir):
    return [re.search(r'S(\d{4}_\d{2})', os.path.basename(p)).group(1)
            for p in sorted(glob.glob(os.path.join(data_dir, 'S*_FINAL.xlsx')))]


def build_queue(data_dir):
    items = []
    for path in sorted(glob.glob(os.path.join(data_dir, 'S*_FINAL.xlsx'))):
        sid = re.search(r'S(\d{4}_\d{2})', os.path.basename(path)).group(1)
        for fl in analyze(path):
            for (sh, row, pos, cid, club, fate) in fl['unresolved']:
                items.append({
                    'sid': sid, 'path': path, 'sheet': sh, 'row': row,
                    'pos': pos, 'club': club, 'fate': fate, 'kind': 'fate',
                    'comp': fl['comp'], 'level': fl['level'],
                    'base': fl['base'], 'phases': fl['phases'],
                    'key': f'{sid}|{sh}|{row}'})
        for tz in scan_torzo(path):
            items.append({
                'sid': sid, 'path': path, 'sheet': tz['sheet'], 'kind': 'torzo',
                'comp': tz['name'], 'level': tz['level'], 'node': tz['node'],
                'pct': tz['pct'],
                'phases': [{'name': tz['name'],
                            'type': f"tabulka · odehráno ~{tz['pct']}%",
                            'rows': [(sh, i, pos, None, club, fate)
                                     for (sh, i, pos, gp, club, fate) in tz['rows']]}],
                'key': f"{sid}|torzo|{tz['node']}"})
    return items


class FlagSolver(tk.Toplevel):
    def __init__(self, master, data_dir):
        super().__init__(master)
        self.title('Almanach — dořešení osudů (flagy)')
        self.geometry('1180x780')
        self.data_dir = data_dir
        self.prog_path = os.path.join(data_dir, 'solve_progress.json')
        self.progress = self._load_progress()
        self.queue = build_queue(data_dir)
        self.wbcache = {}
        self.idx = 0
        self._build_ui()
        self.protocol('WM_DELETE_WINDOW', self._on_close)
        self._goto_first_pending()
        self.show()

    # -- stav --
    def _load_progress(self):
        try:
            return json.load(open(self.prog_path, encoding='utf-8'))
        except Exception:
            return {}
    def _save_progress(self):
        try:
            json.dump(self.progress, open(self.prog_path, 'w', encoding='utf-8'),
                      ensure_ascii=False, indent=0)
        except Exception:
            pass
    def _done_count(self):
        return sum(1 for v in self.progress.values() if v and v != SKIP)

    def _wb(self, path):
        if path not in self.wbcache:
            self.wbcache[path] = openpyxl.load_workbook(path)
        return self.wbcache[path]
    def _fatecol(self, wb, sheet):
        head = [c.value for c in wb[sheet][1]]
        return head.index('season_fate') + 1 if 'season_fate' in head else None

    # -- UI --
    def _build_ui(self):
        top = ttk.Frame(self); top.pack(fill='x', padx=8, pady=6)
        ttk.Button(top, text='❓ Nápověda', command=self.show_help).pack(side='left')
        ttk.Button(top, text='🏷 Opravit názvy klubů',
                   command=self.open_clubs).pack(side='left', padx=8)
        ttk.Button(top, text='🔗 Návaznost klubů (sezóna→sezóna)',
                   command=self.open_continuity).pack(side='left')
        ttk.Button(top, text='💾 Hotovo → zabalit do ZIP',
                   command=self.export_zip).pack(side='right')
        self.count_lbl = ttk.Label(top, text='', font=('TkDefaultFont', 10, 'bold'))
        self.count_lbl.pack(side='right', padx=12)

        self.pbar = ttk.Progressbar(self, mode='determinate')
        self.pbar.pack(fill='x', padx=8)

        mid = ttk.Frame(self); mid.pack(fill='both', expand=True, padx=8, pady=6)
        self.title_lbl = ttk.Label(mid, text='', font=('TkDefaultFont', 13, 'bold'),
                                   wraplength=1100, justify='left')
        self.title_lbl.pack(anchor='w')

        # kompletní tabulky všech fází soutěže (fáze → týmy)
        tf = ttk.Frame(mid); tf.pack(fill='both', expand=True, pady=4)
        cols = ('pos', 'klub', 'osud')
        self.tbl = ttk.Treeview(tf, columns=cols, show='tree headings', height=15)
        self.tbl.heading('#0', text='FÁZE / SKUPINA'); self.tbl.column('#0', width=240)
        for c, w in zip(cols, (50, 330, 200)):
            self.tbl.heading(c, text=c.upper()); self.tbl.column(c, width=w, anchor='w')
        self.tbl.column('pos', anchor='center')
        ysb = ttk.Scrollbar(tf, command=self.tbl.yview); ysb.pack(side='right', fill='y')
        self.tbl.configure(yscrollcommand=ysb.set)
        self.tbl.pack(side='left', fill='both', expand=True)
        self.tbl.tag_configure('q', background='#fff6da')          # sporný (postup/sestup)
        self.tbl.tag_configure('cur', background='#ffe0a3')         # právě řešený tým
        self.tbl.tag_configure('phase', font=('TkDefaultFont', 9, 'bold'),
                               foreground='#1a3050')

        self.team_lbl = ttk.Label(mid, text='', font=('TkDefaultFont', 12, 'bold'),
                                  foreground='#b07d10')
        self.team_lbl.pack(anchor='w')

        self.action = ttk.Frame(mid); self.action.pack(anchor='w', pady=6)
        self.bf = ttk.Frame(self.action)        # varianty osudu (fate)
        self._keys = '1234567890'
        for i, (label, val) in enumerate(VARIANTS):
            ttk.Button(self.bf, text=f'{self._keys[i]}  {label}', width=20,
                       command=lambda v=val: self.choose(v)).grid(
                           row=i // 5, column=i % 5, padx=3, pady=3, sticky='w')
        self.tzf = ttk.Frame(self.action)       # torzo: ponechat / odstranit pořadí
        ttk.Button(self.tzf, text='1  Ponechat pořadí', width=24,
                   command=self.keep_order).grid(row=0, column=0, padx=3, pady=3)
        ttk.Button(self.tzf, text='2  Odstranit pořadí (jen seznam týmů) + poznámka',
                   width=46, command=self.remove_order).grid(row=0, column=1, padx=3, pady=3)
        self.bf.pack(anchor='w')

        cf = ttk.Frame(mid); cf.pack(anchor='w', pady=4)
        ttk.Label(cf, text='Můj koment / vlastní:').pack(side='left')
        self.custom = ttk.Entry(cf, width=46); self.custom.pack(side='left', padx=4)
        self.custom.bind('<Return>', lambda e: self._to_prompt())
        ttk.Button(cf, text='➜ Do promptu pro chat (p)', command=self._to_prompt).pack(side='left')
        ttk.Button(cf, text='Zapsat jako osud', command=self._choose_custom).pack(side='left', padx=6)
        ttk.Button(cf, text='Přeskočit (s)', command=self.skip).pack(side='left', padx=12)

        nav = ttk.Frame(mid); nav.pack(anchor='w', pady=6)
        ttk.Button(nav, text='◀ zpět', command=lambda: self.move(-1)).pack(side='left')
        ttk.Button(nav, text='další ▶', command=lambda: self.move(1)).pack(side='left', padx=6)
        ttk.Button(nav, text='⏭ další nevyřešený (n)',
                   command=self.next_pending).pack(side='left', padx=6)
        ttk.Label(nav, text='   (←/→ = zpět/další · n = další nevyřešený · '
                            's = přeskočit · p = do chatu)',
                  foreground='#888').pack(side='left')
        self.bind('<Key>', self._key)
        self.status = ttk.Label(self, text='', relief='sunken', anchor='w')
        self.status.pack(fill='x', side='bottom')

    def _key(self, e):
        if self.focus_get() is self.custom:
            return
        ch = (e.char or '').lower()
        it = self.queue[self.idx] if self.queue else None
        torzo = bool(it and it['kind'] == 'torzo')
        if ch == 's':
            self.skip()
        elif ch == 'p':
            self._to_prompt()
        elif ch == 'n':
            self.next_pending()
        elif e.keysym == 'Right':
            self.move(1)
        elif e.keysym == 'Left':
            self.move(-1)
        elif torzo:
            if ch == '1':
                self.keep_order()
            elif ch == '2':
                self.remove_order()
        elif ch in self._keys:
            self.choose(VARIANTS[self._keys.index(ch)][1])

    # -- navigace/zobrazení --
    def _goto_first_pending(self):
        for i, it in enumerate(self.queue):
            st = self.progress.get(it['key'])
            if not st or st == SKIP:
                self.idx = i; return
        self.idx = 0

    def show(self):
        n = len(self.queue)
        done = self._done_count()
        chat = sum(1 for v in self.progress.values() if v == CHAT)
        skip = sum(1 for v in self.progress.values() if v == SKIP)
        self.pbar['maximum'] = max(1, n); self.pbar['value'] = done + chat
        self.count_lbl.config(text=f'{done} vyřešeno · {chat} do chatu · {skip} skip  /  {n}')
        if not self.queue:
            self.title_lbl.config(text='Žádné flagy — vše vyřešeno 🎉')
            return
        it = self.queue[self.idx]
        self.title_lbl.config(
            text=f"[{self.idx+1}/{n}]  {it['sid'].replace('_','/')}  ·  {it['comp']}  ({it['level']})")
        self.tbl.delete(*self.tbl.get_children())
        for ph in it['phases']:
            pid = self.tbl.insert('', 'end', text=f"{ph['name']}  [{ph['type']}]",
                                  open=True, tags=('phase',))
            for k, (sh, row, pos, cid, club, fate) in enumerate(ph['rows'], 1):
                dpos = pos if pos not in (None, '') else k
                key = f"{it['sid']}|{sh}|{row}"
                done = self.progress.get(key)
                if done == CHAT:
                    mark = '  ➜ chat'
                elif done and done != SKIP:
                    mark = ' ✔ ' + done
                elif fate in ('postup', 'sestup'):
                    mark = '  ⟵ ?'
                else:
                    mark = ''
                tag = 'cur' if (sh == it['sheet'] and row == it.get('row')) else (
                    'q' if fate in ('postup', 'sestup') else '')
                self.tbl.insert(pid, 'end', values=(dpos, club, f'{fate or "—"}{mark}'),
                                tags=(tag,) if tag else ())
        st = self.progress.get(it['key'])
        stx = ('' if not st else ('  — přeskočeno' if st == SKIP else
               ('  — odloženo do promptu pro chat' if st == CHAT
                else f'  — již: {st}')))
        if it['kind'] == 'torzo':
            self.team_lbl.config(text=f"⚠ TORZO — pořadí je, ale odehráno jen ~{it['pct']}% "
                                      f"jednoho kola. Ponechat, nebo odstranit pořadí?{stx}")
            self.bf.pack_forget(); self.tzf.pack(anchor='w')
        else:
            pre = f"{it['pos']}. " if it['pos'] not in (None, '') else ''
            self.team_lbl.config(text=f"→  {pre}{it['club']}   (teď: {it['fate']}){stx}")
            self.tzf.pack_forget(); self.bf.pack(anchor='w')
        self.custom.delete(0, 'end')

    def move(self, d):
        if self.queue:
            self.idx = (self.idx + d) % len(self.queue)
            self.show()

    def choose(self, val):
        it = self.queue[self.idx]
        if it['kind'] == 'torzo':
            return
        wb = self._wb(it['path']); col = self._fatecol(wb, it['sheet'])
        if col:
            wb[it['sheet']].cell(row=it['row'], column=col).value = val
            wb.save(it['path'])
        self.progress[it['key']] = val
        self._save_progress()
        self.status.config(text=f"Zapsáno: {it['club']} → {val}")
        self._advance()

    def keep_order(self):
        it = self.queue[self.idx]
        if it['kind'] != 'torzo':
            return
        self.progress[it['key']] = 'pořadí ponecháno'
        self._save_progress()
        self.status.config(text=f"Torzo: pořadí ponecháno — {it['comp']}")
        self._advance()

    def remove_order(self):
        it = self.queue[self.idx]
        if it['kind'] != 'torzo':
            return
        wb = self._wb(it['path'])
        for (sh, row, pos, cid, club, fate) in it['phases'][0]['rows']:
            ws = wb[sh]; head = [c.value for c in ws[1]]
            if 'pos' in head:
                ws.cell(row=row, column=head.index('pos') + 1).value = None
        if 'NOTES' in wb.sheetnames:
            nws = wb['NOTES']; nh = [c.value for c in nws[1]]
            ncol = nh.index('node_id') if 'node_id' in nh else None
            scol = nh.index('source_type') if 'source_type' in nh else None
            exists = ncol is not None and any(
                r[ncol] == it['node'] and (scol is None or r[scol] == 'torzo-manual')
                for r in nws.iter_rows(min_row=2, values_only=True))
            if not exists:                       # ať nevznikne duplicitní poznámka
                nr = [None] * len(nh)
                for col, val in (('node_id', it['node']),
                                 ('note_text', f"torzo: pořadí odstraněno "
                                               f"(odehráno ~{it['pct']}% jednoho kola)"),
                                 ('source_type', 'torzo-manual')):
                    if col in nh:
                        nr[nh.index(col)] = val
                nws.append(nr)
        wb.save(it['path'])
        self.progress[it['key']] = 'pořadí odstraněno + poznámka'
        self._save_progress()
        self.status.config(text=f"Torzo: pořadí odstraněno + poznámka — {it['comp']}")
        self._advance()

    def _choose_custom(self):
        v = self.custom.get().strip()
        if v:
            self.choose(v)

    def _build_prompt_block(self, it, comment):
        """Hotový prompt do chatu: dotaz + tabulky (výpis z Excelu) + můj koment."""
        L = ['=' * 64,
             f"Sezóna: {it['sid'].replace('_', '/')}",
             f"Soutěž: {it['comp']}  ({it['level']})"]
        if it['kind'] == 'torzo':
            L.append(f"OTÁZKA (TORZO): tabulka má pořadí, ale odehráno jen ~{it['pct']}% "
                     f"jednoho kola. Ponechat pořadí, nebo odstranit (jen seznam týmů)?")
        else:
            pre = f"{it['pos']}. " if it['pos'] not in (None, '') else ''
            L.append(f"OTÁZKA: jaký osud má  {pre}{it['club']}  (teď v Excelu: {it['fate']})?")
        L.append('')
        L.append('Tabulky soutěže (výpis z Excelu):')
        for ph in it['phases']:
            L.append(f"  [{ph['name']}]  ({ph['type']})")
            for k, (sh, row, pos, cid, club, fate) in enumerate(ph['rows'], 1):
                dpos = pos if pos not in (None, '') else k
                mark = '   <-- TENTO' if (sh == it['sheet'] and row == it.get('row')) else ''
                L.append(f"     {str(dpos):>3}. {str(club)[:34]:34} {fate or '—'}{mark}")
        L.append('')
        L.append(f"MŮJ KOMENT: {comment or '(bez komentáře)'}")
        if it['kind'] == 'torzo':
            L.append('Rozhodni: ponechat pořadí / odstranit pořadí (jen seznam) + poznámka.')
        else:
            L.append('Doplň osud: postup / sestup / udržel se / setrval / zůstal,')
            L.append('  nebo štítek: play off / o udržení / o umístění / kvalifikace / prolínací.')
        L.append('')
        return '\n'.join(L)

    def _to_prompt(self):
        it = self.queue[self.idx]
        block = self._build_prompt_block(it, self.custom.get().strip())
        pf = os.path.join(self.data_dir, 'prompty_k_chatu.txt')
        with open(pf, 'a', encoding='utf-8') as f:
            f.write(block + '\n')
        self.progress[it['key']] = CHAT
        self._save_progress()
        self.status.config(text=f"➜ Do promptu pro chat: {it.get('club', it['comp'])}  (soubor prompty_k_chatu.txt)")
        self._advance()

    def skip(self):
        it = self.queue[self.idx]
        self.progress[it['key']] = SKIP
        self._save_progress()
        self.status.config(text=f"Přeskočeno: {it.get('club', it['comp'])}")
        self._advance()

    def _advance(self):
        # lineárně na další položku (NE skok na další nevyřešený) — 8/9/10 popořadě,
        # zpět šipkou kdykoliv; rozhodnutí jednoho týmu nemění ostatní
        if self.idx < len(self.queue) - 1:
            self.idx += 1
        self.show()

    def next_pending(self):
        for off in range(1, len(self.queue) + 1):
            j = (self.idx + off) % len(self.queue)
            st = self.progress.get(self.queue[j]['key'])
            if not st or st == SKIP:
                self.idx = j; self.show(); return
        self.status.config(text='Žádné další nevyřešené.')

    # -- export / zavření --
    def export_zip(self):
        out = filedialog.asksaveasfilename(
            defaultextension='.zip',
            initialfile=f'almanach_opraveno_{datetime.date.today()}.zip',
            filetypes=[('ZIP', '*.zip')])
        if not out:
            return
        done = self._done_count()
        chat = sum(1 for v in self.progress.values() if v == CHAT)
        skipped = sum(1 for v in self.progress.values() if v == SKIP)
        pf = os.path.join(self.data_dir, 'prompty_k_chatu.txt')
        readme = (
            'ALMANACH — opravené sešity\n'
            f'Vytvořeno: {datetime.datetime.now():%Y-%m-%d %H:%M}\n\n'
            'Co je opravené (season_fate):\n'
            '• U víceFázových soutěží už postup/sestup NEvisí na základní části.\n'
            '  Základní část nese štítek, KAM tým šel: play off / o udržení /\n'
            '  o umístění / kvalifikace / prolínací.\n'
            '• Skupina o udržení: poslední = sestup, ostatní = udržel se.\n'
            '• Jednofázové soutěže (přímý postup) beze změny.\n\n'
            f'Ruční dořešení flagů: {done} rozhodnutí, {chat} odloženo do chatu, '
            f'{skipped} přeskočeno.\n'
            + ('• prompty_k_chatu.txt = hotové prompty (dotaz + tabulky + koment) '
               'k dořešení v chatu.\n' if chat else '') + '\n'
            'Slovník osudů: postup, sestup, udržel se, setrval, zůstal\n'
            '  + fázové štítky: play off, o udržení, o umístění, kvalifikace, prolínací.\n')
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
            for p in sorted(glob.glob(os.path.join(self.data_dir, 'S*_FINAL.xlsx'))):
                z.write(p, os.path.join('data', os.path.basename(p)))
            z.writestr('PRECTI_ME.txt', readme)
            if os.path.exists(pf):
                z.write(pf, 'prompty_k_chatu.txt')
        messagebox.showinfo('ZIP', f'Hotovo:\n{out}\n\n{done} rozhodnutí, '
                                   f'{chat} do chatu, {skipped} přeskočeno.')

    def show_help(self):
        messagebox.showinfo('Nápověda — jak na to', (
            'CO TADY DĚLÁŠ:\n'
            'Procházíš situace (flagy), kde si počítač nebyl jistý, a rozhodneš je.\n\n'
            '1) OSUD TÝMU — nahoře vidíš celé tabulky soutěže. U otázky („→ tým …")\n'
            '   klikni na správnou variantu (postup, sestup, udržel se, …).\n'
            '2) TORZO — neúplná tabulka. Vyber „Ponechat pořadí" nebo „Odstranit pořadí".\n'
            '3) NEVÍŠ? Napiš poznámku a dej „➜ Do promptu pro chat", nebo „Přeskočit".\n\n'
            'POHYB: „◀ zpět" a „další ▶" listují. „další nevyřešený" skočí na to,\n'
            'co ještě nemá rozhodnutí. Rozhodnutí jednoho týmu NEMĚNÍ ostatní.\n\n'
            'NÁZVY KLUBŮ: tlačítko „🏷 Opravit názvy klubů" — najdeš klub, opravíš\n'
            'překlep, uložíš (opraví se všude v té sezóně).\n\n'
            'UKLÁDÁNÍ: vše se ukládá průběžně. Můžeš kdykoliv zavřít a vrátit se.\n'
            'NA KONCI: „💾 Hotovo → zabalit do ZIP" a ten zip pošli zpět.'))

    def open_clubs(self):
        ClubEditor(self, self.data_dir)

    def open_continuity(self):
        ContinuityEditor(self, self.data_dir)

    def _on_close(self):
        self._save_progress()
        for wb in self.wbcache.values():
            try:
                wb.close()
            except Exception:
                pass
        self.destroy()


class ClubEditor(tk.Toplevel):
    """Jednoduchá oprava názvů klubů (překlepy) v jedné sezóně.
    Najdeš klub → napíšeš správný název → Ulož: opraví se v CLUBS i ve všech tabulkách."""
    def __init__(self, master, data_dir):
        super().__init__(master)
        self.title('Opravit názvy klubů')
        self.geometry('780x640')
        self.data_dir = data_dir
        self.seasons = season_ids(data_dir)
        self.wb = None
        self.path = None
        self.clubs = []          # [(club_id, name, city, clubs_row, [(sheet,row)])]
        self._build()
        if self.seasons:
            self.season_var.set(self.seasons[0])
            self.load(self.seasons[0])

    def _build(self):
        top = ttk.Frame(self); top.pack(fill='x', padx=8, pady=6)
        ttk.Label(top, text='Sezóna:').pack(side='left')
        self.season_var = tk.StringVar()
        cb = ttk.Combobox(top, textvariable=self.season_var, values=self.seasons,
                          width=12, state='readonly'); cb.pack(side='left', padx=4)
        cb.bind('<<ComboboxSelected>>', lambda e: self.load(self.season_var.get()))
        ttk.Label(top, text='   Hledat:').pack(side='left')
        self.q = ttk.Entry(top, width=26); self.q.pack(side='left', padx=4)
        self.q.bind('<KeyRelease>', lambda e: self.refresh())

        cols = ('klub', 'mesto', 'vyskytu')
        self.lst = ttk.Treeview(self, columns=cols, show='headings', height=20)
        for c, w in zip(cols, (380, 200, 80)):
            self.lst.heading(c, text=c.upper()); self.lst.column(c, width=w, anchor='w')
        self.lst.column('vyskytu', anchor='center')
        self.lst.pack(fill='both', expand=True, padx=8, pady=4)
        self.lst.bind('<<TreeviewSelect>>', self.on_pick)

        ed = ttk.LabelFrame(self, text='Oprava názvu'); ed.pack(fill='x', padx=8, pady=6)
        self.sel = ttk.Label(ed, text='(vyber klub v seznamu)', foreground='#555')
        self.sel.pack(anchor='w', padx=6, pady=2)
        row = ttk.Frame(ed); row.pack(fill='x', padx=6, pady=4)
        ttk.Label(row, text='Správný název:').pack(side='left')
        self.newname = ttk.Entry(row); self.newname.pack(side='left', fill='x', expand=True, padx=4)
        self.newname.bind('<Return>', lambda e: self.save())
        ttk.Button(row, text='Uložit opravu', command=self.save).pack(side='left')
        self.status = ttk.Label(self, text='', relief='sunken', anchor='w')
        self.status.pack(fill='x', side='bottom')

    def load(self, sid):
        self.path = os.path.join(self.data_dir, f'S{sid}_FINAL.xlsx')
        wb = openpyxl.load_workbook(self.path, data_only=True)
        occ = collections.defaultdict(list)
        for sh in wb.sheetnames:
            if sh in NON_DATA:
                continue
            rr = list(wb[sh].iter_rows(values_only=True))
            if not rr:
                continue
            H = {c: j for j, c in enumerate(rr[0]) if c}
            if 'club_id' not in H or 'row_type' not in H:
                continue
            for i, r in enumerate(rr[1:], start=2):
                if r[H['row_type']] in ('T', 'R') and r[H['club_id']]:
                    occ[r[H['club_id']]].append((sh, i))
        cw = list(wb['CLUBS'].iter_rows(values_only=True)); CH = {c: j for j, c in enumerate(cw[0]) if c}
        self.clubs = []
        for i, r in enumerate(cw[1:], start=2):
            cid = r[CH['club_id']] if 'club_id' in CH else None
            if not cid:
                continue
            self.clubs.append((cid,
                               r[CH.get('clean_name', -1)] if 'clean_name' in CH else '',
                               r[CH.get('city', -1)] if 'city' in CH else '',
                               i, occ.get(cid, [])))
        wb.close()
        self.refresh()
        self.status.config(text=f'{len(self.clubs)} klubů v sezóně {sid}')

    def refresh(self):
        q = self.q.get().strip().lower()
        self.lst.delete(*self.lst.get_children())
        self.iid_club = {}
        for c in sorted(self.clubs, key=lambda x: str(x[1])):
            if q and q not in str(c[1]).lower():
                continue
            iid = self.lst.insert('', 'end', values=(c[1], c[2] or '', len(c[4])))
            self.iid_club[iid] = c

    def on_pick(self, _e=None):
        sel = self.lst.selection()
        if not sel:
            return
        c = self.iid_club.get(sel[0])
        if not c:
            return
        self.cur = c
        self.sel.config(text=f'Klub: {c[1]}   ·   {len(c[4])}× v tabulkách   ·   id {c[0]}')
        self.newname.delete(0, 'end'); self.newname.insert(0, c[1] or '')

    def save(self):
        c = getattr(self, 'cur', None)
        new = self.newname.get().strip()
        if not c or not new or new == c[1]:
            return
        if self.wb is None:
            self.wb = openpyxl.load_workbook(self.path)
        wb = self.wb
        # CLUBS.clean_name
        cws = wb['CLUBS']; ch = [x.value for x in cws[1]]
        if 'clean_name' in ch:
            cws.cell(row=c[3], column=ch.index('clean_name') + 1).value = new
        # všechny výskyty v tabulkách: club_name
        for (sh, row) in c[4]:
            head = [x.value for x in wb[sh][1]]
            if 'club_name' in head:
                wb[sh].cell(row=row, column=head.index('club_name') + 1).value = new
        wb.save(self.path)
        self.status.config(text=f'Opraveno: „{c[1]}" → „{new}"  ({len(c[4])} výskytů)')
        # aktualizuj v paměti + seznam
        self.clubs = [(x[0], new, x[2], x[3], x[4]) if x[0] == c[0] else x for x in self.clubs]
        self.refresh()


class ContinuityEditor(tk.Toplevel):
    """Návaznost klubů ze sezóny na sezónu (prev_club_id) — kvůli změnám názvů/nepřesnostem.
    Vidíš, na který klub z předchozí sezóny je tým navázaný, a můžeš to opravit."""
    def __init__(self, master, data_dir):
        super().__init__(master)
        self.title('Návaznost klubů (sezóna → sezóna)')
        self.geometry('940x680')
        self.data_dir = data_dir
        self.seasons = season_ids(data_dir)
        self.wb = None
        self.path = None
        self.clubs = []          # [(cid, name, prev_id, change_note, clubs_row, [(sheet,row)])]
        self.prev_name = {}      # prev season club_id -> name
        self.label2id = {}
        self._build()
        if self.seasons:
            self.season_var.set(self.seasons[0])
            self.load(self.seasons[0])

    def _build(self):
        top = ttk.Frame(self); top.pack(fill='x', padx=8, pady=6)
        ttk.Label(top, text='Sezóna:').pack(side='left')
        self.season_var = tk.StringVar()
        cb = ttk.Combobox(top, textvariable=self.season_var, values=self.seasons,
                          width=12, state='readonly'); cb.pack(side='left', padx=4)
        cb.bind('<<ComboboxSelected>>', lambda e: self.load(self.season_var.get()))
        self.prev_lbl = ttk.Label(top, text='', foreground='#555'); self.prev_lbl.pack(side='left', padx=10)
        ttk.Label(top, text='   Hledat:').pack(side='left')
        self.q = ttk.Entry(top, width=22); self.q.pack(side='left', padx=4)
        self.q.bind('<KeyRelease>', lambda e: self.refresh())
        self.only_missing = tk.BooleanVar(value=False)
        ttk.Checkbutton(top, text='jen nenavázané', variable=self.only_missing,
                        command=self.refresh).pack(side='left', padx=8)

        cols = ('klub', 'navazba')
        self.lst = ttk.Treeview(self, columns=cols, show='headings', height=18)
        self.lst.heading('klub', text='KLUB (tato sezóna)'); self.lst.column('klub', width=380, anchor='w')
        self.lst.heading('navazba', text='← NAVÁZÁN NA (předchozí sezóna)')
        self.lst.column('navazba', width=420, anchor='w')
        self.lst.pack(fill='both', expand=True, padx=8, pady=4)
        self.lst.bind('<<TreeviewSelect>>', self.on_pick)
        self.lst.tag_configure('miss', foreground='#b07d10')

        ed = ttk.LabelFrame(self, text='Oprava návaznosti'); ed.pack(fill='x', padx=8, pady=6)
        self.sel = ttk.Label(ed, text='(vyber klub v seznamu)', foreground='#555')
        self.sel.pack(anchor='w', padx=6, pady=2)
        r1 = ttk.Frame(ed); r1.pack(fill='x', padx=6, pady=3)
        ttk.Label(r1, text='Navázat na klub z předchozí sezóny:').pack(side='left')
        self.prev_var = tk.StringVar()
        self.prev_cb = ttk.Combobox(r1, textvariable=self.prev_var, width=46)
        self.prev_cb.pack(side='left', fill='x', expand=True, padx=4)
        r2 = ttk.Frame(ed); r2.pack(fill='x', padx=6, pady=3)
        ttk.Label(r2, text='Poznámka ke změně:').pack(side='left')
        self.note = ttk.Entry(r2); self.note.pack(side='left', fill='x', expand=True, padx=4)
        ttk.Button(r2, text='Uložit návaznost', command=self.save).pack(side='left')
        self.status = ttk.Label(self, text='', relief='sunken', anchor='w'); self.status.pack(fill='x', side='bottom')

    def load(self, sid):
        self.path = os.path.join(self.data_dir, f'S{sid}_FINAL.xlsx')
        wb = openpyxl.load_workbook(self.path, data_only=True)
        meta = {}
        if 'META' in wb.sheetnames:
            for r in wb['META'].iter_rows(values_only=True):
                if r and r[0]:
                    meta[str(r[0])] = r[1]
        prev_sid = str(meta.get('prev_season') or '').replace('S', '')
        if not prev_sid and sid in self.seasons:
            i = self.seasons.index(sid)
            prev_sid = self.seasons[i - 1] if i > 0 else ''
        # výskyty club_id v tabulkách (kvůli prev_club_id v řádcích)
        occ = collections.defaultdict(list)
        for sh in wb.sheetnames:
            if sh in NON_DATA:
                continue
            rr = list(wb[sh].iter_rows(values_only=True))
            if not rr:
                continue
            H = {c: j for j, c in enumerate(rr[0]) if c}
            if 'club_id' not in H or 'row_type' not in H:
                continue
            for i, r in enumerate(rr[1:], start=2):
                if r[H['row_type']] in ('T', 'R') and r[H['club_id']]:
                    occ[r[H['club_id']]].append((sh, i))
        cw = list(wb['CLUBS'].iter_rows(values_only=True)); CH = {c: j for j, c in enumerate(cw[0]) if c}
        self.clubs = []
        for i, r in enumerate(cw[1:], start=2):
            cid = r[CH['club_id']] if 'club_id' in CH else None
            if not cid:
                continue
            self.clubs.append((cid, r[CH.get('clean_name', -1)] if 'clean_name' in CH else '',
                               r[CH.get('prev_club_id', -1)] if 'prev_club_id' in CH else None,
                               r[CH.get('change_note', -1)] if 'change_note' in CH else '',
                               i, occ.get(cid, [])))
        wb.close()
        # předchozí sezóna: jména klubů
        self.prev_name = {}
        ppath = os.path.join(self.data_dir, f'S{prev_sid}_FINAL.xlsx')
        if prev_sid and os.path.exists(ppath):
            pw = openpyxl.load_workbook(ppath, data_only=True)
            pcw = list(pw['CLUBS'].iter_rows(values_only=True)); PH = {c: j for j, c in enumerate(pcw[0]) if c}
            for r in pcw[1:]:
                pid = r[PH['club_id']] if 'club_id' in PH else None
                if pid:
                    self.prev_name[pid] = r[PH.get('clean_name', -1)] if 'clean_name' in PH else ''
            pw.close()
        self.prev_lbl.config(text=('← předchozí: ' + prev_sid.replace('_', '/')) if prev_sid
                             else '(žádná předchozí sezóna)')
        opts = ['(žádný / nový klub)'] + [f'{n}  [{i}]' for i, n in
                                          sorted(self.prev_name.items(), key=lambda x: str(x[1]))]
        self.label2id = {f'{n}  [{i}]': i for i, n in self.prev_name.items()}
        self.prev_cb['values'] = opts
        self.refresh()
        miss = sum(1 for c in self.clubs if not c[2])
        self.status.config(text=f'{len(self.clubs)} klubů · {len(self.clubs)-miss} navázáno · {miss} nenavázáno')

    def refresh(self):
        q = self.q.get().strip().lower()
        self.lst.delete(*self.lst.get_children())
        self.iid_club = {}
        for c in sorted(self.clubs, key=lambda x: str(x[1])):
            if q and q not in str(c[1]).lower():
                continue
            if self.only_missing.get() and c[2]:
                continue
            pname = self.prev_name.get(c[2]) if c[2] else None
            nav = pname if pname else ('—  (nenavázáno / nový klub)' if not c[2]
                                       else f'?  (id {c[2]} není v předchozí sezóně)')
            iid = self.lst.insert('', 'end', values=(c[1], nav),
                                  tags=('miss',) if (not pname) else ())
            self.iid_club[iid] = c

    def on_pick(self, _e=None):
        sel = self.lst.selection()
        if not sel:
            return
        c = self.iid_club.get(sel[0])
        if not c:
            return
        self.cur = c
        self.sel.config(text=f'Klub: {c[1]}   ·   id {c[0]}   ·   {len(c[5])}× v tabulkách')
        pname = self.prev_name.get(c[2]) if c[2] else None
        self.prev_var.set(f'{pname}  [{c[2]}]' if pname else '(žádný / nový klub)')
        self.note.delete(0, 'end'); self.note.insert(0, c[3] or '')

    def save(self):
        c = getattr(self, 'cur', None)
        if not c:
            return
        label = self.prev_var.get().strip()
        new_prev = self.label2id.get(label)
        if new_prev is None:
            m = re.search(r'\[([^\]]+)\]\s*$', label)
            new_prev = m.group(1) if (m and m.group(1) in self.prev_name) else None
        note = self.note.get().strip()
        if self.wb is None:
            self.wb = openpyxl.load_workbook(self.path)
        wb = self.wb
        cws = wb['CLUBS']; ch = [x.value for x in cws[1]]
        if 'prev_club_id' in ch:
            cws.cell(row=c[4], column=ch.index('prev_club_id') + 1).value = new_prev
        if 'change_note' in ch and note:
            cws.cell(row=c[4], column=ch.index('change_note') + 1).value = note
        for (sh, row) in c[5]:                    # sjednoť i prev_club_id v tabulkách
            head = [x.value for x in wb[sh][1]]
            if 'prev_club_id' in head:
                wb[sh].cell(row=row, column=head.index('prev_club_id') + 1).value = new_prev
        wb.save(self.path)
        self.clubs = [(x[0], x[1], new_prev, note or x[3], x[4], x[5]) if x[0] == c[0] else x
                      for x in self.clubs]
        nm = self.prev_name.get(new_prev, '(nový / žádný)')
        self.status.config(text=f'Návaznost uložena: {c[1]} ← {nm}')
        self.refresh()


def extract_payload():
    """Při běhu jako .exe (PyInstaller) rozbal přibalené složky vedle .exe (jen poprvé)."""
    if not getattr(sys, 'frozen', False):
        return
    payload = os.path.join(getattr(sys, '_MEIPASS', ''), '_payload')
    if not os.path.isdir(payload):
        return
    ws = _app_workspace()
    try:
        os.makedirs(ws, exist_ok=True)
    except Exception:
        return
    for name in os.listdir(payload):
        src = os.path.join(payload, name)
        dst = os.path.join(ws, name)
        if os.path.isdir(src) and not os.path.exists(dst):
            try:
                shutil.copytree(src, dst)
            except Exception:
                pass


def main():
    extract_payload()
    root = tk.Tk(); root.withdraw()
    data_dir = find_data_dir()
    if not data_dir:
        data_dir = filedialog.askdirectory(title='Vyber složku se sešity S*_FINAL.xlsx')
    if not data_dir or not glob.glob(os.path.join(data_dir, 'S*_FINAL.xlsx')):
        print('Nenašel jsem žádné S*_FINAL.xlsx.'); return
    app = FlagSolver(root, data_dir)
    app.protocol('WM_DELETE_WINDOW', lambda: (app._on_close(), root.destroy()))
    root.mainloop()


if __name__ == '__main__':
    main()
