#!/usr/bin/env python3
"""
build_1948_49_zupni.py — doplní 15 župních soutěží 1948/49 do S1948_49_FINAL.xlsx.

Zdroj: surový XLS od kolegy (jen ZÁKLAD ÚPLNOSTI — víme, co existovalo a kdo
v tom hrál; hodnoty NEJSOU finální pravda, ladí se postupně z archivů/tisku).
Soubor se neukládá do repa. Spuštění:
    python3 build_1948_49_zupni.py <raw.xls> [--write]
Bez --write jen vypíše souhrn parsování.

Mapování úrovní: I.třída / krajský přebor / župní mistrovství = L30,
II.třída = L40, III.třída = L50 (pod divizí L20).
"""
import sys, re, shutil, datetime
import xlrd
import openpyxl

TARGET = 'data/S1948_49_FINAL.xlsx'

# (raw sheet, sheet code, display name, region)
ZUPY = [
    ('Středočeská',   'ZSTRC', 'Středočeská župa',   'REG_STC'),
    ('Kladenská',     'ZKLAD', 'Kladenská župa',     'REG_STC'),
    ('Středolabská',  'ZSLAB', 'Středolabská župa',  'REG_STC'),
    ('Šumavská',      'ZSUMA', 'Šumavská župa',      'REG_JHC'),
    ('Podbrdská',     'ZPODB', 'Podbrdská župa',     'REG_STC'),
    ('Krušnohorská',  'ZKRUS', 'Krušnohorská župa',  'REG_SVC'),
    ('Západočeská',   'ZZAPC', 'Západočeská župa',   'REG_ZPC'),
    ('Severočeská',   'ZSEVC', 'Severočeská župa',   'REG_SVC'),
    ('Východočeská',  'ZVYCH', 'Východočeská župa',  'REG_VYC'),
    ('Orlická',       'ZORLI', 'Orlická župa',       'REG_VYC'),
    ('Horácká',       'ZHORA', 'Horácká župa',       'REG_VYS'),
    ('Západomoravská','ZZAPM', 'Západomoravská župa','REG_SVM'),
    ('Slovácká',      'ZSLOV', 'Slovácká župa',      'REG_SVM'),
    ('Hanácká',       'ZHANA', 'Hanácká župa',       'REG_SVM'),
    ('Slezská',       'ZSLEZ', 'Slezská župa',       'REG_SVM'),
]

STD_HEADER = ['row_type','block_name','pos','club_name','note','GP','W','D','L',
              'GF',':','GA','PTS','comp_path','node_id','level','club_id',
              'prev_club_id','dest_node_id','dest_type','season_fate','tr_id','district']


def s(v):
    if v is None:
        return ''
    if isinstance(v, float):
        # integer-valued float → int text; but for header detection we keep str
        if v == int(v):
            return str(int(v))
        return str(v)
    return str(v).strip()


def num(v):
    """Numerická hodnota buňky nebo None."""
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v) == int(v) else v
    t = str(v).strip().replace(',', '.')
    if re.fullmatch(r'-?\d+(\.0+)?', t):
        return int(float(t))
    return None


def trida_level(c0):
    """Vrátí (level, label) pro řádek třídy, nebo None když to není třída."""
    low = c0.lower()
    if 'oblastní soutěž' in low or low.strip() in ('oblastní soutěž', 'oblastní'):
        return ('L20', c0)
    if not ('třída' in low or 'přebor' in low or 'mistrovství' in low):
        return None
    nodots = low.replace('.', '').replace(' ', '')
    m = re.match(r'(iii|ii|i)', nodots)
    if 'iii' in nodots[:5]:
        lvl = 'L50'
    elif re.match(r'(iii|ii)', nodots):
        lvl = 'L40' if nodots.startswith('ii') and not nodots.startswith('iii') else ('L50' if nodots.startswith('iii') else 'L30')
    else:
        lvl = 'L30'
    # robustnější: rozhodni podle počtu I na začátku
    mm = re.match(r'(i+)', nodots)
    if mm:
        n = len(mm.group(1))
        lvl = {1: 'L30', 2: 'L40', 3: 'L50'}.get(n, 'L30')
    elif 'přebor' in low or 'mistrovství' in low:
        lvl = 'L30'
    return (lvl, c0)


def norm_trida_label(c0):
    """Sjednoť název třídy a vrať (label, extra_text)."""
    low = c0.lower()
    if 'oblastní soutěž' in low:
        return ('Oblastní soutěž', '')
    if 'přebor' in low:
        return ('Krajský přebor', '')
    if 'mistrovství' in low:
        return ('Župní mistrovství', '')
    m = re.match(r'^\s*([IVX]+)\.?\s*(župní\s*)?třída\s*([AB])?\s*', c0, re.I)
    if m:
        roman = m.group(1).upper()
        sub = f" {m.group(3).upper()}" if m.group(3) else ''
        label = (f"{roman}. župní třída{sub}" if m.group(2)
                 else f"{roman}.třída{sub}")
        extra = c0[m.end():].strip()
        return (label, extra)
    return (c0.strip(), '')


def is_group(c0):
    low = c0.lower()
    return low.startswith('skupin') or low.startswith('okrsek')


def parse_zupa(sh):
    """Vrátí seznam tříd: [{label, level, groups:[{label, teams:[...]}], notes:[...]}].
    teams: dict(rank,name,raw,note,GP,W,D,L,GF,GA,PTS,status)."""
    tridy = []
    cur_trida = None
    cur_group = None          # dict or None (=implicitní skupina na úrovni třídy)
    mode = 'std'              # 'std' | 'result' (po finále/kvalifikaci)
    result_label = None

    def ensure_implicit_group():
        nonlocal cur_group
        if cur_trida is None:
            return None
        if cur_group is None:
            # implicitní skupina = sama třída
            g = cur_trida['_implicit']
            return g
        return cur_group

    for r in range(min(sh.nrows, 250)):
        cells = [sh.cell_value(r, c) if c < sh.ncols else '' for c in range(10)]
        c0 = s(cells[0])
        name = s(cells[1])
        # úplně prázdný řádek (sloupce 0-9)
        if not any(s(x) for x in cells):
            continue

        if name:  # ── TÝMOVÝ ŘÁDEK ──
            # ochrana: jméno obsahující skóre (N:N) je ve skutečnosti výsledek zápasu,
            # ne tým (zdroj místo tabulky uvedl jen pár výsledků) → poznámka
            if re.search(r'\d+\s*:\s*\d+', name) and num(cells[2]) is None:
                if cur_trida is not None:
                    cur_trida['notes'].append(f"výsledek: {name}")
                continue
            if cur_trida is None:
                # tým bez třídy → založ implicitní I.třídu
                cur_trida = new_trida('I.třída', 'L30')
                tridy.append(cur_trida)
            g = ensure_implicit_group()
            rank = None
            mrank = re.match(r'(\d+)', c0)
            if mrank:
                rank = int(mrank.group(1))
            # status text místo statistik?
            col2 = cells[2]
            status = None
            GP = W = D = L = GF = GA = PTS = None
            if num(col2) is None and s(col2):
                status = s(col2)
            else:
                GP, W, D, L = num(cells[2]), num(cells[3]), num(cells[4]), num(cells[5])
                GF, GA, PTS = num(cells[6]), num(cells[8]), num(cells[9])
            # note kód z koncovky (N)/(S)/(M)
            note = None
            mnote = re.search(r'\((N|S|M)\)\s*$', name)
            clean = name
            if mnote:
                note = mnote.group(1)
                clean = name[:mnote.start()].strip()
            g['teams'].append(dict(rank=rank, name=clean, raw=name, note=note,
                                   GP=GP, W=W, D=D, L=L, GF=GF, GA=GA, PTS=PTS,
                                   status=status))
            continue

        # ── ŘÁDEK SE SLOUPCEM 0 (hlavička / poznámka / výsledek) ──
        if re.fullmatch(r'[\d]+\.?', c0) or c0 in ('?',):
            continue  # osamocené pořadí bez jména

        low = c0.lower()
        tl = trida_level(c0)
        if tl:
            label, extra = norm_trida_label(c0)
            cur_trida = new_trida(label, tl[0])
            tridy.append(cur_trida)
            cur_group = None
            mode = 'std'
            if extra and len(extra) > 2:
                cur_trida['notes'].append(extra)
            continue
        if is_group(c0):
            cur_group = dict(label=c0, teams=[])
            cur_trida['groups'].append(cur_group)
            mode = 'std'
            continue
        if low.startswith('finále'):
            mode = 'result'; result_label = 'finále'
            continue
        if 'kvalifikace' in low:
            mode = 'result'; result_label = 'kvalifikace'
            continue
        # jinak: poznámka nebo výsledkový řádek
        if cur_trida is not None:
            if mode == 'result':
                cur_trida['notes'].append(f"{result_label}: {c0}")
            else:
                cur_trida['notes'].append(c0)
    return tridy


def new_trida(label, level):
    impl = dict(label=None, teams=[])   # implicitní skupina = třída samotná
    return dict(label=label, level=level, groups=[], notes=[], _implicit=impl)


def main():
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    raw_path = sys.argv[1]
    do_write = '--write' in sys.argv

    rb = xlrd.open_workbook(raw_path)
    parsed = {}
    for raw_name, code, disp, region in ZUPY:
        if raw_name not in rb.sheet_names():
            print(f"!! chybí list {raw_name}"); continue
        parsed[code] = (disp, region, parse_zupa(rb.sheet_by_name(raw_name)))

    # ── souhrn ──
    tot_teams = tot_nodes = tot_notes = 0
    for code in [c for _, c, _, _ in ZUPY if c in parsed]:
        disp, region, tridy = parsed[code]
        zt = zn = znotes = 0
        parts = []
        for t in tridy:
            groups = [t['_implicit']] + t['groups'] if t['_implicit']['teams'] else t['groups']
            real_groups = [g for g in groups if g['teams']]
            nteams = sum(len(g['teams']) for g in real_groups)
            zt += nteams
            zn += 1 + len([g for g in t['groups'] if g['teams']])  # container + named groups
            znotes += len(t['notes'])
            parts.append(f"{t['label']}[{t['level']}]:{nteams}t/{len(real_groups)}g")
        tot_teams += zt; tot_nodes += zn; tot_notes += znotes
        print(f"{code} {disp:20s} týmů={zt:3d} nodů~{zn:2d} pozn={znotes:2d} | " + "  ".join(parts))
    print(f"\nCELKEM: týmů={tot_teams} nodů~{tot_nodes} poznámek={tot_notes}")

    if not do_write:
        print("\n(dry-run — pro zápis přidej --write)")
        return

    write_into_target(parsed)


def write_into_target(parsed):
    shutil.copy(TARGET, TARGET + '.bak')
    wb = openpyxl.load_workbook(TARGET)

    # zjisti max id
    def maxid(prefix):
        mx = 0
        for sn in wb.sheetnames:
            for row in wb[sn].iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        for g in re.findall(prefix + r'(\d+)', v):
                            mx = max(mx, int(g))
        return mx
    node_n = maxid('NODE_S1948_49_')
    club_n = maxid('CLUB_S1948_49_')
    tr_n = maxid('TR_S1948_49_')

    clubs_ws = wb['CLUBS']
    notes_ws = wb['NOTES']
    sys_ws = wb['SYSTEM']

    def add_node(name, ctype, level, region, parent, entry=None, phase=None, note=None):
        nonlocal node_n
        node_n += 1
        nid = f"NODE_S1948_49_{node_n:04d}"
        sys_ws.append([nid, name, ctype, level, region, parent, None, None,
                       '2-1-0', None, note, None, entry, phase])
        return nid

    def add_club(clean, raw, sheet, note, level):
        nonlocal club_n
        club_n += 1
        cid = f"CLUB_S1948_49_{club_n:04d}"
        clubs_ws.append([cid, clean, raw, sheet, note, level, None, None,
                         'Župní soutěž 1948/49 (surový podklad kolegy – základ úplnosti, hodnoty k ověření).', None])
        return cid

    def add_tr():
        nonlocal tr_n
        tr_n += 1
        return f"TR_S1948_49_{tr_n:05d}"

    new_notes = []
    for raw_name, code, disp, region in ZUPY:
        if code not in parsed:
            continue
        disp, region, tridy = parsed[code]
        ws = wb.create_sheet(code)
        ws.append(STD_HEADER)
        used_labels = {}
        for t in tridy:
            # poznámky třídy → NOTES
            for nt in t['notes']:
                new_notes.append(f"{disp} – {t['label']}: {nt}")
            # prázdná třída (jen hlavička, žádné týmy) → jen poznámka, bez nodu
            n_all = len(t['_implicit']['teams']) + sum(len(g['teams']) for g in t['groups'])
            if n_all == 0:
                new_notes.append(f"{disp} – {t['label']}: v podkladu bez tabulek")
                continue
            # kontejner třídy
            tlabel = t['label']
            if tlabel in used_labels:
                used_labels[tlabel] += 1
                tlabel = f"{tlabel} ({used_labels[tlabel]})"
            else:
                used_labels[tlabel] = 1
            cont_name = f"{disp} {tlabel}"
            cont_id = add_node(cont_name, 'league', t['level'], region, None,
                               note='Župní soutěž 1948/49 (podklad kolegy).')
            # skupiny: implicitní (teams přímo) + pojmenované
            groups = []
            if t['_implicit']['teams']:
                groups.append(('', t['_implicit'], cont_id, cont_name))
            for g in t['groups']:
                if not g['teams']:
                    new_notes.append(f"{disp} – {tlabel} / {g['label']}: v podkladu bez týmů")
                    continue
                gnode = add_node(f"{cont_name} / {g['label']}", 'group', t['level'],
                                 region, cont_id)
                groups.append((g['label'], g, gnode, f"{cont_name} / {g['label']}"))
            for glabel, g, gnode, gpath in groups:
                ws.append(['H', gpath if glabel else cont_name, None, None, None,
                           None, None, None, None, None, None, None, None,
                           None, gnode, t['level'], None, None, None, None, None,
                           None, None])
                for tm in g['teams']:
                    cid = add_club(tm['name'], tm['raw'], code, tm['note'], t['level'])
                    fate = tm['status'] if tm['status'] else None
                    ws.append(['T', None, tm['rank'], tm['name'], tm['note'],
                               tm['GP'], tm['W'], tm['D'], tm['L'], tm['GF'], ':',
                               tm['GA'], tm['PTS'], gpath if glabel else cont_name,
                               gnode, t['level'], cid, None, None, None, fate,
                               add_tr(), None])

    for nt in new_notes:
        notes_ws.append([nt])

    # META update
    meta_ws = wb['META']
    metavals = {meta_ws.cell(row=i+1, column=1).value: i+1
                for i in range(meta_ws.max_row)}
    total_clubs = club_n
    # přepočti nody
    total_nodes = sum(1 for row in sys_ws.iter_rows(min_row=2, values_only=True) if row[0])
    if 'total_clubs' in metavals:
        meta_ws.cell(row=metavals['total_clubs'], column=2, value=total_clubs)
    if 'total_nodes' in metavals:
        meta_ws.cell(row=metavals['total_nodes'], column=2, value=total_nodes)
    if 'status' in metavals:
        meta_ws.cell(row=metavals['status'], column=2,
                     value=('Nejvyšší soutěž + kompletní divize (4 kraje + Slovensko) + '
                            'celostátní play-off divize + 15 župních soutěží (I.–III. třída) '
                            'z podkladu kolegy jako základ úplnosti (hodnoty k ruční verifikaci).'))
    meta_ws.append(['zupni_note',
                    f'Doplněno 15 župních soutěží ({datetime.date.today().isoformat()}) ze '
                    f'surového XLS kolegy – jen základ úplnosti, mnohde neúplné, '
                    f'hodnoty se ladí z archivů/dobového tisku.'])

    wb.save(TARGET)
    print(f"\n✓ zapsáno do {TARGET} (záloha {TARGET}.bak)")
    print(f"  nové listy: {', '.join(c for _,c,_,_ in ZUPY if c in parsed)}")
    print(f"  celkem klubů v sezoně: {total_clubs}, nodů: {total_nodes}")


if __name__ == '__main__':
    main()
