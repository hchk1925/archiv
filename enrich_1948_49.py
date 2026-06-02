#!/usr/bin/env python3
"""
enrich_1948_49.py — srovnání 1948/49 s PDF (in-place, zachová TODO/fáze).

A) Nejvyšší soutěž (Státní liga): čísla už sedí s PDF; doplní oficiální sokolské
   názvy do raw_name (clean_name = čitelný klub zůstává) + věrné poznámky z PDF.
B) Středočeská divize (Skupina A + B) — první nižší region, ověřeno (GF=GA,
   body=2V+R): nový list 30STRC + SYSTEM uzly + CLUBS.
Idempotentní.
"""
import openpyxl, re
from openpyxl.utils import get_column_letter

PATH = 'data/S1948_49_FINAL.xlsx'
S = 'S1948_49'

# clean_name → oficiální název z PDF (do raw_name)
RAW = {
    'LTC Praha': 'JTO Sokol Pražský (LTC Praha)',
    'ŠK Bratislava': 'ZJ Sokol NV Bratislava',
    'ATK Praha': 'ATK Praha',
    'AC Stadion České Budějovice': 'JTO Sokol Stadion České Budějovice',
    'Sokol Zbrojovka Židenice': 'ZSJ Zbrojovka Brno-Židenice',
    'AC Sparta Bubeneč': 'JTO Sokol Sparta Bubeneč',
    'Sokol Prostějov': 'JTO Sokol II Prostějov',
    'I. ČLTK Praha': 'JTO Sokol I Praha (I. ČLTK Praha)',
}

NOTES = [
    'Nejvyšší soutěž (Státní liga) i divize zpracovány z PDF (dobový tisk).',
    '31. 3. 1948 sjednocení všech tělovýchovných organizací pod JTO Sokol — odtud sokolské názvy klubů (raw_name).',
    '3. 11. 1948: sloučení JTO Sokol Horácká Slavia Třebíč se ZSJ Zbrojovka Brno-Židenice; nový klub pod hlavičkou Zbrojovky převzal hráče i práva na Státní ligu.',
    '9. 11. 1948: po odchodu všech hráčů ze Stadionu Podolí (do Sparty a I. ČLTK) klub vyřazen ze soutěže — JTO Sokol Stadion Podolí ukončil činnost.',
    'Pražské oddíly (LTC, ATK, Sparta, I. ČLTK) hrály doma na ZS Štvanice; Prostějov v Olomouci; Židenice na ZS Za Lužánkami.',
]

# Středočeská divize — (pos, raw_name, GP,W,D,L,GF,GA,PTS, city)
STRC_A = [
    (1, 'JTO Sokol Černošice', 6, 5, 1, 0, 39, 25, 11, 'Černošice'),
    (2, 'JTO Sokol Záběhlice-Zbraslav', 6, 4, 0, 2, 27, 19, 8, 'Praha'),
    (3, 'JTO Sokol Slavia Praha', 6, 3, 0, 3, 26, 18, 6, 'Praha'),
    (4, 'JTO Sokol Libeň', 6, 3, 0, 3, 25, 28, 6, 'Praha'),
    (5, 'JTO Sokol Roudnice nad Labem', 6, 3, 0, 3, 18, 23, 6, 'Roudnice nad Labem'),
    (6, 'JTO Sokol Kročehlavy', 6, 2, 1, 3, 24, 18, 5, 'Kladno'),
    (7, 'JTO Sokol II Smíchov', 6, 0, 0, 6, 15, 43, 0, 'Praha'),
]
STRC_B = [
    (1, 'JTO Sokol Hostivař', 6, 6, 0, 0, 67, 22, 12, 'Praha'),
    (2, 'ZSJ Kara Starý Kolín', 6, 5, 0, 1, 53, 28, 10, 'Starý Kolín'),
    (3, 'JTO Sokol Říčany', 6, 4, 0, 2, 37, 41, 8, 'Říčany'),
    (4, 'JTO Sokol Velké Popovice', 6, 3, 0, 3, 37, 38, 6, 'Velké Popovice'),
    (5, 'JTO Sokol Modrobílí Kolín', 6, 2, 0, 4, 40, 39, 4, 'Kolín'),
    (6, 'JTO Sokol Vršovice-Bohemians', 6, 1, 0, 5, 33, 43, 2, 'Praha'),
    (7, 'JTO Sokol Čelákovice', 6, 0, 0, 6, 22, 78, 0, 'Čelákovice'),
]

HDR10 = ['row_type', 'block_name', 'pos', 'club_name', 'note', 'GP', 'W', 'D', 'L',
         'GF', ':', 'GA', 'PTS', 'comp_path', 'node_id', 'level', 'club_id',
         'prev_club_id', 'dest_node_id', 'dest_type', 'season_fate', 'tr_id', 'district']


def clean(raw):
    return re.sub(r'^JTO ', '', raw)


def top_league(wb):
    ws = wb['10_liga']
    H = {c.value: i for i, c in enumerate(ws[1])}
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, H['club_name'] + 1).value
        if nm in RAW:
            # raw_name v 10_liga není sloupec; ulož do CLUBS. Tady jen ponech.
            pass
    cl = wb['CLUBS']
    Hc = {c.value: i for i, c in enumerate(cl[1])}
    for r in range(2, cl.max_row + 1):
        nm = cl.cell(r, Hc['clean_name'] + 1).value
        if nm in RAW:
            cl.cell(r, Hc['raw_name'] + 1).value = RAW[nm]
    # NOTES (přepiš na věrné PDF)
    no = wb['NOTES']
    no.delete_rows(2, no.max_row)
    for t in NOTES:
        no.append([t])
    # META
    me = wb['META']
    mk = {me.cell(r, 1).value: r for r in range(2, me.max_row + 1)}
    def setm(k, v):
        if k in mk:
            me.cell(mk[k], 2).value = v
        else:
            me.append([k, v])
    setm('source', 'PDF sources/CZE1 + CZE2_3_nizsi (dobový tisk)')
    setm('note', 'Šestý ročník Státní ligy, 8 týmů jednokolově (mistr LTC Praha). '
                 'Sokolské názvy (31.3.1948 sjednocení pod JTO Sokol). Stadion Podolí '
                 'vyřazen 9.11.1948. Divize zpracovány po krajích.')


def add_strc(wb):
    if '30STRC' in wb.sheetnames:
        return
    ws = wb.create_sheet('30STRC')
    ws.append(HDR10)
    NODE_DIV = f'NODE_{S}_0004'
    NODE_A = f'NODE_{S}_0005'
    NODE_B = f'NODE_{S}_0006'
    cl = wb['CLUBS']
    sy = wb['SYSTEM']
    cidx = 13
    tr = 101

    def block(bname, node, teams):
        nonlocal cidx, tr
        ws.append(['H', bname, '', '', '', '', '', '', '', '', '', '', '', '',
                   node, 'L20', '', '', '', '', '', '', ''])
        for (pos, raw, gp, w, d, l, gf, ga, pts, city) in teams:
            cid = f'CLUB_{S}_{cidx:04d}'
            ws.append(['T', '', pos, clean(raw), '', gp, w, d, l, gf, ':', ga, pts,
                       bname, node, 'L20', cid, '', '', '', '', f'TR_{S}_{tr:05d}', ''])
            cl.append([cid, clean(raw), raw, '30STRC', '', 'L20', '', '',
                       'Středočeská divize (PDF). prev_club_id zatím prázdný.', city])
            cidx += 1
            tr += 1

    block('Středočeská divize / Skupina A', NODE_A, STRC_A)
    block('Středočeská divize / Skupina B', NODE_B, STRC_B)
    sy.append([NODE_DIV, '30_Středočeská divize', 'league', 'L20', 'REG_STC', '', '', '', '2-1-0', '', '', '', '', ''])
    sy.append([NODE_A, '30_Středočeská divize Skupina A', 'group', 'L20', 'REG_STC', NODE_DIV, '', '', '2-1-0', '', '', '', '', ''])
    sy.append([NODE_B, '30_Středočeská divize Skupina B', 'group', 'L20', 'REG_STC', NODE_DIV, '', '', '2-1-0', '', '', '', '', ''])
    for j, wdt in enumerate([6, 30, 5, 26, 5, 4, 4, 4, 4, 5, 3, 5, 5, 30, 20, 6, 20, 20, 18, 12, 12, 18, 12], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    print("  ✔ 30STRC (Středočeská divize A+B) + uzly + 14 klubů")


def main():
    wb = openpyxl.load_workbook(PATH)
    top_league(wb)
    print("  ✔ Státní liga: raw_name (sokolské názvy) + poznámky + META")
    add_strc(wb)
    wb.save(PATH)
    print(f"=== {PATH} srovnáno s PDF ===")


if __name__ == '__main__':
    main()
