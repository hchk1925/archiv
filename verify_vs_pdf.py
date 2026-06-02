#!/usr/bin/env python3
"""
verify_vs_pdf.py — křížová kontrola NÁRODNÍCH soutěží (nejvyšší + 1./2. liga)
proti originálním PDF. Pořadí pravdy: PDF > xlsx.

Zdroje PDF:
  sources/CZE1/CZ-YYYY-YY.pdf        — nejvyšší soutěž
  sources/CZE2_3_nizsi/CZ-YYYY-YY-DS.pdf — 1. liga, 2. liga, kvalifikace

Kontroluje listy s národním tierem (název `^\\d0_`: 10_liga, 20_*, 30_2liga…);
krajské listy (30PRAH, 30STRC…) vynechává — ty PDF DS nepokrývají systematicky.

Join přes (GP, GF, GA) v poolu řádků obou PDF; 2. průchod přes ověřené V/R/P.
Hlásí a (s --write) opravuje:
  - V/R/P proti PDF (éra <2000, kde platí 2-1-0)
  - PTS=2·V+R v ne-přenosových blocích (<2000)
  - drobné překlepy GF:GA (když V/R/P jednoznačně sedí)
Řádky bez jednoznačné shody nechává (hlásí „k ověření") — nehádá.
"""
import openpyxl, glob, os, re, sys
from pdfminer.high_level import extract_text

WRITE = '--write' in sys.argv
CZE1, CZE23 = 'sources/CZE1', 'sources/CZE2_3_nizsi'
NAT_SHEET = re.compile(r'^\d0_')          # národní tier (podtržítko za číslem)
REGIONAL = re.compile(r'oblast|Plzeň|kraj', re.I)   # vyloučit krajské listy
ROW_FLEX = re.compile(r'(?<!\d)(\d{1,2})\s+((?:\d{1,3}\s+){2,6}?)(\d{1,3}):(\d{1,3})(?!\d)')
CARRY = re.compile(r'udržen|umíst|prolín|baráž|nadstavb|playoff|play-off', re.I)


def derive_wdl(outs):
    n = len(outs)
    if n == 3:                      # V R P
        return outs[0], outs[1], outs[2]
    if n == 4:                      # V VP PP P (bez remíz)
        return outs[0] + outs[1], 0, outs[2] + outs[3]
    if n == 5:                      # V VP R PP P
        return outs[0] + outs[1], outs[2], outs[3] + outs[4]
    if n == 6:                      # V VP VN PN PP P (nájezdy, bez remíz)
        return outs[0] + outs[1] + outs[2], 0, outs[3] + outs[4] + outs[5]
    return None


def _parse_one(path, rows):
    if not os.path.exists(path):
        return
    for line in extract_text(path).splitlines():
        m = ROW_FLEX.search(line)             # max 1 řádek tabulky na textový řádek
        if not m:
            continue
        outs = [int(x) for x in m.group(2).split()]
        rows.append({'gp': int(m.group(1)), 'outs': outs, 'wdl': derive_wdl(outs),
                     'gf': int(m.group(3)), 'ga': int(m.group(4)), 'used': False})


def parse_pdf(sid):
    yr = sid[1:].replace('_', '-')
    rows = []
    _parse_one(os.path.join(CZE1, f'CZ-{yr}.pdf'), rows)
    _parse_one(os.path.join(CZE23, f'CZ-{yr}-DS.pdf'), rows)
    return rows or None


def hcol(ws, name):
    for j, c in enumerate(next(ws.iter_rows(max_row=1)), 1):
        if c.value == name:
            return j
    return None


def process_sheet(ws, pdf, yr):
    cRT, cGP = hcol(ws, 'row_type'), hcol(ws, 'GP')
    cW, cD, cL = hcol(ws, 'W'), hcol(ws, 'D'), hcol(ws, 'L')
    cGF, cGA, cPTS = hcol(ws, 'GF'), hcol(ws, 'GA'), hcol(ws, 'PTS')
    if None in (cRT, cGP, cW, cD, cL, cGF, cGA, cPTS):
        return [], [], [], 0
    wdldiff, gafix, fixed = [], [], 0
    unmatched, blk = [], ''
    for row in ws.iter_rows():
        if row[cRT - 1].value == 'H' and row[1].value:
            blk = str(row[1].value)
        if row[cRT - 1].value != 'T':
            continue
        gp, gf, ga, nm = (row[cGP-1].value, row[cGF-1].value,
                          row[cGA-1].value, row[3].value)
        if not all(isinstance(x, int) for x in (gp, gf, ga)):
            continue
        cand = next((r for r in pdf if not r['used'] and r['gp'] == gp
                     and r['gf'] == gf and r['ga'] == ga), None)
        if cand is None:
            unmatched.append((row, nm, gp))
            continue
        cand['used'] = True
        if yr < 2000 and cand['wdl'] is not None:
            w, d, l = cand['wdl']
            xw, xd, xl = row[cW-1].value, row[cD-1].value, row[cL-1].value
            if (xw, xd, xl) != (w, d, l) and w + d + l == gp:
                wdldiff.append((nm, (xw, xd, xl), (w, d, l)))
                if WRITE:
                    row[cW-1].value, row[cD-1].value, row[cL-1].value = w, d, l
                    row[cPTS-1].value = 2 * w + d
                    fixed += 1
            cw, cd = row[cW-1].value, row[cD-1].value
            if (isinstance(cw, int) and isinstance(cd, int)
                    and not CARRY.search(blk) and row[cPTS-1].value != 2 * cw + cd):
                wdldiff.append((nm + ' [PTS]', row[cPTS-1].value, 2 * cw + cd))
                if WRITE:
                    row[cPTS-1].value = 2 * cw + cd
                    fixed += 1
    # PASS 2 (<2000): napáruj no-match přes ověřené V/R/P, převezmi GF:GA z PDF
    still = []
    for row, nm, gp in unmatched:
        done = False
        if yr < 2000:
            wdl = (row[cW-1].value, row[cD-1].value, row[cL-1].value)
            cands = [r for r in pdf if not r['used'] and r['gp'] == gp
                     and r['wdl'] == wdl]
            if len(cands) == 1:
                p = cands[0]
                gf, ga = row[cGF-1].value, row[cGA-1].value
                # věrohodný překlep: jedna půlka sedí přesně, nebo malý oboustranný rozdíl
                plausible = (p['gf'] == gf or p['ga'] == ga
                             or abs(p['gf'] - gf) + abs(p['ga'] - ga) <= 4)
                if plausible:
                    p['used'] = True
                    gafix.append((nm, f"{gf}:{ga}", f"{p['gf']}:{p['ga']}"))
                    if WRITE:
                        row[cGF-1].value, row[cGA-1].value = p['gf'], p['ga']
                        fixed += 1
                    done = True
        if not done:
            still.append((nm, gp, row[cGF-1].value, row[cGA-1].value))
    return still, wdldiff, gafix, fixed


def main():
    tot = {'nomatch': 0, 'fix': 0, 'wrote': 0}
    report = []
    for path in sorted(glob.glob('data/S*_FINAL.xlsx')):
        sid = re.search(r'S\d{4}_\d{2}', os.path.basename(path)).group(0)
        pdf = parse_pdf(sid)
        if pdf is None:
            continue
        yr = int(sid[1:5])
        wb = openpyxl.load_workbook(path)
        snm, swd, sga, sfx = [], [], [], 0
        for shname in wb.sheetnames:
            if not NAT_SHEET.match(shname) or REGIONAL.search(shname):
                continue
            nm, wd, ga, fx = process_sheet(wb[shname], pdf, yr)
            snm += [(shname,) + x for x in nm]
            swd += [(shname,) + x for x in wd]
            sga += [(shname,) + x for x in ga]
            sfx += fx
        if WRITE and sfx:
            wb.save(path)
        wb.close()
        tot['nomatch'] += len(snm); tot['fix'] += len(swd) + len(sga); tot['wrote'] += sfx
        if snm or swd or sga:
            report.append((sid, snm, swd, sga, sfx))

    for sid, snm, swd, sga, sfx in report:
        print(f"\n=== {sid[1:]} ===")
        for sh, nm, old, new in swd:
            print(f"   ~ V/R/P [{sh}] '{nm}': {old} -> PDF {new} {'✔' if WRITE else ''}")
        for sh, nm, old, new in sga:
            print(f"   ~ GF:GA [{sh}] '{nm}': {old} -> PDF {new} {'✔' if WRITE else ''}")
        for sh, nm, gp, gf, ga in snm:
            print(f"   ✗ k ověření [{sh}]: '{nm}' GP{gp} {gf}:{ga}")
    print(f"\n=== SOUHRN === k ověření: {tot['nomatch']} | "
          f"opraveno (V/R/P+PTS+GF/GA): {tot['fix']} | zapsáno: {tot['wrote']}"
          f"  (WRITE={WRITE})")


if __name__ == '__main__':
    main()
