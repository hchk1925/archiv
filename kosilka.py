#!/usr/bin/env python3
"""
kosilka.py — KOŠILKY soutěží.

Košilka = „karta" soutěže: jedno kanonické místo, kde je
  • seznam všech týmů soutěže (každý jednou) a jejich OSUD za sezónu,
  • seznam a pořadí FÁZÍ soutěže (základní část → play off → finále → o udržení…),
  • NÁVAZNOST x-1 ↔ x+1 (předchozí / následující sezóna).

Osud (final outcome) je jeden z:
  setrval · postup · sestup · zánik · reorganizace   (+ cíl u reorganizace/postupu/sestupu)

Fyzicky: v každém sešitu S….xlsx jeden list „KOSILKA", kde jsou pod sebou
košilky všech soutěží dané sezóny. Vícefázové soutěže = JEDNA soutěž (dle kořene
parent_node_id); jednotlivé fáze zůstávají jako podklad (P-řádky + původní tabulky).

Generuje se z existujících dat (SYSTEM + tabulky + CLUBS.prev_club_id). Pouští se:
    python kosilka.py             # přegeneruje KOSILKA ve všech sezónách
    python kosilka.py 1986_87     # jen jedna sezóna
"""
import collections
import glob
import os
import re
import sys

import openpyxl

NON_DATA = {'META', 'CLUBS', 'SYSTEM', 'NOTES', 'TODO', 'CHANGES', 'README', 'KOSILKA'}
PLAYOFF_T = {'playoff_round', 'playoff', 'final_group', 'final_series', 'series'}
RELEG_T = {'relegation_group', 'relegation_playoff', 'relegation_series'}
CLASS_T = {'classification'}
QUAL_T = {'qualification_group', 'qualification_series', 'qualification', 'baraz'}

KOSILKA_SHEET = 'KOSILKA'
FATES = ['setrval', 'postup', 'sestup', 'zánik', 'reorganizace']
KOS_COLS = ['row_type', 'comp_id', 'comp_name', 'level', 'phase_order', 'phase_name',
            'club_id', 'club_name', 'pos', 'fate', 'fate_target',
            'prev_club_id', 'next_club_id', 'note']


def _lvl(v):
    m = re.search(r'(\d+)', str(v or ''))
    return int(m.group(1)) if m else None


def _poskey(p):
    m = re.match(r'\d+', str(p) if p is not None else '')
    return int(m.group()) if m else 999


def _gpn(v):
    m = re.match(r'\d+', str(v) if v is not None else '')
    return int(m.group()) if m else 0


def _sibling_label(node):
    t = str(node.get('competition_type') or '')
    nm = str(node.get('name') or '').lower()
    if t in PLAYOFF_T or any(w in nm for w in
            ('play off', 'playoff', 'play-off', 'finále', 'čtvrtfinále',
             'semifinále', 'předkolo')):
        return 'play off'
    if t in RELEG_T or 'udržení' in nm or 'záchran' in nm:
        return 'o udržení'
    if t in CLASS_T or 'o umístění' in nm or re.search(r'\bo \d+\.', nm):
        return 'o umístění'
    if t in QUAL_T or 'kvalif' in nm or 'baráž' in nm or 'prolín' in nm:
        return 'prolínací' if ('baráž' in nm or 'prolín' in nm) else 'kvalifikace'
    return None


def _root_of(nodes, nid):
    seen = set()
    while nid in nodes and nodes[nid].get('parent') in nodes and nodes[nid].get('parent') not in seen:
        seen.add(nid)
        nid = nodes[nid]['parent']
    return nid


def read_season(path):
    """Vrátí (nodes, standings, clubs). standings[nid] = [{cid,pos,name,gp}].
    clubs[cid] = {name, prev}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    nodes = {}
    if 'SYSTEM' in wb.sheetnames:
        sr = list(wb['SYSTEM'].iter_rows(values_only=True)); SH = {c: j for j, c in enumerate(sr[0]) if c}
        for r in sr[1:]:
            nid = r[SH['node_id']] if 'node_id' in SH else None
            if not nid:
                continue
            g = lambda k: (r[SH[k]] if k in SH and SH[k] < len(r) else None)
            nodes[nid] = {'name': g('name') or '', 'level': g('level') or '',
                          'parent': g('parent_node_id'),
                          'competition_type': g('competition_type') or '',
                          'phase_order': g('phase_order')}
    clubs = {}
    if 'CLUBS' in wb.sheetnames:
        cw = list(wb['CLUBS'].iter_rows(values_only=True)); CH = {c: j for j, c in enumerate(cw[0]) if c}
        for r in cw[1:]:
            cid = r[CH['club_id']] if 'club_id' in CH else None
            if cid:
                clubs[cid] = {'name': r[CH.get('clean_name', -1)] if 'clean_name' in CH else '',
                              'prev': r[CH.get('prev_club_id', -1)] if 'prev_club_id' in CH else None}
    standings = collections.defaultdict(list)
    for sh in wb.sheetnames:
        if sh in NON_DATA:
            continue
        rr = list(wb[sh].iter_rows(values_only=True))
        if not rr:
            continue
        H = {c: j for j, c in enumerate(rr[0]) if c}
        if 'node_id' not in H or 'row_type' not in H:
            continue
        for r in rr[1:]:
            if r[H['row_type']] != 'T':
                continue
            nid = r[H['node_id']]
            if not nid:
                continue
            g = lambda k: (r[H[k]] if k in H else None)
            standings[nid].append({'cid': g('club_id'), 'pos': g('pos'),
                                   'name': g('club_name'), 'gp': g('GP')})
    wb.close()
    return nodes, standings, clubs


def competitions(nodes, standings):
    """Seskup fáze pod kořen. Vrátí list (root, ordered_phase_nodes)."""
    groups = collections.defaultdict(list)
    for nid in standings:
        if standings[nid]:
            groups[_root_of(nodes, nid)].append(nid)
    out = []
    for root, nids in groups.items():
        def pk(p, _root=root):
            po = nodes.get(p, {}).get('phase_order')
            try:
                pon = int(re.match(r'\d+', str(po)).group())
            except Exception:
                pon = 999
            return (0 if p == _root else 1, pon, _lvl(nodes.get(p, {}).get('level')) or 0,
                    str(nodes.get(p, {}).get('name', p)))
        out.append((root, sorted(nids, key=pk)))
    out.sort(key=lambda x: (_lvl(nodes.get(x[0], {}).get('level')) or 9999,
                            str(nodes.get(x[0], {}).get('name', ''))))
    return out


def roster(nodes, standings, root, phase_nodes):
    """Kanonický seznam týmů soutěže (každý jednou) s celkovým pořadím."""
    base = [p for p in phase_nodes if _sibling_label(nodes.get(p, {})) is None]
    if standings.get(root):
        primary = [root]
    elif base:
        primary = base
    else:
        primary = phase_nodes
    seen = {}; order = []
    for p in sorted(primary, key=lambda x: str(nodes.get(x, {}).get('name', x))):
        for r in sorted(standings[p], key=lambda x: _poskey(x['pos'])):
            key = r['cid'] if r['cid'] else f'_{id(r)}'
            if key in seen:
                if _gpn(r['gp']) > _gpn(seen[key]['gp']):
                    seen[key].update(r)
                continue
            seen[key] = dict(r); order.append(key)
    return [seen[k] for k in order]


def base_levels(nodes, standings):
    """{club_id: level kořenové soutěže ze ZÁKLADNÍ fáze}."""
    out = {}
    def consider(only_base):
        for nid, rows in standings.items():
            nd = nodes.get(nid, {})
            if only_base and _sibling_label(nd) is not None:
                continue
            lv = nodes.get(_root_of(nodes, nid), {}).get('level', '') or nd.get('level', '')
            for r in rows:
                cid = r['cid']
                if not cid:
                    continue
                cur = out.get(cid)
                if cur is None or (_lvl(lv) or 9999) < (_lvl(cur) or 9999):
                    out[cid] = lv
    consider(True)
    saved = dict(out)
    consider(False)
    for k in list(out):
        if k in saved:
            out[k] = saved[k]
    return out


def derive_fate(cid, cur_level, succ_ids, next_levels, last_season):
    """Osud z návaznosti. Bez nástupce = člověk doplní (zánik vs reorganizace)."""
    if succ_ids:
        nl = min((_lvl(next_levels.get(s)) or 9999) for s in succ_ids)
        cl = _lvl(cur_level) or 9999
        if nl < cl:
            return 'postup', next_levels.get(succ_ids[0], '')
        if nl > cl:
            return 'sestup', next_levels.get(succ_ids[0], '')
        return 'setrval', ''
    if last_season:
        return '', ''                     # poslední sezóna – budoucnost neznámá
    return '', ''                         # nenavázáno → k doplnění (zánik/reorganizace)


def build_one(path, prev_path, next_path, last_season):
    nodes, standings, clubs = read_season(path)
    next_clubs, next_lv = {}, {}
    if next_path and os.path.exists(next_path):
        n_nodes, n_st, next_clubs = read_season(next_path)
        next_lv = base_levels(n_nodes, n_st)
    succ_index = collections.defaultdict(list)        # prev_club_id -> [next club ids]
    for ncid, info in next_clubs.items():
        if info.get('prev'):
            succ_index[info['prev']].append(ncid)
    cur_lv = base_levels(nodes, standings)

    wb = openpyxl.load_workbook(path)
    if KOSILKA_SHEET in wb.sheetnames:
        del wb[KOSILKA_SHEET]
    ws = wb.create_sheet(KOSILKA_SHEET)
    ws.append(KOS_COLS)
    idx = {c: i for i, c in enumerate(KOS_COLS)}

    def row(**kw):
        r = [None] * len(KOS_COLS)
        for k, v in kw.items():
            r[idx[k]] = v
        ws.append(r)

    n_comp = n_team = 0
    for root, phase_nodes in competitions(nodes, standings):
        nd = nodes.get(root, {})
        row(row_type='C', comp_id=root, comp_name=nd.get('name', root), level=nd.get('level', ''))
        n_comp += 1
        for po, p in enumerate(phase_nodes, 1):
            pn = nodes.get(p, {})
            row(row_type='P', comp_id=root, phase_order=po, phase_name=pn.get('name', p),
                level=pn.get('level', ''))
        for i, t in enumerate(roster(nodes, standings, root, phase_nodes), 1):
            cid = t['cid']
            succ = succ_index.get(cid, [])
            fate, target = derive_fate(cid, nd.get('level', ''), succ, next_lv, last_season)
            row(row_type='T', comp_id=root, comp_name=nd.get('name', root),
                level=nd.get('level', ''), club_id=cid, club_name=t.get('name', ''),
                pos=(t['pos'] if t['pos'] not in (None, '') else i),
                fate=fate, fate_target=target,
                prev_club_id=clubs.get(cid, {}).get('prev'),
                next_club_id=(succ[0] if succ else None))
            n_team += 1
    wb.save(path)
    wb.close()
    return n_comp, n_team


def build_all(data_dir):
    paths = sorted(glob.glob(os.path.join(data_dir, 'S*_FINAL.xlsx')))
    sids = [re.search(r'S(\d{4}_\d{2})', os.path.basename(p)).group(1) for p in paths]
    tot_c = tot_t = 0
    for i, (sid, path) in enumerate(zip(sids, paths)):
        prev_path = paths[i - 1] if i > 0 else None
        next_path = paths[i + 1] if i < len(paths) - 1 else None
        last = (i == len(paths) - 1)
        c, t = build_one(path, prev_path, next_path, last)
        tot_c += c; tot_t += t
        print(f'  {sid}: {c} soutěží, {t} týmů')
    print(f'Hotovo. {len(paths)} sezón, {tot_c} košilek soutěží, {tot_t} týmů.')


def _find_data_dir():
    for c in ('xlsx', 'data', '.'):
        if glob.glob(os.path.join(c, 'S*_FINAL.xlsx')):
            return c
    return '.'


def main():
    dd = _find_data_dir()
    if len(sys.argv) > 1:
        sid = sys.argv[1]
        paths = sorted(glob.glob(os.path.join(dd, 'S*_FINAL.xlsx')))
        sids = [re.search(r'S(\d{4}_\d{2})', os.path.basename(p)).group(1) for p in paths]
        if sid not in sids:
            print('Neznámá sezóna:', sid); return
        i = sids.index(sid)
        c, t = build_one(paths[i], paths[i - 1] if i > 0 else None,
                         paths[i + 1] if i < len(paths) - 1 else None, i == len(paths) - 1)
        print(f'{sid}: {c} soutěží, {t} týmů.')
    else:
        build_all(dd)


if __name__ == '__main__':
    main()
