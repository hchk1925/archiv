#!/usr/bin/env python3
"""
test_desktop.py — headless smoke test desktopové appky (Tkinter).

Spusť pod virtuálním displejem:
    xvfb-run -a python3 test_desktop.py
Ověří: čisté helpery, sestavení okna, naplnění stromu, zobrazení tabulky +
komentáře u torza a zápis řešení do xlsx (s automatickým obnovením sešitu).
"""
import os
import sys

FAIL = []


def check(cond, msg):
    print(('  OK  ' if cond else ' FAIL ') + msg)
    if not cond:
        FAIL.append(msg)


def main():
    import desktop
    import app as core
    # v headless nesmí naskočit modální dialog (čekal by na klik)
    desktop.messagebox.showinfo = lambda *a, **k: None
    desktop.messagebox.showerror = lambda *a, **k: None
    desktop.messagebox.showwarning = lambda *a, **k: None

    # 1) čisté helpery bez GUI
    sid = next(s for s in core.list_season_ids()
               if core.load_season(s)['todo'])
    d = core.load_season(sid)
    check(isinstance(desktop.resolved_set(d), set), 'resolved_set vrací set')
    nt = desktop.node_todo_map(d)
    check(all(k.startswith('NODE_') for k in nt), 'node_todo_map klíčuje podle NODE_')

    # 2) sestavení okna
    win = desktop.AlmanachDesktop()
    win.update()
    check(len(win.item_map) > 0, 'strom naplněn položkami')

    # 3) vyber sezónu s torzem a najdi uzel s mezerou
    torzo_sid = next(s for s in core.list_season_ids()
                     if core.audit_notes_by_node(core.load_season(s)))
    win.season_var.set(torzo_sid)
    win.load_sid(torzo_sid)
    win.update()
    gap_node = next(iter(core.audit_notes_by_node(win.d)))
    win.show_node(gap_node)
    win.update()
    check(len(win.table.get_children()) > 0, 'tabulka týmů se zobrazila')
    txt = win.comment.get('1.0', 'end').strip()
    check('⚑' in txt or 'chybí' in txt or 'máme jen' in txt,
          'komentář „co chybí" se zobrazil u torza')
    check(win.cur_num is not None, 'u torza je předvyplněný formulář řešení')

    # 4) zápis řešení do xlsx + obnovení
    path = os.path.join(core.DATA, f'S{torzo_sid}_FINAL.xlsx')
    backup = open(path, 'rb').read()
    try:
        win.stav.set('ANO')
        win.varianta.set('doplnit týmy')
        win.pozn.delete(0, 'end')
        win.pozn.insert(0, 'DESKTOP-TEST')
        num = win.cur_num
        win.save()
        import openpyxl
        ws = openpyxl.load_workbook(path, read_only=True)['TODO']
        H = {c.value: j for j, c in enumerate(next(ws.iter_rows(max_row=1)), 0)}
        ok = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[H['#']]) == str(num):
                ok = (row[H['vyřešeno? (ANO/ne)']] == 'ANO'
                      and 'DESKTOP-TEST' in str(row[H['poznámka kolegy']]))
                break
        check(ok, 'Uložit do xlsx zapsalo stav + poznámku')
    finally:
        open(path, 'wb').write(backup)
    check(open(path, 'rb').read() == backup, 'sešit obnoven do původního stavu')

    # 5) PDF export builder (bez dialogu)
    pdf = core.build_full_pdf(win.d, torzo_sid)
    check(pdf[:5] == b'%PDF-', 'build_full_pdf z desktopu → PDF')

    # 6) ÚROVNĚ (matice úroveň × ČR/SK) — sestavení, přesun úrovně, zápis do SYSTEM
    import openpyxl
    check(desktop.region_of('Západoslovenský krajský přebor') == 'SK',
          'ČR/SK: Slovensko z názvu')
    check(desktop.region_of('Středočeský, 3. úroveň') == 'CZ',
          'ČR/SK: Čechy z názvu')
    check(desktop.chip_name('Středočeský, 3. úroveň (krajský přebor), skupina A')
          == 'Středočeský', 'štítek: bez úrovně a skupiny')
    osid = next(s for s in core.list_season_ids()
                if len({n['level'] for n in core.load_season(s)['system']}) > 1)
    od = core.load_season(osid)
    oc = desktop.OrgChart(win, osid, od)
    oc.update()
    check(len(oc.inner.winfo_children()) > 0, 'matice: buňky vykresleny')
    check(len(oc._levels()) > 1, 'matice: víc úrovní (řádků)')
    nid = next(iter(oc.model))
    cur = oc.model[nid]['level']
    other = next(l for l in oc._levels() if l != cur)
    oc.set_level([nid], other)
    check(oc.model[nid]['level'] == other, 'matice: změna úrovně v modelu')
    check(nid in oc._changes(), 'matice: změna zaevidována')
    opath = os.path.join(core.DATA, f'S{osid}_FINAL.xlsx')
    obackup = open(opath, 'rb').read()
    try:
        oc.save()
        ws = openpyxl.load_workbook(opath, read_only=True)['SYSTEM']
        H = {c.value: j for j, c in enumerate(next(ws.iter_rows(max_row=1)), 0)}
        got = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[H['node_id']] == nid:
                got = row[H['level']]
                break
        check(got == other, 'matice: úroveň zapsána do SYSTEM listu')
    finally:
        open(opath, 'wb').write(obackup)
    check(open(opath, 'rb').read() == obackup, 'matice: sešit obnoven')
    oc.destroy()

    win.destroy()
    print('\n' + ('VŠE OK' if not FAIL else f'SELHALO: {len(FAIL)}'))
    sys.exit(1 if FAIL else 0)


if __name__ == '__main__':
    main()
