#!/usr/bin/env python3
"""
reconcile_1948_49_kval.py — KVAL list (z Wikipedie) je duplikát divizní Finálové
skupiny (30DIV) — identické V-R-P i týmy, jen mylně označeno "kvalifikace o ligu"
a bez GF/GA. PDF žádnou samostatnou kvalifikaci o Státní ligu nemá.

Řešení (sloučení):
  - přepojí chain z duplicitních KVAL klubů na divizní id
    (0009→0058 Královo Pole, 0010→0052 Vítkovice, 0011→0033 Meteor ČB, 0012→0040 Pardubice),
  - odstraní list KVAL, CLUBS 0009–0012 a SYSTEM uzly 0002/0003,
  - doplní poznámku.
Idempotentní (když KVAL chybí, nic nedělá).
"""
import openpyxl

S48 = 'data/S1948_49_FINAL.xlsx'
S49 = 'data/S1949_50_FINAL.xlsx'
MAP = {  # KVAL duplikát → divizní club_id
    'CLUB_S1948_49_0009': 'CLUB_S1948_49_0058',  # Královo Pole
    'CLUB_S1948_49_0010': 'CLUB_S1948_49_0052',  # Vítkovice
    'CLUB_S1948_49_0011': 'CLUB_S1948_49_0033',  # Meteor ČB
    'CLUB_S1948_49_0012': 'CLUB_S1948_49_0040',  # Pardubice
}


def repoint_chain():
    wb = openpyxl.load_workbook(S49)
    ws = wb['CLUBS']
    hdr = [c.value for c in ws[1]]
    pi = hdr.index('prev_club_id')
    n = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, pi + 1).value
        if v in MAP:
            ws.cell(r, pi + 1).value = MAP[v]
            n += 1
    if n:
        wb.save(S49)
    print(f"  ✔ 1949/50: přepojeno {n} chain linků na divizní id")


def cleanup():
    wb = openpyxl.load_workbook(S48)
    if 'KVAL' not in wb.sheetnames:
        print("  · KVAL už není — nic")
        wb.close()
        return
    wb.remove(wb['KVAL'])
    # CLUBS: smaž 0009–0012
    cl = wb['CLUBS']
    ci = [c.value for c in cl[1]].index('club_id')
    for r in range(cl.max_row, 1, -1):
        if cl.cell(r, ci + 1).value in MAP:
            cl.delete_rows(r, 1)
    # SYSTEM: smaž uzly 0002/0003 (Kvalifikace o ligu)
    sy = wb['SYSTEM']
    for r in range(sy.max_row, 1, -1):
        if sy.cell(r, 1).value in (f'NODE_{S48[5:13]}_0002', f'NODE_{S48[5:13]}_0003'):
            sy.delete_rows(r, 1)
    # NOTES
    no = wb['NOTES']
    note = ('„Kvalifikace o ligu" z Wikipedie = divizní Finálová skupina (list 30DIV) '
            '— sloučeno (identické V-R-P, GF/GA doplněno z PDF). Vítěz Královo Pole '
            'postoupil do Státní ligy 1949/50. Samostatná kvalifikace o Státní ligu '
            'v PDF neexistuje.')
    if note not in {no.cell(r, 1).value for r in range(2, no.max_row + 1)}:
        no.append([note])
    # META
    me = wb['META']
    mk = {me.cell(r, 1).value: r for r in range(2, me.max_row + 1)}
    cnt = sum(1 for r in range(2, cl.max_row + 1) if cl.cell(r, 1).value)
    if 'total_clubs' in mk:
        me.cell(mk['total_clubs'], 2).value = str(cnt)
    wb.save(S48)
    print(f"  ✔ 1948/49: KVAL list + kluby 0009–0012 + uzly 0002/0003 odstraněny "
          f"(klubů nyní {cnt})")


if __name__ == '__main__':
    repoint_chain()
    cleanup()
