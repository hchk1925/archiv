#!/usr/bin/env python3
"""Show concrete examples of broken prev links, fate inconsistencies, missing links."""
import openpyxl, glob, os, re, sys
from collections import defaultdict

FILES = sorted(glob.glob('data/S*_FINAL.xlsx'))
NON_DATA = {'NOTES', 'SYSTEM', 'SERIES', 'CLUBS', 'META', 'PYRAMIDA'}


def load_clubs(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['CLUBS']
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0]); idx = {h: i for i, h in enumerate(hdr)}
    clubs = {}
    for r in rows[1:]:
        if not r or r[idx['club_id']] is None:
            continue
        clubs[r[idx['club_id']]] = {
            'clean_name': r[idx['clean_name']], 'sheet': r[idx['sheet']],
            'prev_club_id': r[idx['prev_club_id']],
            'city': r[idx['city']] if 'city' in idx else None}
    wb.close()
    return clubs


def load_fate_detail(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    cid_fate = defaultdict(lambda: defaultdict(set))  # cid -> fate -> {sheets}
    cid_name = {}
    for sh in wb.sheetnames:
        if sh in NON_DATA:
            continue
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hdr = list(rows[0]); idx = {h: i for i, h in enumerate(hdr)}
        if 'club_id' not in idx or 'season_fate' not in idx:
            continue
        for r in rows[1:]:
            if not r:
                continue
            cid = r[idx['club_id']]; fate = r[idx['season_fate']]
            nm = r[idx['club_name']] if 'club_name' in idx else None
            if cid and fate:
                cid_fate[cid][str(fate).strip()].add(sh)
                if nm:
                    cid_name[cid] = nm
    wb.close()
    return cid_fate, cid_name


mode = sys.argv[1] if len(sys.argv) > 1 else 'broken'

if mode == 'broken':
    prev = None
    for path in FILES:
        clubs = load_clubs(path)
        if prev is not None:
            prev_ids = set(prev.keys())
            for cid, c in clubs.items():
                p = c['prev_club_id']
                if p and p not in prev_ids:
                    print(f"{os.path.basename(path)[1:8]} BROKEN {cid} '{c['clean_name']}' "
                          f"[{c['sheet']}] -> {p} (not in prev season)")
        prev = clubs

elif mode == 'inconst':
    for path in FILES:
        cid_fate, cid_name = load_fate_detail(path)
        bad = {cid: d for cid, d in cid_fate.items() if len(d) > 1}
        if bad:
            print(f"\n=== {os.path.basename(path)[1:8]}  ({len(bad)} inconsistent) ===")
            for cid, d in sorted(bad.items())[:8]:
                parts = []
                for fate, shs in d.items():
                    parts.append(f"{fate}({','.join(sorted(shs))})")
                print(f"  {cid} '{cid_name.get(cid,'?')}': " + " | ".join(parts))

elif mode == 'missing':
    prev = None
    for path in FILES:
        clubs = load_clubs(path)
        if prev is not None:
            pbyname = defaultdict(list)
            for pid, pc in prev.items():
                pbyname[str(pc['clean_name']).strip().lower()].append((pid, pc))
            for cid, c in clubs.items():
                if c['prev_club_id']:
                    continue
                nm = str(c['clean_name']).strip().lower()
                if nm in pbyname:
                    cands = pbyname[nm]
                    print(f"{os.path.basename(path)[1:8]} MISS {cid} '{c['clean_name']}' "
                          f"[{c['sheet']}/{c['city']}] -> cand: "
                          + "; ".join(f"{p} [{pc['sheet']}/{pc['city']}]" for p, pc in cands))
        prev = clubs
