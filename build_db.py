#!/usr/bin/env python3
"""
build_db.py — Konsolidace 65 sezónních workbooků do jedné databáze.

Vstup:  data/S*_FINAL.xlsx
Výstup: almanach.sqlite  + CSV exporty do export/

Tabulky:
  seasons        — z META (1 řádek / sezóna)
  clubs          — CLUBS všech sezón + odvozený chain_id (řetěz prev_club_id)
  competitions   — SYSTEM (pyramida soutěží) všech sezón
  standings      — T-řádky (tabulky) ze všech datových listů
  registrations  — R-řádky (registrované kluby bez tabulky)
  series         — list SERIES
  notes          — list NOTES

Re-runnable: DB se vždy přestaví od nuly.
"""
import openpyxl, glob, os, re, sqlite3, csv, sys
from collections import defaultdict

DATA = sys.argv[1] if len(sys.argv) > 1 else 'data'
DBF = 'almanach.sqlite'
EXPORT = 'export'
NON_DATA = {'NOTES', 'SYSTEM', 'SERIES', 'CLUBS', 'META', 'PYRAMIDA'}


def season_of(path):
    m = re.search(r'S(\d{4})_(\d{2})', os.path.basename(path))
    return f"S{m.group(1)}_{m.group(2)}"


def hidx(row0):
    return {h: i for i, h in enumerate(row0) if h is not None}


def cell(row, H, name):
    i = H.get(name)
    return row[i] if i is not None and i < len(row) else None


def main():
    files = sorted(glob.glob(os.path.join(DATA, 'S*_FINAL.xlsx')))
    if not files:
        print("Žádné sezónní soubory.")
        return
    if os.path.exists(DBF):
        os.remove(DBF)
    db = sqlite3.connect(DBF)
    c = db.cursor()
    c.executescript("""
    CREATE TABLE seasons(season_id TEXT PRIMARY KEY, season_label TEXT,
        era_note TEXT, scoring_system TEXT, source TEXT, prev_season TEXT,
        total_clubs INT, total_nodes INT, note TEXT);
    CREATE TABLE clubs(season_id TEXT, club_id TEXT, clean_name TEXT,
        raw_name TEXT, sheet TEXT, level TEXT, district TEXT,
        prev_club_id TEXT, change_note TEXT, city TEXT,
        chain_id TEXT, chain_seq INT,
        PRIMARY KEY(season_id, club_id));
    CREATE TABLE competitions(season_id TEXT, node_id TEXT, name TEXT,
        competition_type TEXT, level TEXT, region TEXT, parent_node_id TEXT,
        feeds_into TEXT, feeds_into_loser TEXT, scoring TEXT, status TEXT,
        note TEXT, prev_node_id TEXT, entry TEXT, phase_order INT,
        PRIMARY KEY(season_id, node_id));
    CREATE TABLE standings(season_id TEXT, sheet TEXT, node_id TEXT,
        club_id TEXT, pos TEXT, club_name TEXT, GP INT, W INT, D INT, L INT,
        GF INT, GA INT, PTS INT, season_fate TEXT, comp_path TEXT,
        level TEXT, prev_club_id TEXT, tr_id TEXT, district TEXT);
    CREATE TABLE registrations(season_id TEXT, sheet TEXT, node_id TEXT,
        club_id TEXT, club_name TEXT, level TEXT, prev_club_id TEXT,
        district TEXT, comp_path TEXT);
    CREATE TABLE series(season_id TEXT, series_id TEXT, node_id TEXT,
        club_id TEXT, raw_name TEXT, clean_name TEXT, side TEXT,
        game_scores TEXT, series_score TEXT, dest_node_id TEXT, note TEXT);
    CREATE TABLE notes(season_id TEXT, node_id TEXT, sheet TEXT,
        note_text TEXT, source_type TEXT, fake_club_ids TEXT);
    """)

    clubs_rows = []          # pro chain výpočet
    for path in files:
        sid = season_of(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

        # META
        meta = {}
        if 'META' in wb.sheetnames:
            for r in wb['META'].iter_rows(values_only=True):
                if r and r[0]:
                    meta[str(r[0])] = r[1]
        c.execute("INSERT OR REPLACE INTO seasons VALUES(?,?,?,?,?,?,?,?,?)", (
            sid, meta.get('season_label'), meta.get('era_note'),
            meta.get('scoring_system'), meta.get('source'),
            meta.get('prev_season'),
            int(meta['total_clubs']) if str(meta.get('total_clubs','')).isdigit() else None,
            int(meta['total_nodes']) if str(meta.get('total_nodes','')).isdigit() else None,
            meta.get('note')))

        # CLUBS
        ws = wb['CLUBS']
        rows = list(ws.iter_rows(values_only=True))
        H = hidx(rows[0])
        for r in rows[1:]:
            if not r or cell(r, H, 'club_id') is None:
                continue
            rec = (sid, cell(r, H, 'club_id'), cell(r, H, 'clean_name'),
                   cell(r, H, 'raw_name'), cell(r, H, 'sheet'),
                   cell(r, H, 'level'), cell(r, H, 'district'),
                   cell(r, H, 'prev_club_id'), cell(r, H, 'change_note'),
                   cell(r, H, 'city'))
            clubs_rows.append(rec)

        # SYSTEM
        if 'SYSTEM' in wb.sheetnames:
            ws = wb['SYSTEM']; rr = list(ws.iter_rows(values_only=True))
            SH = hidx(rr[0])
            for r in rr[1:]:
                if not r or cell(r, SH, 'node_id') is None:
                    continue
                ph = cell(r, SH, 'phase_order')
                c.execute("INSERT OR REPLACE INTO competitions VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
                    sid, cell(r, SH, 'node_id'), cell(r, SH, 'name'),
                    cell(r, SH, 'competition_type'), cell(r, SH, 'level'),
                    cell(r, SH, 'region'), cell(r, SH, 'parent_node_id'),
                    cell(r, SH, 'feeds_into'), cell(r, SH, 'feeds_into_loser'),
                    cell(r, SH, 'scoring'), cell(r, SH, 'status'),
                    cell(r, SH, 'note'), cell(r, SH, 'prev_node_id'),
                    cell(r, SH, 'entry'),
                    int(ph) if str(ph).strip().isdigit() else None))

        # SERIES
        if 'SERIES' in wb.sheetnames:
            ws = wb['SERIES']; rr = list(ws.iter_rows(values_only=True))
            if rr:
                ZH = hidx(rr[0])
                for r in rr[1:]:
                    if not r or cell(r, ZH, 'series_id') is None:
                        continue
                    c.execute("INSERT INTO series VALUES(?,?,?,?,?,?,?,?,?,?,?)", (
                        sid, cell(r, ZH, 'series_id'), cell(r, ZH, 'node_id'),
                        cell(r, ZH, 'club_id'), cell(r, ZH, 'raw_name'),
                        cell(r, ZH, 'clean_name'), cell(r, ZH, 'side'),
                        cell(r, ZH, 'game_scores'), cell(r, ZH, 'series_score'),
                        cell(r, ZH, 'dest_node_id'), cell(r, ZH, 'note')))

        # NOTES
        if 'NOTES' in wb.sheetnames:
            ws = wb['NOTES']; rr = list(ws.iter_rows(values_only=True))
            if rr:
                NH = hidx(rr[0])
                for r in rr[1:]:
                    if not r or all(x is None for x in r):
                        continue
                    c.execute("INSERT INTO notes VALUES(?,?,?,?,?,?)", (
                        sid, cell(r, NH, 'node_id'), cell(r, NH, 'sheet'),
                        cell(r, NH, 'note_text'), cell(r, NH, 'source_type'),
                        cell(r, NH, 'fake_club_ids')))

        # datové listy → standings / registrations
        for sh in wb.sheetnames:
            if sh in NON_DATA:
                continue
            ws = wb[sh]; rr = list(ws.iter_rows(values_only=True))
            if not rr:
                continue
            DH = hidx(rr[0])
            if 'row_type' not in DH or 'club_id' not in DH:
                continue
            for r in rr[1:]:
                if not r:
                    continue
                rt = cell(r, DH, 'row_type')
                cid = cell(r, DH, 'club_id')
                if rt == 'T' and cid:
                    def num(x):
                        try: return int(x)
                        except (TypeError, ValueError): return None
                    c.execute("INSERT INTO standings VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
                        sid, sh, cell(r, DH, 'node_id'), cid,
                        cell(r, DH, 'pos'), cell(r, DH, 'club_name'),
                        num(cell(r, DH, 'GP')), num(cell(r, DH, 'W')),
                        num(cell(r, DH, 'D')), num(cell(r, DH, 'L')),
                        num(cell(r, DH, 'GF')), num(cell(r, DH, 'GA')),
                        num(cell(r, DH, 'PTS')), cell(r, DH, 'season_fate'),
                        cell(r, DH, 'comp_path'), cell(r, DH, 'level'),
                        cell(r, DH, 'prev_club_id'), cell(r, DH, 'tr_id'),
                        cell(r, DH, 'district')))
                elif rt == 'R' and cid:
                    c.execute("INSERT INTO registrations VALUES(?,?,?,?,?,?,?,?,?)", (
                        sid, sh, cell(r, DH, 'node_id'), cid,
                        cell(r, DH, 'club_name'), cell(r, DH, 'level'),
                        cell(r, DH, 'prev_club_id'), cell(r, DH, 'district'),
                        cell(r, DH, 'comp_path')))
        wb.close()

    # ── chain_id: zpětný řetěz prev_club_id (cesta, ne strom) ──
    # prev_club_id je zpětná funkce (≤1 předchůdce). Dopředu ale větví:
    # A-tým + B-tým + chybně slinkovaný klub míří na stejného předchůdce.
    # Dle D20 (B-tým = vlastní identita) zdědí chain JEN „primární nástupce"
    # předchůdce; ostatní sourozenci zakládají vlastní chain.
    B_SUFFIX = re.compile(r'(\bB\b|\bII+\b|\bIII\b|\bIV\b|\sC)\s*$')

    valid = {(s, cid) for (s, cid, *_rest) in clubs_rows}
    info = {}
    for (s, cid, cn, rn, sh, lv, di, pcid, chg, cty) in clubs_rows:
        m = re.match(r'L(\d+)', str(lv or ''))
        info[(s, cid)] = {
            'prev': pcid, 'name': str(cn or ''),
            'lvl': int(m.group(1)) if m else 9999,
        }

    def prev_key(key):
        p = info[key]['prev']
        if not p:
            return None
        pm = re.search(r'CLUB_(S\d{4}_\d{2})_', str(p))
        ps = pm.group(1) if pm else None
        return (ps, p) if ps and (ps, p) in valid else None

    # děti každého předchůdce → urči primárního nástupce
    children = defaultdict(list)
    for key in valid:
        pk = prev_key(key)
        if pk is not None:
            children[pk].append(key)

    primary_of_parent = {}
    for pk, kids in children.items():
        non_b = [k for k in kids if not B_SUFFIX.search(info[k]['name'])]
        pool = non_b or kids
        pool.sort(key=lambda k: (info[k]['lvl'], k[1]))
        primary_of_parent[pk] = pool[0]

    # přiřazení chainů odzadu (nejstarší sezóna první)
    chain_of = {}
    chain_seq = defaultdict(int)
    enriched = []
    for rec in sorted(clubs_rows, key=lambda x: (x[0], x[1])):
        s, cid, cn, rn, sh, lv, di, pcid, chn, cty = rec
        key = (s, cid)
        pk = prev_key(key)
        if pk is not None and pk in chain_of and primary_of_parent.get(pk) == key:
            chain_id = chain_of[pk]            # primární nástupce dědí
        else:
            chain_id = f"CHAIN_{cid}"           # vlastní (B-tým / větev / kořen)
        chain_of[key] = chain_id
        chain_seq[chain_id] += 1
        enriched.append(rec + (chain_id, chain_seq[chain_id]))

    c.executemany("INSERT OR REPLACE INTO clubs VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                  enriched)
    c.executescript("""
    CREATE INDEX ix_clubs_chain ON clubs(chain_id);
    CREATE INDEX ix_clubs_prev  ON clubs(prev_club_id);
    CREATE INDEX ix_clubs_name  ON clubs(clean_name);
    CREATE INDEX ix_std_club    ON standings(season_id, club_id);
    CREATE INDEX ix_std_node    ON standings(season_id, node_id);
    CREATE INDEX ix_comp_parent ON competitions(season_id, parent_node_id);
    """)

    # ── pohledy ──
    c.executescript("""
    CREATE VIEW v_club_history AS
      SELECT chain_id, season_id, club_id, clean_name, level, city,
             prev_club_id, chain_seq
      FROM clubs ORDER BY chain_id, season_id;
    CREATE VIEW v_promotions AS
      SELECT season_id, club_id, club_name, sheet, level, season_fate
      FROM standings WHERE season_fate IN ('postup','sestup','reorganizace',
             'slouceni','zanik');
    CREATE VIEW v_pyramid AS
      SELECT season_id, node_id, name, competition_type, level,
             parent_node_id, feeds_into, feeds_into_loser
      FROM competitions ORDER BY season_id, level, node_id;
    """)
    db.commit()

    # ── CSV exporty ──
    os.makedirs(EXPORT, exist_ok=True)
    for tbl in ['seasons', 'clubs', 'competitions', 'standings',
                'registrations', 'series', 'notes']:
        cur = c.execute(f"SELECT * FROM {tbl}")
        cols = [d[0] for d in cur.description]
        with open(os.path.join(EXPORT, f"{tbl}.csv"), 'w', newline='',
                  encoding='utf-8') as fh:
            w = csv.writer(fh)
            w.writerow(cols)
            w.writerows(cur.fetchall())

    # statistiky
    def q1(sql):
        return c.execute(sql).fetchone()[0]
    stats = {
        'sezón': q1("SELECT COUNT(*) FROM seasons"),
        'klub-záznamů': q1("SELECT COUNT(*) FROM clubs"),
        'unikátních řetězů (chain_id)': q1("SELECT COUNT(DISTINCT chain_id) FROM clubs"),
        'soutěžních uzlů': q1("SELECT COUNT(*) FROM competitions"),
        'tabulkových řádků (standings)': q1("SELECT COUNT(*) FROM standings"),
        'registrací (R)': q1("SELECT COUNT(*) FROM registrations"),
        'série': q1("SELECT COUNT(*) FROM series"),
        'poznámek': q1("SELECT COUNT(*) FROM notes"),
    }
    longest = c.execute("""SELECT chain_id,
        COUNT(DISTINCT season_id) ns, MIN(season_id), MAX(season_id),
        (SELECT clean_name FROM clubs c2 WHERE c2.chain_id=c1.chain_id
         ORDER BY season_id DESC LIMIT 1) last_name
        FROM clubs c1 GROUP BY chain_id ORDER BY ns DESC LIMIT 8""").fetchall()
    maxlen = q1("SELECT MAX(n) FROM (SELECT chain_id,"
                "COUNT(DISTINCT season_id) n FROM clubs GROUP BY chain_id)")
    stats['nejdelší řetěz (sezón)'] = maxlen
    db.close()

    print("=== Konsolidovaná DB hotova:", DBF, "===")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    print("\nNejdelší řetězy klubů (kontinuita napříč sezónami):")
    for ch, n, s0, s1, nm in longest:
        print(f"  {n:>2} sezón  {s0}–{s1}  {nm}  [{ch}]")
    print(f"\nCSV exporty v {EXPORT}/  (7 souborů)")


if __name__ == '__main__':
    main()
